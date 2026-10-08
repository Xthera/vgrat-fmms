#!/usr/bin/env python3

"""
VGrat FMS - BUILD FUNDS DATA

Builds:

    data/funds.json
    data/bid_history.json

MASTER UNIVERSE
===============

Funds Links.xlsm
    Column A = Prudential URL
    Column B = exact PruAccess Fund Name

Every populated Column A row becomes one fund record.

IMPORTANT:
    This script is located in the repository root.

Therefore:

    ROOT = Path(__file__).resolve().parent

NOT:

    Path(__file__).resolve().parent.parent


RESEARCH INPUT
==============

Research Funds.xlsx is already researched upstream.

This script DOES NOT perform AI research, web research, Morningstar
research, manager research, or any other research.

Research Funds.xlsx contains:

    Worksheet 1:
        Fund Research
        Per-fund researched Geography/Sector values.

    Worksheet 2:
        Authoritative Geography master categories.

    Worksheet 3:
        Authoritative Sector master categories.

The worksheet order is preserved as the authoritative source.

The master Geography/Sector category lists are published into the
top-level "research.masterCategories" block in funds.json.

Per-fund research is published as:

    "research": {
        "geographic1": "...",
        "geographic2": "...",
        "sector1": "...",
        "sector2": "..."
    }


HOLDINGS RESOLUTION
===================

Stage order:

    baseline
        ->
    recovery1
        ->
    recovery2
        ->
    recovery3

Rules:

- Baseline success is retained.
- Baseline no_holdings_section is retained.
- Failed baseline funds may be resolved by Recovery 1.
- Remaining failed funds may be resolved by Recovery 2.
- Remaining failed funds may be resolved by Recovery 3.
- Nothing is inferred or fabricated.
- Recovery stages never replace an already-resolved fund.


BID HISTORY
===========

Read from:

    output_pruaccess/funds/<row>_<id>/

Expected files:

    prudential_fund.json
    bid_history.json

Historical observations are copied from the extracted PruAccess data.

No raw-date carry-forward is performed here.


DIVIDEND
========

The Prudential source record is authoritative.

Rules:

    dividendRate non-empty
        -> hasDividend = true

    dividendRate empty/missing
        -> hasDividend = false

The existing dividendRate string is preserved.

dividendHistory is published alongside it: the same
dividendRate string parsed into a list, newest first:

    [{"exDate": "2026-09-30", "rate": 1.25}, ...]

dividendUnit:
    Extracted by scripts/test_pruaccess.py from the Prudential
    source (API unit field, otherwise the unit printed next to
    the latest payout on the fund page) and published with:

        dividendUnitSource    "prudential-api:<field>" |
                              "prudential-page" | "default"
        dividendUnitEvidence  the source text it was read from

    Only when Prudential shows no unit is DEFAULT_DIVIDEND_UNIT
    applied, and dividendUnitSource is then "default".

No dividend calculation is performed.


TEXT CLEAN-UP
=============

fundObjective is published as plain text. Some source records
contain escaped HTML (e.g. "<p>... &quot;ADRs&quot; ...</p>");
tags are removed and entities decoded.

Holdings: some factsheet parses place the holding weight at the
end of the name ("SK HYNIX INC 6.7%") and record a different,
incorrect weightPercent. When a name ends with a percentage, that
percentage is the published weight: it is removed from the name
and used as weightPercent / weightText.


FAILED FUNDS
============

A fund that fails this run (Prudential metadata, holdings or BID
history) no longer stops the build:

    - successful funds are published with this run's data;
    - a failed fund keeps its last good record from the previous
      data/funds.json and data/bid_history.json, with
      updateStatus.status = "stale" and the reasons;
    - a failed fund with no earlier good data is left out;
    - when only the holdings failed, the fund's details and prices are
      still published, with holdings from the last good run (or none);
    - failures are listed in the log and the GitHub Actions run
      summary, and in funds.json summary.failedFunds.

ALLOW_UNRESOLVED=1 publishes the partial records instead (old option).


OUTPUT
======

data/funds.json
data/bid_history.json

Both files contain the same master fund identity block and fund count.
CLOSED FUNDS

    Closed Funds.xlsx (optional) lists funds Prudential has closed.
    scripts/test_pruaccess.py saves their past prices to
    output_pruaccess/closed/<code>/; they are published with
    closed = true, shown as Closed on the site, and never ranked.

"""

from __future__ import annotations

import json
import html
import os
import re
import sys
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


# =============================================================================
# CONFIGURATION
# =============================================================================

# IMPORTANT:
# build_funds_data.py is in the repository root.
ROOT = Path(__file__).resolve().parent

# ---------------------------------------------------------------------------
# MASTER INPUTS
# ---------------------------------------------------------------------------

EXCEL_FILE = ROOT / "Funds Links.xlsm"

RESEARCH_EXCEL_FILE = ROOT / "Research Funds.xlsx"

RESEARCH_FUND_SHEET_NAME = "Fund Research"

# Worksheet 2 = Geography master categories
RESEARCH_GEOGRAPHY_SHEET_INDEX = 1

# Worksheet 3 = Sector master categories
RESEARCH_SECTOR_SHEET_INDEX = 2

# ---------------------------------------------------------------------------
# HOLDINGS
# ---------------------------------------------------------------------------

BASELINE_FILE = ROOT / "output_holdings" / "all_holdings.json"

RECOVERY_STAGES = [
    (
        "recovery1",
        ROOT / "output_holdings_recovery",
    ),
    (
        "recovery2",
        ROOT / "output_holdings_recovery_2",
    ),
    (
        "recovery3",
        ROOT / "output_holdings_recovery_3",
    ),
]

# ---------------------------------------------------------------------------
# PRUACCESS
# ---------------------------------------------------------------------------

PRUACCESS_FUNDS_DIR = (
    ROOT / "output_pruaccess" / "funds"
)

PRUACCESS_PRODUCTION_DIR = (
    ROOT / "output_pruaccess" / "production"
)

# ---------------------------------------------------------------------------
# OUTPUT
# ---------------------------------------------------------------------------

DATA_DIR = ROOT / "data"

FUNDS_OUT = DATA_DIR / "funds.json"

BID_OUT = DATA_DIR / "bid_history.json"

# ---------------------------------------------------------------------------
# GENERAL
# ---------------------------------------------------------------------------

SCHEMA_VERSION = 1

MAX_HOLDINGS = 10

# Applied when the Prudential source publishes no dividendUnit.
DEFAULT_DIVIDEND_UNIT = "% per payout"

ALLOW_UNRESOLVED = (
    os.environ.get(
        "ALLOW_UNRESOLVED",
        "",
    )
    .strip()
    .lower()
    in {
        "1",
        "true",
        "yes",
    }
)


# =============================================================================
# BASIC HELPERS
# =============================================================================

def clean_text(value: Any) -> str:
    """
    Normalize whitespace while preserving the actual value content.
    """

    if value is None:
        return ""

    return re.sub(
        r"\s+",
        " ",
        str(value).replace(
            "\xa0",
            " ",
        ),
    ).strip()


def clean_raw_text(value: Any) -> str:
    """
    Preserve structured strings such as dividendRate while removing
    surrounding whitespace.

    Example:

        20260331=2.50&20260228=2.50

    remains exactly that string.
    """

    if value is None:
        return ""

    return str(value).strip()


def clean_html_text(value: Any) -> str:
    """
    Convert (possibly escaped) HTML to plain text.

        "<p>Invest in ADRs (&quot;ADRs&quot;)</p>"
            -> 'Invest in ADRs ("ADRs")'
    """

    if value is None:
        return ""

    text = str(value)

    # Decode up to twice: some records are double-escaped.
    for _ in range(2):
        decoded = html.unescape(text)
        if decoded == text:
            break
        text = decoded

    text = re.sub(
        r"<br\s*/?>|</p\s*>",
        " ",
        text,
        flags=re.IGNORECASE,
    )

    text = re.sub(
        r"<[^>]+>",
        "",
        text,
    )

    return clean_text(
        text
    )


def utc_now_iso() -> str:
    return (
        datetime.now(
            timezone.utc
        )
        .replace(
            microsecond=0
        )
        .isoformat()
    )


def load_json(path: Path) -> Any:
    try:
        return json.loads(
            path.read_text(
                encoding="utf-8"
            )
        )
    except Exception as error:
        raise RuntimeError(
            f"Could not read {path}: {error}"
        ) from error


def write_json_atomic(
    path: Path,
    data: Any,
    compact: bool = False,
) -> None:

    path.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    temporary = path.with_suffix(
        path.suffix + ".tmp"
    )

    if compact:
        text = json.dumps(
            data,
            ensure_ascii=False,
            separators=(
                ",",
                ":",
            ),
        )
    else:
        text = json.dumps(
            data,
            ensure_ascii=False,
            indent=2,
        )

    temporary.write_text(
        text + "\n",
        encoding="utf-8",
    )

    temporary.replace(path)


def normalize_key(value: Any) -> str:
    return re.sub(
        r"[^a-z0-9]+",
        "",
        clean_text(value).lower(),
    )


# =============================================================================
# EXCEL MASTER UNIVERSE
# =============================================================================

def read_excel_funds() -> dict[int, dict]:
    """
    Read Funds Links.xlsm.

    Column A:
        Prudential URL

    Column B:
        exact PruAccess Fund Name

    Column A controls the master universe.
    """

    if not EXCEL_FILE.exists():
        raise FileNotFoundError(
            f"Excel file not found: {EXCEL_FILE}"
        )

    print(
        f"Master workbook: {EXCEL_FILE}"
    )

    workbook = load_workbook(
        EXCEL_FILE,
        read_only=True,
        keep_vba=True,
        data_only=True,
    )

    funds: dict[int, dict] = {}

    try:
        worksheet = workbook.active

        for row in range(
            2,
            worksheet.max_row + 1,
        ):

            prudential_url = clean_text(
                worksheet.cell(
                    row=row,
                    column=1,
                ).value
            )

            pruaccess_name = clean_text(
                worksheet.cell(
                    row=row,
                    column=2,
                ).value
            )

            if not prudential_url:
                continue

            funds[row] = {
                "excelRow": row,
                "prudentialUrl": prudential_url,
                "pruAccessName": pruaccess_name,
            }

    finally:
        workbook.close()

    if not funds:
        raise RuntimeError(
            "No populated URLs were found in "
            "Funds Links.xlsm Column A."
        )

    print(
        f"Excel fund universe: {len(funds)}"
    )

    return funds


# =============================================================================
# RESEARCH FUNDS.XLSX
# =============================================================================

def find_header_row(
    worksheet,
    required_fields: set[str],
    max_scan_rows: int = 20,
) -> int | None:
    """
    Locate a header row by normalized column names.
    """

    max_row = min(
        worksheet.max_row,
        max_scan_rows,
    )

    for row in range(
        1,
        max_row + 1,
    ):

        headers = {
            normalize_key(
                worksheet.cell(
                    row=row,
                    column=column,
                ).value
            )
            for column in range(
                1,
                worksheet.max_column + 1,
            )
        }

        headers.discard("")

        if required_fields.issubset(
            headers
        ):
            return row

    return None


def build_header_map(
    worksheet,
    header_row: int,
) -> dict[str, int]:
    result = {}

    for column in range(
        1,
        worksheet.max_column + 1,
    ):

        value = worksheet.cell(
            row=header_row,
            column=column,
        ).value

        key = normalize_key(value)

        if key:
            result[key] = column

    return result


def first_matching_column(
    header_map: dict[str, int],
    candidates: list[str],
) -> int | None:

    for candidate in candidates:

        key = normalize_key(
            candidate
        )

        if key in header_map:
            return header_map[key]

    return None


def find_fund_research_columns(
    worksheet,
) -> tuple[int, dict[str, int]]:
    """
    Identify the Fund Research worksheet header and columns.

    Supported semantic names are deliberately broad because the
    workbook may use slightly different display labels.
    """

    possible_fund_names = {
        "fundname",
        "fund",
        "fundtitle",
        "prulinkfundname",
        "prudentialfundname",
        "fundresearchname",
        "fundidentifier",
    }

    header_row = find_header_row(
        worksheet,
        {
            "fundname"
        },
    )

    if header_row is None:

        for row in range(
            1,
            min(
                worksheet.max_row,
                20,
            )
            + 1,
        ):

            header_map = build_header_map(
                worksheet,
                row,
            )

            if (
                first_matching_column(
                    header_map,
                    list(
                        possible_fund_names
                    ),
                )
                is not None
            ):
                header_row = row
                break

    if header_row is None:
        raise RuntimeError(
            "Could not identify the header row "
            "in Research Funds.xlsx / Fund Research."
        )

    header_map = build_header_map(
        worksheet,
        header_row,
    )

    fund_column = first_matching_column(
        header_map,
        [
            "Fund Name",
            "Fund",
            "Fund Title",
            "PRULink Fund Name",
            "Prudential Fund Name",
            "Fund Identifier",
        ],
    )

    if fund_column is None:
        raise RuntimeError(
            "Could not identify the fund-name column "
            "in Research Funds.xlsx / Fund Research."
        )

    columns = {
        "fund": fund_column,
    }

    geography_candidates = {
        "geographic1": [
            "Geographic 1",
            "Geography 1",
            "Geographic1",
            "Geography1",
            "Region 1",
            "Geographic Exposure 1",
            "Geography Exposure 1",
        ],
        "geographic2": [
            "Geographic 2",
            "Geography 2",
            "Geographic2",
            "Geography2",
            "Region 2",
            "Geographic Exposure 2",
            "Geography Exposure 2",
        ],
        "sector1": [
            "Sector 1",
            "Sector1",
            "Top Sector 1",
            "Sector Exposure 1",
        ],
        "sector2": [
            "Sector 2",
            "Sector2",
            "Top Sector 2",
            "Sector Exposure 2",
        ],
    }

    for field, candidates in geography_candidates.items():

        column = first_matching_column(
            header_map,
            candidates,
        )

        if column is not None:
            columns[field] = column

    return (
        header_row,
        columns,
    )


def read_fund_research() -> tuple[
    dict[str, dict],
    str,
]:
    """
    Read already-researched per-fund values.

    No research is performed here.
    """

    if not RESEARCH_EXCEL_FILE.exists():
        raise FileNotFoundError(
            f"Research workbook not found: "
            f"{RESEARCH_EXCEL_FILE}"
        )

    workbook = load_workbook(
        RESEARCH_EXCEL_FILE,
        read_only=True,
        data_only=True,
    )

    try:

        if (
            RESEARCH_FUND_SHEET_NAME
            not in workbook.sheetnames
        ):
            raise RuntimeError(
                "Research workbook does not contain "
                f"worksheet '{RESEARCH_FUND_SHEET_NAME}'. "
                f"Available sheets: {workbook.sheetnames}"
            )

        worksheet = workbook[
            RESEARCH_FUND_SHEET_NAME
        ]

        (
            header_row,
            columns,
        ) = find_fund_research_columns(
            worksheet
        )

        research: dict[str, dict] = {}

        for row in range(
            header_row + 1,
            worksheet.max_row + 1,
        ):

            fund_name = clean_text(
                worksheet.cell(
                    row=row,
                    column=columns["fund"],
                ).value
            )

            if not fund_name:
                continue

            record = {
                "fundName": fund_name,
                "geographic1": None,
                "geographic2": None,
                "sector1": None,
                "sector2": None,
            }

            for field in (
                "geographic1",
                "geographic2",
                "sector1",
                "sector2",
            ):

                column = columns.get(
                    field
                )

                if column is None:
                    continue

                value = clean_text(
                    worksheet.cell(
                        row=row,
                        column=column,
                    ).value
                )

                if value:
                    record[field] = value

            # Exact normalized fund name is the primary lookup.
            research[
                normalize_key(
                    fund_name
                )
            ] = record

        return (
            research,
            worksheet.title,
        )

    finally:
        workbook.close()


def read_master_category_sheet(
    workbook,
    sheet_index: int,
    category_type: str,
) -> tuple[list[str], str]:
    """
    Read the authoritative category list from a worksheet.

    The first useful single category column is selected.

    Duplicate values are removed while preserving worksheet order.
    """

    if sheet_index >= len(
        workbook.sheetnames
    ):
        raise RuntimeError(
            f"Research workbook does not contain "
            f"worksheet {sheet_index + 1} "
            f"for {category_type} master categories."
        )

    sheet_name = workbook.sheetnames[
        sheet_index
    ]

    worksheet = workbook[
        sheet_name
    ]

    preferred_headers = {
        "geography": [
            "Geography",
            "Geographic",
            "Geographic Category",
            "Region",
            "Category",
        ],
        "sector": [
            "Sector",
            "Sector Category",
            "Industry",
            "Category",
        ],
    }

    header_row = None
    category_column = None

    # Look for a clearly named header.
    for row in range(
        1,
        min(
            worksheet.max_row,
            20,
        )
        + 1,
    ):

        header_map = build_header_map(
            worksheet,
            row,
        )

        category_column = first_matching_column(
            header_map,
            preferred_headers[
                category_type
            ],
        )

        if category_column is not None:
            header_row = row
            break

    # Fallback:
    # choose the first non-empty column if the worksheet has no obvious
    # header. This still uses the supplied master worksheet itself.
    if category_column is None:

        for column in range(
            1,
            worksheet.max_column + 1,
        ):

            values = []

            for row in range(
                1,
                min(
                    worksheet.max_row,
                    20,
                )
                + 1,
            ):

                value = clean_text(
                    worksheet.cell(
                        row=row,
                        column=column,
                    ).value
                )

                if value:
                    values.append(
                        value
                    )

            if values:
                category_column = column
                header_row = 0
                break

    if category_column is None:
        raise RuntimeError(
            f"Could not identify a category column "
            f"in Research Funds.xlsx worksheet "
            f"'{sheet_name}'."
        )

    start_row = (
        header_row + 1
        if header_row
        else 1
    )

    categories = []
    seen = set()

    for row in range(
        start_row,
        worksheet.max_row + 1,
    ):

        value = clean_text(
            worksheet.cell(
                row=row,
                column=category_column,
            ).value
        )

        if not value:
            continue

        key = normalize_key(
            value
        )

        if key in seen:
            continue

        seen.add(key)
        categories.append(
            value
        )

    if not categories:
        raise RuntimeError(
            f"No {category_type} master categories "
            f"were found in worksheet '{sheet_name}'."
        )

    return (
        categories,
        sheet_name,
    )


def read_research_workbook() -> dict:
    """
    Read:

        Fund Research
        Geography master sheet
        Sector master sheet
    """

    if not RESEARCH_EXCEL_FILE.exists():
        raise FileNotFoundError(
            f"Research workbook not found: "
            f"{RESEARCH_EXCEL_FILE}"
        )

    (
        fund_research,
        fund_research_sheet,
    ) = read_fund_research()

    workbook = load_workbook(
        RESEARCH_EXCEL_FILE,
        read_only=True,
        data_only=True,
    )

    try:

        (
            geography_categories,
            geography_sheet,
        ) = read_master_category_sheet(
            workbook,
            RESEARCH_GEOGRAPHY_SHEET_INDEX,
            "geography",
        )

        (
            sector_categories,
            sector_sheet,
        ) = read_master_category_sheet(
            workbook,
            RESEARCH_SECTOR_SHEET_INDEX,
            "sector",
        )

    finally:
        workbook.close()

    return {
        "fundResearch": fund_research,
        "fundResearchWorksheet": fund_research_sheet,
        "geographyCategories": geography_categories,
        "geographyWorksheet": geography_sheet,
        "sectorCategories": sector_categories,
        "sectorWorksheet": sector_sheet,
    }


def get_research_for_fund(
    fund_research: dict[str, dict],
    prudential: dict,
    excel: dict,
) -> dict:
    """
    Match Research Funds.xlsx against the authoritative fund identity.

    Matching priority:

        1. fundName from Prudential source
        2. PruAccess name
        3. URL-derived identity is NOT used

    This prevents accidental fuzzy matches.
    """

    candidates = [
        clean_text(
            prudential.get(
                "fundName"
            )
        ),
        clean_text(
            excel.get(
                "pruAccessName"
            )
        ),
    ]

    for candidate in candidates:

        if not candidate:
            continue

        record = fund_research.get(
            normalize_key(
                candidate
            )
        )

        if record is not None:

            return {
                "geographic1": record.get(
                    "geographic1"
                ),
                "geographic2": record.get(
                    "geographic2"
                ),
                "sector1": record.get(
                    "sector1"
                ),
                "sector2": record.get(
                    "sector2"
                ),
            }

    return {
        "geographic1": None,
        "geographic2": None,
        "sector1": None,
        "sector2": None,
    }


# =============================================================================
# HOLDINGS
# =============================================================================

def clean_holdings(
    holdings: Any,
    label: str,
) -> list[dict]:
    """
    Validate holdings extracted by the frozen holdings pipelines.
    """

    if (
        not isinstance(
            holdings,
            list,
        )
        or not holdings
    ):
        raise ValueError(
            f"{label}: holdings list is empty or invalid."
        )

    if len(holdings) > MAX_HOLDINGS:
        raise ValueError(
            f"{label}: more than {MAX_HOLDINGS} holdings."
        )

    cleaned = []

    for position, item in enumerate(
        holdings,
        start=1,
    ):

        if not isinstance(
            item,
            dict,
        ):
            raise ValueError(
                f"{label}: holding {position} "
                "is not an object."
            )

        rank = item.get(
            "rank"
        )

        name = clean_text(
            item.get(
                "name"
            )
        )

        weight = item.get(
            "weightPercent"
        )

        weight_text = clean_text(
            item.get(
                "weightText"
            )
        )

        # Weight parsed into the name: "SK HYNIX INC 6.7%"
        embedded = re.match(
            r"^(.*?)\s+(\d+(?:\.\d+)?)\s*%$",
            name,
        )

        if embedded and embedded.group(1).strip():

            name = embedded.group(1).strip()

            weight = float(
                embedded.group(2)
            )

            weight_text = f"{weight:.1f}%"

        if rank != position:
            raise ValueError(
                f"{label}: rank {rank} "
                f"at position {position}."
            )

        if not name:
            raise ValueError(
                f"{label}: holding {position} "
                "has no name."
            )

        if (
            isinstance(
                weight,
                bool,
            )
            or not isinstance(
                weight,
                (
                    int,
                    float,
                ),
            )
            or not 0 <= weight <= 100
        ):
            raise ValueError(
                f"{label}: holding {position} "
                f"has invalid weight {weight!r}."
            )

        cleaned.append(
            {
                "rank": position,
                "name": name,
                "weightPercent": weight,
                "weightText": weight_text,
            }
        )

    return cleaned


def holdings_block(
    result: dict,
    stage: str,
    label: str,
) -> dict | None:
    """
    Convert one holdings pipeline result into the final published
    topHoldings block.

    Returns None for an unresolved/failed result.
    """

    status = result.get(
        "status"
    )

    common = {
        "source": stage,
        "parser": result.get(
            "holdingsParser"
        ),
        "factsheetUrl": result.get(
            "factsheetUrl"
        ),
        "factsheetDocumentDate": result.get(
            "factsheetDocumentDate"
        ),
        "factsheetDataAsAt": result.get(
            "factsheetDataAsAt"
        ),
    }

    if status == "success":

        holdings = clean_holdings(
            result.get(
                "topHoldings"
            ),
            label,
        )

        return {
            "status": "published",
            **common,
            "count": len(
                holdings
            ),
            "holdings": holdings,
        }

    if status == "no_holdings_section":

        return {
            "status": "no_holdings_section",
            **common,
            "count": 0,
            "holdings": [],
        }

    return None


def iter_stage_result_files(
    stage_dir: Path,
):
    """
    Support the existing recovery output structure.

    Preferred:

        <stage>/funds/*/top_holdings.json

    Also supports:

        <stage>/funds/*.json

    without changing the recovery pipelines.
    """

    funds_dir = (
        stage_dir / "funds"
    )

    if not funds_dir.exists():
        return

    preferred = sorted(
        funds_dir.glob(
            "*/top_holdings.json"
        )
    )

    if preferred:
        for path in preferred:
            yield path
        return

    for path in sorted(
        funds_dir.glob(
            "*.json"
        )
    ):
        if path.name in {
            "run_summary.json",
            "all_holdings.json",
        }:
            continue

        yield path


def resolve_holdings(
    excel_funds: dict[int, dict],
) -> dict[int, dict]:
    """
    Resolve holdings in strict order:

        baseline -> recovery1 -> recovery2 -> recovery3
    """

    resolved: dict[int, dict] = {}

    def accept(
        result: dict,
        stage: str,
        source_path: Path | None = None,
    ) -> None:

        if not isinstance(
            result,
            dict,
        ):
            return

        row_value = result.get(
            "excelRow"
        )

        if row_value is None:
            return

        try:
            row = int(
                row_value
            )
        except (
            TypeError,
            ValueError,
        ):
            return

        if row not in excel_funds:
            raise RuntimeError(
                f"{stage}: row {row} "
                "is not present in Funds Links.xlsm."
            )

        # First successful resolution wins.
        if row in resolved:
            return

        result_url = clean_text(
            result.get(
                "prudentialUrl"
            )
        )

        master_url = excel_funds[row][
            "prudentialUrl"
        ]

        if (
            result_url
            and result_url != master_url
        ):
            raise RuntimeError(
                f"{stage}: Prudential URL mismatch "
                f"for Excel row {row}."
            )

        block = holdings_block(
            result,
            stage,
            f"{stage} row {row}",
        )

        if block is None:
            return

        block["fundName"] = (
            clean_text(
                result.get(
                    "fundName"
                )
            )
            or None
        )

        if source_path is not None:
            block["sourceFile"] = str(
                source_path.relative_to(
                    ROOT
                )
            )

        resolved[row] = block

    # -----------------------------------------------------------------------
    # BASELINE
    # -----------------------------------------------------------------------

    if not BASELINE_FILE.exists():
        raise FileNotFoundError(
            f"Baseline holdings output not found: "
            f"{BASELINE_FILE}"
        )

    baseline = load_json(
        BASELINE_FILE
    )

    baseline_funds = baseline.get(
        "funds",
        [],
    )

    if not isinstance(
        baseline_funds,
        list,
    ):
        raise RuntimeError(
            f"Invalid baseline output: "
            f"{BASELINE_FILE}"
        )

    for result in baseline_funds:
        accept(
            result,
            "baseline",
            BASELINE_FILE,
        )

    # -----------------------------------------------------------------------
    # RECOVERY 1 / 2 / 3
    # -----------------------------------------------------------------------

    for stage, stage_dir in RECOVERY_STAGES:

        if not stage_dir.exists():
            print(
                f"INFO: {stage_dir} not found; "
                f"{stage} skipped."
            )
            continue

        found_any = False

        for path in iter_stage_result_files(
            stage_dir
        ):

            found_any = True

            try:
                result = load_json(
                    path
                )
            except Exception as error:
                print(
                    f"WARNING: Could not read "
                    f"{path}: {error}"
                )
                continue

            # A stage file may itself contain a fund object.
            if isinstance(
                result,
                dict,
            ):
                accept(
                    result,
                    stage,
                    path,
                )

        if not found_any:
            print(
                f"INFO: No fund result files found "
                f"for {stage}."
            )

    return resolved


# =============================================================================
# PRUDENTIAL SOURCE DATA
# =============================================================================

def parse_dividend_history(
    dividend_rate: str,
) -> list[dict]:
    """
    "20260930=1.25&20260630=1.25"
        -> [{"exDate": "2026-09-30", "rate": 1.25}, ...]

    Newest first. Unparseable parts are skipped.
    """

    history = []

    for part in dividend_rate.split(
        "&"
    ):

        date, _, rate = part.partition(
            "="
        )

        date = date.strip()
        rate = rate.strip()

        if not re.fullmatch(
            r"\d{8}",
            date,
        ):
            continue

        try:
            value = float(
                rate
            )
        except ValueError:
            continue

        history.append(
            {
                "exDate": (
                    f"{date[0:4]}-{date[4:6]}-{date[6:8]}"
                ),
                "rate": value,
            }
        )

    history.sort(
        key=lambda item: item["exDate"],
        reverse=True,
    )

    return history


def extract_dividend_fields(
    source: dict,
) -> dict:
    """
    Preserve Prudential's dividend values.

    Critical rule:

        non-empty dividendRate -> hasDividend = true
    """

    dividend_rate = clean_raw_text(
        source.get(
            "dividendRate"
        )
    )

    dividend_unit = clean_text(
        source.get(
            "dividendUnit"
        )
    )

    has_dividend = bool(
        dividend_rate
    )

    if dividend_unit:

        unit_source = (
            clean_text(
                source.get(
                    "dividendUnitSource"
                )
            )
            or "prudential"
        )

        unit_evidence = (
            clean_text(
                source.get(
                    "dividendUnitEvidence"
                )
            )
            or None
        )

    elif has_dividend:

        unit_source = "default"
        unit_evidence = None

    else:

        unit_source = None
        unit_evidence = None

    return {
        "hasDividend": has_dividend,
        "dividendRate": (
            dividend_rate
            if dividend_rate
            else None
        ),
        "dividendUnit": (
            dividend_unit
            or (
                DEFAULT_DIVIDEND_UNIT
                if has_dividend
                else None
            )
        ),
        "dividendUnitSource": unit_source,
        "dividendUnitEvidence": unit_evidence,
        "dividendHistory": (
            parse_dividend_history(
                dividend_rate
            )
            if has_dividend
            else []
        ),
    }


# ============================================================================
# PAYMENT MODE (Cash / SRS / CPF-OA / CPF-SA)
# ============================================================================
#
# Payment mode is extracted ONLY from explicit Prudential payment-mode
# fields in the raw ilpfunds.json record, plus explicitly named payment
# flags. We do not infer a payment mode merely because an unrelated field
# happens to contain the words "cash", "CPF", or "SRS".
#
# Supported explicit value examples:
#
#     "paymentMode": "Cash, SRS, CPF-OA"
#     "paymentModes": ["Cash", "SRS"]
#     "payMode": "CPFIS-OA"
#     "paymentMethod": "Cash"
#
# Supported explicit flag examples:
#
#     "isSrs": true
#     "srs": "Y"
#     "cpfisOa": true
#     "cpfOa": "Yes"
#     "cpfisSa": true
#     "cpfSa": "Y"
#     "isCash": true
#
# Result on each fund:
#
#     paymentModes
#         Ordered list containing only modes explicitly supported by the
#         Prudential source record.
#
#     paymentModeSource
#         "prudential-api:<field names>" when a mode was found.
#
#     paymentModeEvidence
#         The exact source fields/values used for the extraction, useful for
#         checking the result against the raw Prudential response.
#
# No AI, web research, or inference is performed here.

PAYMENT_MODE_ORDER = (
    "Cash",
    "SRS",
    "CPF-OA",
    "CPF-SA",
    "CPF",
)

# Explicit payment-mode field names. These are normalized by removing
# punctuation/case before comparison.
PAYMENT_VALUE_FIELDS = {
    "paymentmode",
    "paymentmodes",
    "paymode",
    "paymodes",
    "paymentmethod",
    "paymentmethods",
    "paymenttype",
    "paymenttypes",
    "allowedpaymentmode",
    "allowedpaymentmodes",
    "eligiblepaymentmode",
    "eligiblepaymentmodes",
    "availablepaymentmode",
    "availablepaymentmodes",
}

# Explicit flag field names. The name itself identifies the payment mode.
PAYMENT_FLAG_FIELDS = {
    # Cash
    "cash",
    "iscash",
    "cancash",
    "cashallowed",
    "casheligible",
    "iscashallowed",
    "iscasheligible",
    # SRS
    "srs",
    "issrs",
    "cansrs",
    "srsallowed",
    "srseligible",
    "issrsallowed",
    "issrseligible",
    # CPF-OA
    "cpfoa",
    "cpfisoa",
    "iscpfoa",
    "iscpfisoa",
    "cancpfoa",
    "cancpfisoa",
    "cpfoaallowed",
    "cpfisoaallowed",
    "cpfoaeligible",
    "cpfisoaeligible",
    # CPF-SA
    "cpfsa",
    "cpfissa",
    "iscpfsa",
    "iscpfissa",
    "cancpfsa",
    "cancpfissa",
    "cpfsaallowed",
    "cpfissaallowed",
    "cpfsaeligible",
    "cpfissaeligible",
}

PAYMENT_DIAGNOSTICS = {
    "raw_keys": set(),
    "funds_with_modes": 0,
    "field_usage": {},
    "mode_counts": {},
}

_FALSE_WORDS = {
    "",
    "n",
    "no",
    "false",
    "0",
    "none",
    "nil",
    "na",
    "n/a",
    "-",
    "not applicable",
    "not available",
    "null",
}

_TRUE_WORDS = {
    "y",
    "yes",
    "true",
    "1",
    "t",
}


def _payment_field_name(value: Any) -> str:
    """Normalize a raw API field name for exact payment-field matching."""

    return re.sub(
        r"[^a-z0-9]",
        "",
        str(value).lower(),
    )


def _is_truthy_payment_flag(value: Any) -> bool:
    """Return True only for values that explicitly mean enabled/allowed."""

    if isinstance(value, bool):
        return value

    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return value != 0

    text = clean_raw_text(value).strip().lower()

    if text in _FALSE_WORDS:
        return False

    if text in _TRUE_WORDS:
        return True

    # Do not treat arbitrary non-empty strings as a true flag. This avoids
    # turning unrelated text such as "Not supported" into a payment mode.
    return False


def _payment_modes_from_text(value: Any) -> set[str]:
    """
    Extract modes from the value of an explicit payment-mode field.

    This parser is deliberately conservative. It recognizes Cash, SRS,
    CPF-OA/CPFIS-OA, CPF-SA/CPFIS-SA, and generic CPF/CPFIS. It does not
    classify an arbitrary mention of these words elsewhere in the API.
    """

    if isinstance(value, dict):
        values = list(value.values())
        text = " ".join(clean_raw_text(item) for item in values)
    elif isinstance(value, (list, tuple, set)):
        text = " ".join(clean_raw_text(item) for item in value)
    else:
        text = clean_raw_text(value)

    if not text:
        return set()

    # Normalize separators but retain enough structure for token matching.
    normalized = re.sub(r"[\s_/]+", "-", text.upper())
    compact = re.sub(r"[^A-Z0-9]", "", text.upper())

    modes: set[str] = set()

    # Cash / SRS can be represented as standalone values or within a
    # comma/semicolon-delimited payment-mode string.
    if re.search(r"(?<![A-Z])CASH(?![A-Z])", normalized) or compact == "CASH":
        modes.add("Cash")

    if re.search(r"(?<![A-Z])SRS(?![A-Z])", normalized) or compact == "SRS":
        modes.add("SRS")

    # CPF Ordinary Account.
    if (
        re.search(r"CPF(?:IS)?-?(?:OA|ORDINARY(?:-ACCOUNT)?)", normalized)
        or compact in {"CPFOA", "CPFISOA", "CPFORDINARY", "CPFISORDINARY"}
    ):
        modes.add("CPF-OA")

    # CPF Special Account.
    if (
        re.search(r"CPF(?:IS)?-?(?:SA|SPECIAL(?:-ACCOUNT)?)", normalized)
        or compact in {"CPFSA", "CPFISSA", "CPFSPECIAL", "CPFISSPECIAL"}
    ):
        modes.add("CPF-SA")

    # Generic CPF/CPFIS is only returned when the explicit payment field
    # itself says CPF but does not identify OA or SA.
    if not ({"CPF-OA", "CPF-SA"} & modes):
        if re.search(r"(?<![A-Z])CPFIS?(?![A-Z])", normalized) or compact in {
            "CPF",
            "CPFIS",
        }:
            modes.add("CPF")

    return modes


def _payment_mode_from_flag_name(field_name: str) -> str | None:
    """Map an explicitly named payment flag to its canonical mode."""

    name = _payment_field_name(field_name)

    if name in {
        "cash",
        "iscash",
        "cancash",
        "cashallowed",
        "casheligible",
        "iscashallowed",
        "iscasheligible",
    }:
        return "Cash"

    if name in {
        "srs",
        "issrs",
        "cansrs",
        "srsallowed",
        "srseligible",
        "issrsallowed",
        "issrseligible",
    }:
        return "SRS"

    if name in {
        "cpfoa",
        "cpfisoa",
        "iscpfoa",
        "iscpfisoa",
        "cancpfoa",
        "cancpfisoa",
        "cpfoaallowed",
        "cpfisoaallowed",
        "cpfoaeligible",
        "cpfisoaeligible",
    }:
        return "CPF-OA"

    if name in {
        "cpfsa",
        "cpfissa",
        "iscpfsa",
        "iscpfissa",
        "cancpfsa",
        "cancpfissa",
        "cpfsaallowed",
        "cpfissaallowed",
        "cpfsaeligible",
        "cpfissaeligible",
    }:
        return "CPF-SA"

    return None


def _record_payment_diagnostic(field_name: str, modes: set[str]) -> None:
    """Record which raw field produced which payment modes."""

    key = str(field_name)

    usage = PAYMENT_DIAGNOSTICS["field_usage"]
    usage[key] = usage.get(key, 0) + 1

    for mode in modes:
        counts = PAYMENT_DIAGNOSTICS["mode_counts"]
        counts[mode] = counts.get(mode, 0) + 1


def extract_payment_modes(raw: Any) -> dict:
    """
    Extract payment modes from explicit Prudential API fields only.

    Priority:
        1. Explicit payment-mode value fields.
        2. Explicitly named payment flags.

    A generic raw field containing "CPF", "SRS", or "Cash" is ignored
    unless its field name itself is a recognized payment field/flag.
    """

    empty = {
        "paymentModes": [],
        "paymentModeSource": None,
        "paymentModeEvidence": None,
    }

    if not isinstance(raw, dict):
        return empty

    PAYMENT_DIAGNOSTICS["raw_keys"].update(str(key) for key in raw.keys())

    modes: set[str] = set()
    used: list[tuple[str, Any]] = []

    # -----------------------------------------------------------------------
    # 1. Explicit payment-mode value fields
    # -----------------------------------------------------------------------

    for key, value in raw.items():
        field_name = _payment_field_name(key)

        if field_name not in PAYMENT_VALUE_FIELDS:
            continue

        if value is None:
            continue

        found = _payment_modes_from_text(value)

        # Some APIs return an object such as:
        # {"cash": true, "srs": true, "cpfOa": true}
        # under paymentMode/paymentModes. Handle that structure explicitly.
        if isinstance(value, dict):
            for nested_key, nested_value in value.items():
                nested_mode = _payment_mode_from_flag_name(str(nested_key))
                if nested_mode and _is_truthy_payment_flag(nested_value):
                    found.add(nested_mode)

        if found:
            modes |= found
            used.append((str(key), value))
            _record_payment_diagnostic(str(key), found)

    # -----------------------------------------------------------------------
    # 2. Explicit payment flags
    # -----------------------------------------------------------------------

    for key, value in raw.items():
        field_name = _payment_field_name(key)

        if field_name not in PAYMENT_FLAG_FIELDS:
            continue

        if not _is_truthy_payment_flag(value):
            continue

        mode = _payment_mode_from_flag_name(str(key))

        if mode:
            found = {mode}
            modes |= found
            used.append((str(key), value))
            _record_payment_diagnostic(str(key), found)

    if not modes:
        return empty

    PAYMENT_DIAGNOSTICS["funds_with_modes"] += 1

    return {
        "paymentModes": [
            mode
            for mode in PAYMENT_MODE_ORDER
            if mode in modes
        ],
        "paymentModeSource": (
            "prudential-api:"
            + ",".join(key for key, _ in used)
        ),
        "paymentModeEvidence": "; ".join(
            f"{key}={clean_raw_text(value)[:120]}"
            for key, value in used
        ),
    }


def print_payment_mode_diagnostics(fund_count: int) -> None:
    """Print a useful payment-mode extraction summary for the build log."""

    print()
    print("Payment mode:")
    print(
        f"  funds with payment mode: "
        f"{PAYMENT_DIAGNOSTICS['funds_with_modes']} of {fund_count}"
    )

    mode_counts = PAYMENT_DIAGNOSTICS["mode_counts"]

    if mode_counts:
        print("  modes:")
        for mode in PAYMENT_MODE_ORDER:
            if mode in mode_counts:
                print(f"    {mode}: {mode_counts[mode]}")

    field_usage = PAYMENT_DIAGNOSTICS["field_usage"]

    if field_usage:
        print("  source fields:")
        for field_name, count in sorted(
            field_usage.items(),
            key=lambda item: (-item[1], item[0].lower()),
        ):
            print(f"    {field_name}: {count}")
    else:
        keys = sorted(PAYMENT_DIAGNOSTICS["raw_keys"])
        print(
            "  No explicit Prudential payment-mode field was found. "
            "Raw API fields available:"
        )

        for start in range(0, len(keys), 6):
            print("    " + ", ".join(keys[start:start + 6]))


def normalize_prudential_fund(
    source: dict,
) -> dict:
    """
    Preserve the extracted Prudential fund metadata while ensuring
    dividend fields are correct.
    """

    if not isinstance(
        source,
        dict,
    ):
        raise ValueError(
            "Prudential fund record is not an object."
        )

    fund = {
        key: value
        for key, value in source.items()
        if key not in {
            "raw",
            "fundIdentifier",
            "fundCode",
            "fundName",
        }
    }

    dividend_fields = extract_dividend_fields(
        source
    )

    fund.update(
        dividend_fields
    )

    fund.update(
        extract_payment_modes(
            source.get("raw")
        )
    )

    if "fundObjective" in fund:

        fund["fundObjective"] = clean_html_text(
            fund.get(
                "fundObjective"
            )
        )

    return fund


def load_prudential_source_records() -> dict[int, dict]:
    """
    Load prudential_fund.json records from the existing PruAccess
    extraction output.

    Preferred location:

        output_pruaccess/funds/<row>_<id>/prudential_fund.json

    A production directory is also checked as a compatibility
    fallback.
    """

    result: dict[int, dict] = {}

    # -----------------------------------------------------------------------
    # Primary source
    # -----------------------------------------------------------------------

    if PRUACCESS_FUNDS_DIR.exists():

        for directory in sorted(
            PRUACCESS_FUNDS_DIR.iterdir()
        ):

            if not directory.is_dir():
                continue

            match = re.match(
                r"^(\d+)_",
                directory.name,
            )

            if not match:
                continue

            row = int(
                match.group(1)
            )

            path = (
                directory
                / "prudential_fund.json"
            )

            if not path.exists():
                continue

            try:
                data = load_json(
                    path
                )
            except Exception as error:
                print(
                    f"WARNING: Could not read "
                    f"{path}: {error}"
                )
                continue

            if isinstance(
                data,
                dict,
            ):
                result[row] = data

    # -----------------------------------------------------------------------
    # Production fallback
    # -----------------------------------------------------------------------

    if PRUACCESS_PRODUCTION_DIR.exists():

        for path in sorted(
            PRUACCESS_PRODUCTION_DIR.rglob(
                "prudential_fund.json"
            )
        ):

            match = re.search(
                r"(?:^|/)(\d+)_",
                str(
                    path.parent
                ).replace(
                    "\\",
                    "/",
                ),
            )

            if not match:
                continue

            row = int(
                match.group(1)
            )

            if row in result:
                continue

            try:
                data = load_json(
                    path
                )
            except Exception:
                continue

            if isinstance(
                data,
                dict,
            ):
                result[row] = data

    return result


# =============================================================================
# BID HISTORY
# =============================================================================

def extract_observations(
    history: dict,
) -> list:
    """
    Accept the existing PruAccess observation layout.

    The extraction pipeline historically uses:

        observations: [
            {
                "date": "...",
                "bidPrice": ...
            }
        ]

    Compatibility is also provided for "bid".
    """

    observations = history.get(
        "observations"
    )

    if not isinstance(
        observations,
        list,
    ):
        return []

    normalized = []

    for item in observations:

        if not isinstance(
            item,
            dict,
        ):
            raise ValueError(
                "BID observation is not an object."
            )

        date = clean_text(
            item.get(
                "date"
            )
        )

        if not date:
            raise ValueError(
                "BID observation has no date."
            )

        if "bidPrice" in item:
            bid = item.get(
                "bidPrice"
            )
        else:
            bid = item.get(
                "bid"
            )

        if (
            isinstance(
                bid,
                bool,
            )
            or not isinstance(
                bid,
                (
                    int,
                    float,
                ),
            )
        ):
            raise ValueError(
                f"Invalid BID value on {date}: "
                f"{bid!r}"
            )

        normalized.append(
            {
                "date": date,
                "bidPrice": bid,
            }
        )

    return normalized


def validate_bid_history(
    history: dict,
    label: str,
) -> list[dict]:
    """
    Validate chronological BID observations.

    No raw-date carry-forward is introduced.
    """

    if not isinstance(
        history,
        dict,
    ):
        raise ValueError(
            f"{label}: BID history is not an object."
        )

    observations = extract_observations(
        history
    )

    if not observations:
        raise ValueError(
            f"{label}: no BID observations."
        )

    expected_count = history.get(
        "observationCount"
    )

    if (
        expected_count is not None
        and expected_count
        != len(observations)
    ):
        raise ValueError(
            f"{label}: observationCount mismatch. "
            f"Expected {expected_count}, "
            f"got {len(observations)}."
        )

    seen = set()
    previous_date = None

    for item in observations:

        date = item["date"]

        if date in seen:
            raise ValueError(
                f"{label}: duplicate BID date "
                f"{date}."
            )

        if (
            previous_date is not None
            and date < previous_date
        ):
            raise ValueError(
                f"{label}: BID observations "
                "are not chronological."
            )

        seen.add(
            date
        )

        previous_date = date

    return observations


def load_bid_history() -> dict[int, dict]:
    """
    Load bid_history.json from:

        output_pruaccess/funds/<row>_<id>/

    """

    result: dict[int, dict] = {}

    if not PRUACCESS_FUNDS_DIR.exists():
        print(
            f"WARNING: PruAccess funds directory "
            f"not found: {PRUACCESS_FUNDS_DIR}"
        )
        return result

    for directory in sorted(
        PRUACCESS_FUNDS_DIR.iterdir()
    ):

        if not directory.is_dir():
            continue

        match = re.match(
            r"^(\d+)_",
            directory.name,
        )

        if not match:
            continue

        row = int(
            match.group(1)
        )

        path = (
            directory
            / "bid_history.json"
        )

        if not path.exists():
            continue

        try:
            history = load_json(
                path
            )
        except Exception as error:
            print(
                f"WARNING: Could not read "
                f"{path}: {error}"
            )
            continue

        if not isinstance(
            history,
            dict,
        ):
            continue

        result[row] = history

    return result


# =============================================================================
# IDENTITY
# =============================================================================

def build_identity(
    excel: dict,
    prudential: dict,
    holdings: dict | None,
) -> dict:

    fund_name = (
        clean_text(
            prudential.get(
                "fundName"
            )
        )
        or clean_text(
            (holdings or {}).get(
                "fundName"
            )
        )
        or None
    )

    fund_identifier = (
        clean_text(
            prudential.get(
                "fundIdentifier"
            )
        )
        or None
    )

    fund_code = (
        clean_text(
            prudential.get(
                "fundCode"
            )
        )
        or None
    )

    return {
        "excelRow": excel[
            "excelRow"
        ],
        "fundIdentifier": fund_identifier,
        "fundCode": fund_code,
        "fundName": fund_name,
        "pruAccessName": (
            excel.get(
                "pruAccessName"
            )
            or None
        ),
        "prudentialUrl": excel[
            "prudentialUrl"
        ],
    }


# =============================================================================
# MAIN BUILD
# =============================================================================

# =============================================================================
# FAILED FUNDS: SKIP, KEEP LAST GOOD DATA, REPORT
# =============================================================================
#
# A fund that fails any part of this run (Prudential metadata, holdings or
# BID history) does not stop the build. Instead:
#
#   - funds that succeeded are published with this run's data;
#   - a failed fund keeps its record from the last successful run
#     (data/funds.json + data/bid_history.json before this run), marked
#     with updateStatus.status = "stale";
#   - a failed fund that has never succeeded is left out;
#   - when only the holdings failed, details and prices are still updated,
#     with holdings from the last good run (or published without holdings);
#   - every failure is listed in the log and in the GitHub Actions run
#     summary.
#
# ALLOW_UNRESOLVED=1 keeps the old behaviour of publishing the partial
# record of a failed fund instead.

ROW_MESSAGE_PATTERN = re.compile(r"^Row (\d+): (.*)$")


def friendly_utc(value: str | None) -> str:
    """'2026-10-06T17:07:52+00:00' -> '06 Oct 2026' (or 'an earlier run')."""

    try:
        return datetime.fromisoformat(str(value)).strftime("%d %b %Y")
    except (TypeError, ValueError):
        return "an earlier run"


def failures_by_row(messages: list[str]) -> dict[int, list[str]]:
    """Group 'Row N: reason' messages by Excel row."""

    grouped: dict[int, list[str]] = {}

    for message in messages:
        match = ROW_MESSAGE_PATTERN.match(message)

        if match:
            grouped.setdefault(int(match.group(1)), []).append(match.group(2))

    return grouped


def load_previous_output(path: Path) -> dict:
    """The last published output, or an empty structure if unavailable."""

    try:
        data = load_json(path)
    except Exception as error:
        print(f"Previous output not available ({path.name}): {error}")
        return {}

    return data if isinstance(data, dict) else {}


def previous_index(previous: dict) -> dict[str, dict]:
    """Index last-run records by Prudential URL and PruAccess name."""

    index: dict[str, dict] = {}

    for record in previous.get("funds") or []:
        if not isinstance(record, dict):
            continue

        for key in (
            clean_text(record.get("prudentialUrl")),
            "name:" + normalize_key(record.get("pruAccessName")),
        ):
            if key and key != "name:":
                index.setdefault(key, record)

    return index


def find_previous(index: dict[str, dict], excel: dict) -> dict | None:
    return (
        index.get(clean_text(excel.get("prudentialUrl")))
        or index.get("name:" + normalize_key(excel.get("pruAccessName")))
    )


def apply_failure_policy(
    excel_funds: dict[int, dict],
    fund_records: list[dict],
    bid_records: list[dict],
    unresolved_messages: list[str],
    generated_at: str,
) -> tuple[list[dict], list[dict], list[dict]]:
    """
    Keep successful funds, replace failed funds with their last good
    record (or leave them out), and return the failure report rows.
    """

    failed = failures_by_row(unresolved_messages)

    previous_funds = load_previous_output(FUNDS_OUT)
    previous_bids = load_previous_output(BID_OUT)
    previous_fund_index = previous_index(previous_funds)
    previous_bid_index = previous_index(previous_bids)
    previous_generated = clean_text(previous_funds.get("generatedAtUtc")) or None

    kept_funds: list[dict] = []
    kept_bids: list[dict] = []
    report: list[dict] = []

    bids_by_row = {record.get("excelRow"): record for record in bid_records}

    for record in fund_records:
        row = record.get("excelRow")
        bid_record = bids_by_row.get(row)
        reasons = failed.get(row)

        if not reasons:
            status = {"status": "updated", "lastSuccessfulUpdateUtc": generated_at}

            kept_funds.append({**record, "updateStatus": status})

            if bid_record is not None:
                kept_bids.append({**bid_record, "updateStatus": status})

            continue

        excel = excel_funds.get(row, {})
        name = record.get("fundName") or excel.get("pruAccessName") or excel.get("prudentialUrl")
        old_fund = find_previous(previous_fund_index, excel)
        old_bid = find_previous(previous_bid_index, excel)

        holdings_only = all(reason.startswith("holdings unresolved") for reason in reasons)

        if holdings_only and bid_record is not None:
            # Only the holdings failed: details and prices are fresh, so
            # publish them. Holdings come from the last good run when there
            # is one; otherwise the fund is published without holdings
            # (shown as "not available" on the site) and tried again next run.
            old_holdings = (old_fund or {}).get("topHoldings") or {}
            reuse = old_holdings.get("status") in ("published", "no_holdings_section")
            status = {
                "status": "updated",
                "lastSuccessfulUpdateUtc": generated_at,
                "reasons": reasons,
            }

            if reuse:
                new_record = {**record, "topHoldings": old_holdings}
                status["holdingsFromEarlierRun"] = True
                action = "Updated; holdings kept from the last good run"
                outcome = "updated_holdings_kept"
            else:
                new_record = record
                status["holdingsMissing"] = True
                action = "Published without holdings (tried again next run)"
                outcome = "updated_no_holdings"

            kept_funds.append({**new_record, "updateStatus": status})
            kept_bids.append({**bid_record, "updateStatus": status})

        elif old_fund is not None and old_bid is not None:
            old_status = old_fund.get("updateStatus") or {}
            last_good = old_status.get("lastSuccessfulUpdateUtc") or previous_generated
            status = {
                "status": "stale",
                "lastSuccessfulUpdateUtc": last_good,
                "failedUpdateUtc": generated_at,
                "reasons": reasons,
            }

            kept_funds.append({**old_fund, "excelRow": row, "updateStatus": status})
            kept_bids.append({**old_bid, "excelRow": row, "updateStatus": status})

            action = f"Kept last good data from {friendly_utc(last_good)}"
            outcome = "kept_last_good"
        else:
            action = "Left out (no earlier successful data)"
            outcome = "left_out"

        report.append({
            "excelRow": row,
            "fund": name,
            "reasons": reasons,
            "action": action,
            "outcome": outcome,
        })

    return kept_funds, kept_bids, report


def write_failure_report(
    report: list[dict],
    published: int,
    total: int,
    closed_published: int = 0,
    closed_report: list[dict] | tuple = (),
) -> None:
    """Print the failure report and add it to the GitHub run summary."""

    updated = published - sum(1 for item in report if item["outcome"] == "kept_last_good")
    report = list(report) + list(closed_report)

    print()
    print("=" * 72)
    print("FAILED FUNDS (skipped, build continues)")
    print("=" * 72)

    for item in report:
        print(f" - Row {item['excelRow']}: {item['fund']}")

        for reason in item["reasons"]:
            print(f"     reason: {reason}")

        print(f"     action: {item['action']}")

    lines = [
        "## Funds data build",
        "",
        f"- Funds in Funds Links.xlsm: **{total}**",
        f"- Updated this run: **{updated}**",
        f"- With a failure: **{len(report)}**",
        f"- Published in total: **{published}**",
    ]

    if closed_published or closed_report:
        lines.append(f"- Closed funds published (past prices only): **{closed_published}**")

    lines.append("")

    if report:
        lines += [
            "### Failed funds",
            "",
            "| Row | Fund | Reason | What happened |",
            "| --- | --- | --- | --- |",
        ]

        for item in report:
            reason = "<br>".join(r.replace("|", "\\|") for r in item["reasons"])
            fund = str(item["fund"]).replace("|", "\\|")
            lines.append(f"| {item['excelRow']} | {fund} | {reason} | {item['action']} |")
    else:
        lines.append("All funds updated successfully.")

    summary_path = os.environ.get("GITHUB_STEP_SUMMARY")

    if summary_path:
        try:
            with open(summary_path, "a", encoding="utf-8") as handle:
                handle.write("\n".join(lines) + "\n")
        except OSError as error:
            print(f"Could not write the GitHub run summary: {error}")


# =============================================================================
# CLOSED FUNDS
# =============================================================================
#
# Closed Funds.xlsx lists funds Prudential has closed. Their past prices
# are fetched from PruAccess by scripts/test_pruaccess.py into
# output_pruaccess/closed/<code>/bid_history.json. Each one is published
# with closed = true, its last price, and no returns or holdings, so the
# site can show its history marked "Closed" and leave it out of rankings.
#
# If a fetch fails (or returns fewer prices than last time) the last
# published record is kept; a closed fund that has never been fetched is
# left out. Either way it is listed in the run summary.

CLOSED_FUNDS_FILE = ROOT / "Closed Funds.xlsx"

CLOSED_FUNDS_SHEET = "ClosedFunds"

CLOSED_OUTPUT_DIR = ROOT / "output_pruaccess" / "closed"


def read_closed_funds() -> list[dict]:
    """Rows of Closed Funds.xlsx (empty list if the file is absent)."""

    if not CLOSED_FUNDS_FILE.exists():
        return []

    workbook = load_workbook(
        CLOSED_FUNDS_FILE,
        read_only=True,
        data_only=True,
    )

    try:
        sheet = (
            workbook[CLOSED_FUNDS_SHEET]
            if CLOSED_FUNDS_SHEET in workbook.sheetnames
            else workbook[workbook.sheetnames[0]]
        )

        funds = []

        for row_number, row in enumerate(
            sheet.iter_rows(min_row=2, max_col=8, values_only=True),
            start=2,
        ):
            values = list(row) + [None] * (8 - len(row))
            name = clean_text(values[0])

            if not name:
                continue

            closed_date = values[6]

            if isinstance(closed_date, datetime):
                closed_date = closed_date.strftime("%d-%b-%Y")

            funds.append({
                "closedRow": row_number,
                "pruAccessName": name,
                "fundCode": clean_text(values[1]),
                "currency": clean_text(values[2]),
                "assetClass": clean_text(values[3]),
                "riskClassification": clean_text(values[4]),
                "paysDividend": clean_text(values[5]).lower() in ("yes", "y", "true", "1"),
                "closedDate": clean_text(closed_date),
                "notes": clean_text(values[7]),
            })
    finally:
        workbook.close()

    return funds


def closed_identifier(closed: dict) -> str:
    code = clean_text(closed.get("fundCode")) or normalize_key(closed.get("pruAccessName"))
    return f"CLOSED-{code.upper()}"


def load_closed_outputs() -> dict[str, dict]:
    """
    This run's closed-fund fetch results, keyed by normalized fund code
    and by "name:" + normalized PruAccess name.
    Each value: {"history": ..., "failed": error text or None}.
    """

    results: dict[str, dict] = {}

    if not CLOSED_OUTPUT_DIR.exists():
        return results

    for directory in sorted(CLOSED_OUTPUT_DIR.iterdir()):
        if not directory.is_dir():
            continue

        entry: dict = {"history": None, "failed": None, "meta": {}}

        failure_path = directory / "failure.json"
        history_path = directory / "bid_history.json"
        meta_path = directory / "closed_fund.json"

        try:
            if meta_path.exists():
                entry["meta"] = load_json(meta_path) or {}
        except Exception:
            pass

        if failure_path.exists():
            try:
                failure = load_json(failure_path) or {}
                entry["meta"] = {**failure, **entry["meta"]}
                entry["failed"] = clean_text(failure.get("error")) or "PruAccess fetch failed"
            except Exception as error:
                entry["failed"] = f"failure.json unreadable: {error}"
        elif history_path.exists():
            try:
                entry["history"] = load_json(history_path)
            except Exception as error:
                entry["failed"] = f"bid_history.json unreadable: {error}"
        else:
            continue

        meta = entry["meta"]
        history = entry["history"] or {}

        code = normalize_key(meta.get("fundCode") or history.get("fundCode"))
        name = normalize_key(
            meta.get("pruAccessName")
            or history.get("fundName")
            or history.get("pruAccessFundName")
        )

        if code:
            results.setdefault(code, entry)
        if name:
            results.setdefault("name:" + name, entry)

    return results


def iso_to_display(iso: str, pattern: str) -> str:
    try:
        return datetime.strptime(iso[:10], "%Y-%m-%d").strftime(pattern)
    except (TypeError, ValueError):
        return ""


def build_closed_records(
    closed: dict,
    history: dict,
    observations: list[dict],
    research: dict,
    generated_at: str,
) -> tuple[dict, dict]:
    """Fund and BID records for one closed fund."""

    identifier = closed_identifier(closed)
    first = observations[0]
    last = observations[-1]
    currency = (
        clean_text(closed.get("currency"))
        or clean_text(history.get("currency"))
        or "SGD"
    )
    last_price_date = iso_to_display(last["date"], "%d-%b-%Y")
    closed_date = clean_text(closed.get("closedDate")) or last_price_date

    identity = {
        "excelRow": None,
        "closedRow": closed.get("closedRow"),
        "fundIdentifier": identifier,
        "fundCode": clean_text(closed.get("fundCode")) or None,
        "fundName": closed["pruAccessName"],
        "pruAccessName": closed["pruAccessName"],
        "prudentialUrl": None,
    }

    status = {
        "status": "updated",
        "lastSuccessfulUpdateUtc": generated_at,
    }

    fund_block = {
        "fundCurrency": currency,
        "unitCurrency": currency,
        "assetClass": clean_text(closed.get("assetClass")) or "Closed fund",
        "assetSubClass": None,
        "riskClassification": clean_text(closed.get("riskClassification")) or None,
        "bidPrice": f"${last['bidPrice']:.4f}",
        "offerPrice": None,
        "valuationDate": iso_to_display(last["date"], "%d/%m/%Y"),
        "inceptionDate": iso_to_display(first["date"], "%d-%b-%Y"),
        "closedDate": closed_date,
        "lastPriceDate": last_price_date,
        "firstPriceDate": iso_to_display(first["date"], "%d-%b-%Y"),
        "cumulativeYtd": "-",
        "cumulative1m": "-",
        "cumulative3m": "-",
        "cumulative6m": "-",
        "cumulative1y": "-",
        "cumulative3y": "-",
        "cumulative5y": "-",
        "annualised3y": "-",
        "annualised5y": "-",
        "annualised10y": "-",
        "annualisedSinceLaunch": "-",
        "factsheetUrl": None,
        "prospectusUrl": None,
        "productHighlightSheetUrl": None,
        "annualReportUrl": None,
        "fundObjective": clean_text(closed.get("notes")) or None,
        "investmentManager": None,
        "hasDividend": bool(closed.get("paysDividend")),
        "dividendRate": None,
        "dividendUnit": None,
        "dividendUnitSource": None,
        "dividendUnitEvidence": None,
        "dividendHistory": [],
        "paymentModes": [],
        "paymentModeSource": None,
        "paymentModeEvidence": None,
    }

    fund_record = {
        **identity,
        "closed": True,
        "fund": fund_block,
        "research": research,
        "topHoldings": {
            "status": "closed",
            "holdings": [],
        },
        "updateStatus": status,
    }

    bid_record = {
        **identity,
        "closed": True,
        "bidHistory": {
            "status": "success",
            "priceType": "BID",
            "pruAccessFundId": history.get("fundId"),
            "currency": currency,
            "startDate": history.get("startDate"),
            "endDate": history.get("endDate"),
            "observationCount": len(observations),
            "observations": observations,
        },
        "updateStatus": status,
    }

    return fund_record, bid_record


def build_closed_funds(
    fund_research: dict,
    generated_at: str,
) -> tuple[list[dict], list[dict], list[dict]]:
    """
    Closed fund records for funds.json and bid_history.json, plus report
    rows for any closed fund that could not be refreshed.
    """

    try:
        closed_funds = read_closed_funds()
    except Exception as error:
        print(f"WARNING: could not read {CLOSED_FUNDS_FILE.name}: {error}")
        return [], [], [{
            "excelRow": "Closed",
            "fund": CLOSED_FUNDS_FILE.name,
            "reasons": [f"could not read the file: {error}"],
            "action": "No closed funds published",
            "outcome": "left_out",
        }]

    if not closed_funds:
        return [], [], []

    outputs = load_closed_outputs()

    previous_funds = load_previous_output(FUNDS_OUT)
    previous_bids = load_previous_output(BID_OUT)
    previous_generated = clean_text(previous_funds.get("generatedAtUtc")) or None

    def previous_closed(previous: dict, identifier: str) -> dict | None:
        for record in previous.get("funds") or []:
            if (
                isinstance(record, dict)
                and record.get("closed")
                and record.get("fundIdentifier") == identifier
            ):
                return record
        return None

    fund_records: list[dict] = []
    bid_records: list[dict] = []
    report: list[dict] = []

    print()
    print("=" * 72)
    print(f"CLOSED FUNDS ({len(closed_funds)})")
    print("=" * 72)

    for closed in closed_funds:
        identifier = closed_identifier(closed)
        name = closed["pruAccessName"]
        label = f"Closed fund {name}"

        output = (
            outputs.get(normalize_key(closed.get("fundCode")))
            if closed.get("fundCode")
            else None
        ) or outputs.get("name:" + normalize_key(name))

        old_fund = previous_closed(previous_funds, identifier)
        old_bid = previous_closed(previous_bids, identifier)
        old_count = (
            ((old_bid or {}).get("bidHistory") or {}).get("observationCount") or 0
        )

        reason = None

        if output is None:
            reason = "no PruAccess result for this fund in this run"
        elif output["failed"]:
            reason = output["failed"]
        else:
            try:
                observations = validate_bid_history(output["history"], label)
            except Exception as error:
                reason = str(error)
            else:
                if len(observations) < old_count:
                    reason = (
                        f"PruAccess returned {len(observations)} prices, "
                        f"fewer than the {old_count} published before"
                    )

        if reason is None:
            research = get_research_for_fund(fund_research, {}, {"pruAccessName": name})
            fund_record, bid_record = build_closed_records(
                closed,
                output["history"],
                observations,
                research,
                generated_at,
            )
            fund_records.append(fund_record)
            bid_records.append(bid_record)
            print(
                f" - {name}: {len(observations)} prices, "
                f"{observations[0]['date']} to {observations[-1]['date']}"
            )
            continue

        if old_fund is not None and old_bid is not None:
            last_good = (
                (old_fund.get("updateStatus") or {}).get("lastSuccessfulUpdateUtc")
                or previous_generated
            )
            # A closed fund's prices no longer change, so the last good
            # record is kept as it is; the failure is only reported.
            fund_records.append({**old_fund, "closedRow": closed.get("closedRow")})
            bid_records.append({**old_bid, "closedRow": closed.get("closedRow")})
            action = f"Kept last good data from {friendly_utc(last_good)}"
            outcome = "kept_last_good"
            print(f" - {name}: FAILED ({reason}); {action}")
        else:
            action = "Left out (no earlier successful data)"
            outcome = "left_out"
            print(f" - {name}: FAILED ({reason}); {action}")

        report.append({
            "excelRow": "Closed",
            "fund": name,
            "reasons": [reason],
            "action": action,
            "outcome": outcome,
        })

    return fund_records, bid_records, report


def main() -> int:

    print(
        "=" * 72
    )
    print(
        "VGRAT FMS - BUILD FUNDS DATA"
    )
    print(
        "=" * 72
    )

    print(
        f"Repository root: {ROOT}"
    )

    print(
        f"Funds Links.xlsm: {EXCEL_FILE}"
    )

    print(
        f"Research Funds.xlsx: {RESEARCH_EXCEL_FILE}"
    )

    print(
        f"ALLOW_UNRESOLVED: {ALLOW_UNRESOLVED}"
    )

    # -----------------------------------------------------------------------
    # MASTER UNIVERSE
    # -----------------------------------------------------------------------

    print()
    print(
        "Reading master fund universe..."
    )

    excel_funds = read_excel_funds()

    # -----------------------------------------------------------------------
    # RESEARCH INPUT
    # -----------------------------------------------------------------------

    print()
    print(
        "Reading Research Funds.xlsx..."
    )

    research_data = (
        read_research_workbook()
    )

    fund_research = (
        research_data[
            "fundResearch"
        ]
    )

    geography_categories = (
        research_data[
            "geographyCategories"
        ]
    )

    sector_categories = (
        research_data[
            "sectorCategories"
        ]
    )

    print(
        f"Fund Research records: "
        f"{len(fund_research)}"
    )

    print(
        f"Geography master categories: "
        f"{len(geography_categories)}"
    )

    print(
        f"Sector master categories: "
        f"{len(sector_categories)}"
    )

    # -----------------------------------------------------------------------
    # HOLDINGS
    # -----------------------------------------------------------------------

    print()
    print(
        "Resolving holdings..."
    )

    holdings_by_row = (
        resolve_holdings(
            excel_funds
        )
    )

    # -----------------------------------------------------------------------
    # PRUDENTIAL SOURCE
    # -----------------------------------------------------------------------

    print()
    print(
        "Loading Prudential fund metadata..."
    )

    prudential_by_row = (
        load_prudential_source_records()
    )

    print(
        f"Prudential source records: "
        f"{len(prudential_by_row)}"
    )

    # -----------------------------------------------------------------------
    # BID
    # -----------------------------------------------------------------------

    print()
    print(
        "Loading PruAccess BID history..."
    )

    bid_by_row = (
        load_bid_history()
    )

    print(
        f"BID history records: "
        f"{len(bid_by_row)}"
    )

    # -----------------------------------------------------------------------
    # BUILD RECORDS
    # -----------------------------------------------------------------------

    fund_records = []

    bid_records = []

    generated_at = (
        utc_now_iso()
    )

    # Summary counters
    holdings_published = 0
    holdings_no_section = 0
    holdings_unresolved = 0

    research_matched = 0
    research_missing = 0

    prudential_matched = 0
    prudential_missing = 0

    dividend_funds = 0

    payment_mode_funds = 0

    bid_success = 0
    bid_missing = 0
    total_bid_observations = 0

    holdings_by_source: dict[
        str,
        int,
    ] = {}

    unresolved_messages = []

    # -----------------------------------------------------------------------
    # Every populated master row becomes exactly one record in each output.
    # -----------------------------------------------------------------------

    for row in sorted(
        excel_funds
    ):

        excel = excel_funds[
            row
        ]

        prudential = (
            prudential_by_row.get(
                row,
                {},
            )
        )

        holdings = (
            holdings_by_row.get(
                row
            )
        )

        bid_history_source = (
            bid_by_row.get(
                row
            )
        )

        identity = build_identity(
            excel,
            prudential,
            holdings,
        )

        # ================================================================
        # FUND DATA
        # ================================================================

        if prudential:

            prudential_matched += 1

            fund_info = (
                normalize_prudential_fund(
                    prudential
                )
            )

            dividend_rate = (
                clean_raw_text(
                    fund_info.get(
                        "dividendRate"
                    )
                )
            )

            if dividend_rate:
                dividend_funds += 1

            if fund_info.get("paymentModes"):
                payment_mode_funds += 1

        else:

            prudential_missing += 1

            fund_info = None

            unresolved_messages.append(
                f"Row {row}: Prudential fund metadata missing."
            )

        # ================================================================
        # RESEARCH
        # ================================================================

        research = (
            get_research_for_fund(
                fund_research,
                prudential,
                excel,
            )
        )

        if any(
            value
            for value in research.values()
        ):
            research_matched += 1
        else:
            research_missing += 1

        # ================================================================
        # HOLDINGS
        # ================================================================

        if holdings is None:

            holdings_unresolved += 1

            unresolved_messages.append(
                f"Row {row}: holdings unresolved "
                "after Recovery 3."
            )

            top_holdings = {
                "status": "unresolved",
                "source": None,
                "parser": None,
                "factsheetUrl": None,
                "factsheetDocumentDate": None,
                "factsheetDataAsAt": None,
                "count": 0,
                "holdings": [],
            }

        else:

            top_holdings = {
                key: value
                for key, value in holdings.items()
                if key != "fundName"
            }

            source = clean_text(
                holdings.get(
                    "source"
                )
            )

            if source:

                holdings_by_source[
                    source
                ] = (
                    holdings_by_source.get(
                        source,
                        0,
                    )
                    + 1
                )

            if (
                holdings.get(
                    "status"
                )
                == "published"
            ):
                holdings_published += 1
            elif (
                holdings.get(
                    "status"
                )
                == "no_holdings_section"
            ):
                holdings_no_section += 1

        # ================================================================
        # FUND RECORD
        # ================================================================

        fund_record = {
            **identity,
            "fund": fund_info,
            "research": research,
            "topHoldings": top_holdings,
        }

        fund_records.append(
            fund_record
        )

        # ================================================================
        # BID HISTORY
        # ================================================================

        if bid_history_source is None:

            bid_missing += 1

            unresolved_messages.append(
                f"Row {row}: BID history missing."
            )

            bid_history = {
                "status": "unresolved",
                "priceType": "BID",
                "pruAccessFundId": None,
                "currency": None,
                "startDate": None,
                "endDate": None,
                "observationCount": 0,
                "observations": [],
            }

        else:

            try:

                observations = (
                    validate_bid_history(
                        bid_history_source,
                        f"row {row} BID",
                    )
                )

                bid_success += 1

                total_bid_observations += (
                    len(
                        observations
                    )
                )

                bid_history = {
                    "status": "success",
                    "priceType": "BID",
                    "pruAccessFundId": (
                        bid_history_source.get(
                            "fundId"
                        )
                    ),
                    "currency": (
                        bid_history_source.get(
                            "currency"
                        )
                    ),
                    "startDate": (
                        bid_history_source.get(
                            "startDate"
                        )
                    ),
                    "endDate": (
                        bid_history_source.get(
                            "endDate"
                        )
                    ),
                    "observationCount": (
                        len(
                            observations
                        )
                    ),
                    "observations": observations,
                }

            except Exception as error:

                bid_missing += 1

                unresolved_messages.append(
                    f"Row {row}: BID history invalid: "
                    f"{error}"
                )

                bid_history = {
                    "status": "unresolved",
                    "priceType": "BID",
                    "pruAccessFundId": None,
                    "currency": None,
                    "startDate": None,
                    "endDate": None,
                    "observationCount": 0,
                    "observations": [],
                }

        bid_records.append(
            {
                **identity,
                "bidHistory": bid_history,
            }
        )

    print_payment_mode_diagnostics(
        len(fund_records)
    )

    # -----------------------------------------------------------------------
    # FAILED FUNDS
    # -----------------------------------------------------------------------

    failure_report: list[dict] = []

    if unresolved_messages:

        print()
        print(
            "=" * 72
        )
        print(
            "UNRESOLVED ITEMS"
        )
        print(
            "=" * 72
        )

        for message in (
            unresolved_messages
        ):
            print(
                f" - {message}"
            )

    if ALLOW_UNRESOLVED:

        # Old behaviour: publish failed funds with their partial data.
        if unresolved_messages:
            print()
            print("ALLOW_UNRESOLVED is set: publishing partial records.")

    else:

        # Successful funds are saved; failed funds keep their last good
        # data (or are left out) and are reported.
        fund_records, bid_records, failure_report = apply_failure_policy(
            excel_funds,
            fund_records,
            bid_records,
            unresolved_messages,
            generated_at,
        )

    # Closed funds (Closed Funds.xlsx): past prices only, never ranked.
    closed_fund_records, closed_bid_records, closed_report = build_closed_funds(
        fund_research,
        generated_at,
    )

    write_failure_report(
        failure_report,
        len(fund_records),
        len(excel_funds),
        closed_published=len(closed_fund_records),
        closed_report=closed_report,
    )

    if not fund_records:

        print()
        print("BUILD FAILED: no fund could be published.")
        print("The existing published JSON is left unchanged.")

        return 1

    fund_records = fund_records + closed_fund_records
    bid_records = bid_records + closed_bid_records

    # -----------------------------------------------------------------------
    # OUTPUT ENVELOPES
    # -----------------------------------------------------------------------

    research_envelope = {
        "source": (
            "Research Funds.xlsx"
        ),
        "fundResearchWorksheet": (
            research_data[
                "fundResearchWorksheet"
            ]
        ),
        "masterCategoryWorksheets": {
            "geography": (
                research_data[
                    "geographyWorksheet"
                ]
            ),
            "sector": (
                research_data[
                    "sectorWorksheet"
                ]
            ),
        },
        "masterCategories": {
            "geography": geography_categories,
            "sector": sector_categories,
        },
    }

    funds_output = {
        "schemaVersion": SCHEMA_VERSION,
        "generatedAtUtc": generated_at,
        "source": EXCEL_FILE.name,
        "fundCount": len(
            fund_records
        ),
        "research": research_envelope,
        "summary": {
            "fundCount": len(
                fund_records
            ),
            "holdingsPublished": (
                holdings_published
            ),
            "noHoldingsSection": (
                holdings_no_section
            ),
            "holdingsUnresolved": (
                holdings_unresolved
            ),
            "holdingsBySource": (
                holdings_by_source
            ),
            "researchMatched": (
                research_matched
            ),
            "researchMissing": (
                research_missing
            ),
            "prudentialMatched": (
                prudential_matched
            ),
            "prudentialMissing": (
                prudential_missing
            ),
            "dividendFunds": (
                dividend_funds
            ),
            "paymentModeFunds": (
                payment_mode_funds
            ),
            "paymentModeCounts": dict(
                PAYMENT_DIAGNOSTICS["mode_counts"]
            ),
            "bidHistoryFunds": (
                bid_success
            ),
            "bidHistoryUnresolved": (
                bid_missing
            ),
            "totalBidObservations": (
                total_bid_observations
            ),
            "unresolvedCount": (
                len(
                    unresolved_messages
                )
            ),
            "failedFundCount": len(failure_report) + len(closed_report),
            "staleFunds": [
                item["fund"]
                for item in failure_report
                if item["outcome"] == "kept_last_good"
            ],
            "skippedFunds": [
                item["fund"]
                for item in failure_report
                if item["outcome"] == "left_out"
            ],
            "fundsWithoutHoldings": [
                item["fund"]
                for item in failure_report
                if item["outcome"] == "updated_no_holdings"
            ],
            "failedFunds": failure_report + closed_report,
            "closedFundCount": len(closed_fund_records),
            "closedFunds": [
                record["fundName"] for record in closed_fund_records
            ],
        },
        "funds": fund_records,
    }

    bid_output = {
        "schemaVersion": SCHEMA_VERSION,
        "generatedAtUtc": generated_at,
        "source": EXCEL_FILE.name,
        "fundCount": len(
            bid_records
        ),
        "summary": {
            "fundCount": len(
                bid_records
            ),
            "bidHistoryFunds": (
                bid_success
            ),
            "bidHistoryUnresolved": (
                bid_missing
            ),
            "totalObservations": (
                total_bid_observations
            ),
        },
        "funds": bid_records,
    }

    # -----------------------------------------------------------------------
    # WRITE
    # -----------------------------------------------------------------------

    write_json_atomic(
        FUNDS_OUT,
        funds_output,
    )

    write_json_atomic(
        BID_OUT,
        bid_output,
        compact=True,
    )

    # -----------------------------------------------------------------------
    # FINAL SUMMARY
    # -----------------------------------------------------------------------

    print()
    print(
        "=" * 72
    )
    print(
        "BUILD COMPLETE"
    )
    print(
        "=" * 72
    )

    print(
        f"Fund count:              {len(fund_records)}"
    )

    print(
        f"Holdings published:     {holdings_published}"
    )

    print(
        f"No holdings section:    {holdings_no_section}"
    )

    print(
        f"Holdings unresolved:    {holdings_unresolved}"
    )

    print(
        f"Holdings by source:     {holdings_by_source}"
    )

    print(
        f"Research matched:       {research_matched}"
    )

    print(
        f"Research missing:       {research_missing}"
    )

    print(
        f"Geography categories:   {len(geography_categories)}"
    )

    print(
        f"Sector categories:      {len(sector_categories)}"
    )

    print(
        f"Prudential matched:     {prudential_matched}"
    )

    print(
        f"Prudential missing:     {prudential_missing}"
    )

    print(
        f"Dividend funds:         {dividend_funds}"
    )

    print(
        f"Payment-mode funds:      {payment_mode_funds}"
    )

    print(
        f"BID history funds:      {bid_success}"
    )

    print(
        f"BID unresolved:         {bid_missing}"
    )

    print(
        f"BID observations:       {total_bid_observations}"
    )

    print()
    print(
        f"Wrote: {FUNDS_OUT}"
    )

    print(
        f"Wrote: {BID_OUT}"
    )

    print(
        "=" * 72
    )

    return 0


# =============================================================================
# ENTRY POINT
# =============================================================================

if __name__ == "__main__":

    try:
        raise SystemExit(
            main()
        )

    except Exception as error:

        print(
            f"\nFATAL BUILD ERROR: {error}",
            file=sys.stderr,
        )

        raise SystemExit(
            1
        )
