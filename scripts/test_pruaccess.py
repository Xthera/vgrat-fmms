#!/usr/bin/env python3

"""
PruAccess historical BID price extractor.

MASTER SOURCE
=============

Funds Links.xlsm

Column A:
    Prudential fund URL

Column B:
    Exact PruAccess fund name

WORKFLOW
========

1. Read every populated fund from Funds Links.xlsm.
2. Use Prudential fund URL as the master universe.
3. Retrieve current Prudential fund data.
4. Retrieve official Prudential inception date.
5. Match the PruAccess fund name exactly.
6. Build historical date windows of no more than 10 years.
7. Submit every window through PruAccess using Playwright.
8. Establish the authenticated/browser session.
9. Use the same browser-context cookies for direct HTTP pagination.
10. Retrieve every expected pagination page.
11. Retrieve pagination pages in parallel.
12. Retry failed pages.
13. Validate every page individually.
14. Extract BID observations only.
15. Never fabricate missing dates.
16. Never interpolate missing prices.
17. Never estimate missing prices.
18. Reconstruct the complete history chronologically.
19. Reject duplicate historical dates.
20. Validate inception coverage.

INCEPTION RULE
==============

The oldest available PruAccess historical BID observation must be:

    - on the official Prudential inception date, OR
    - no more than 7 calendar days after the official Prudential
      inception date.

Calendar-date gaps after the first observation are allowed.

The extractor does NOT require continuous daily observations.

No missing date or price is fabricated.

If the oldest observation is more than 7 calendar days after
the official Prudential inception date, the fund fails validation.

PRICE TYPE
==========

Historical price type is BID.

The selected PruAccess Fund Price Type is never changed by the
extractor.

CURRENT BID
===========

The current Prudential BID is stored separately from historical
PruAccess BID observations.

VALIDATION
==========

Every page must succeed.

Every expected pagination page must be retrieved.

Every window must succeed.

Every window must pass validation.

Every fund must pass full-history validation.

A fund is only saved as successful after every window and the
complete historical reconstruction have passed.

If any Excel-master fund fails, the complete production run fails.

NO SYNTHETIC DATA
=================

No synthetic data.
No estimated data.
No interpolation.
No fabricated observations.
No silent deduplication.
No fuzzy matching.
"""

from __future__ import annotations

import asyncio
import json
import re
import time
from concurrent.futures import ThreadPoolExecutor
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Any
from zoneinfo import ZoneInfo

from bs4 import BeautifulSoup
from openpyxl import load_workbook
from playwright.async_api import (
    Browser,
    BrowserContext,
    Page,
    async_playwright,
)


# ============================================================
# CONFIGURATION
# ============================================================

EXCEL_FILE = Path("Funds Links.xlsm")

OUTPUT_DIR = Path("output_pruaccess")

FUNDS_OUTPUT_DIR = (
    OUTPUT_DIR
    / "funds"
)

PRUACCESS_URL = (
    "https://pruaccess.prudential.com.sg/"
    "prulinkfund/viewFundPerformance.do"
)

PAGE_SIZE = 20

MAX_PAGES = 1000

PAGE_RETRY_COUNT = 4

PAGE_RETRY_DELAY_SECONDS = 5

PAGE_TIMEOUT_MS = 120000

PARALLEL_PAGE_WORKERS = 6

INITIAL_PAGE_TIMEOUT_MS = 120000

PRUDENTIAL_WAIT_MS = 3000

PRUACCESS_INITIAL_WAIT_MS = 1500

PRUACCESS_RESULT_WAIT_MS = 2500

BROWSER_HEADLESS = True

MAX_WINDOW_YEARS = 10

# IMPORTANT:
# Historical PruAccess data may begin up to 7 calendar days
# after the official Prudential inception date.
MAX_INCEPTION_DELAY_DAYS = 7

SINGAPORE_TZ = ZoneInfo(
    "Asia/Singapore"
)


# ============================================================
# PRUACCESS SELECTORS
# ============================================================

# IMPORTANT:
#
# PruAccess contains both:
#
#   <div id="startDate" ...>
#   <input id="startDate" name="startDate" ...>
#
# Therefore "#startDate" is NOT a unique selector.
#
# We must explicitly target the actual input element.
#
PRUACCESS_START_DATE_INPUT = (
    'input[name="startDate"]'
)

PRUACCESS_END_DATE_INPUT = (
    'input[name="endDate"]'
)


# ============================================================
# BASIC HELPERS
# ============================================================

def clean_text(
    value: Any,
) -> str:

    if value is None:
        return ""

    return " ".join(
        str(value)
        .replace("\xa0", " ")
        .split()
    ).strip()


def normalize_name(
    value: Any,
) -> str:

    text = clean_text(
        value
    )

    return re.sub(
        r"\s+",
        " ",
        text,
    ).strip().casefold()


def parse_prudential_date(
    value: Any,
) -> date:

    if isinstance(
        value,
        datetime,
    ):
        return value.date()

    if isinstance(
        value,
        date,
    ):
        return value

    text = clean_text(
        value
    )

    formats = [
        "%d-%b-%Y",
        "%d/%b/%Y",
        "%d/%m/%Y",
        "%Y-%m-%d",
        "%Y/%m/%d",
        "%d-%m-%Y",
        "%d %b %Y",
    ]

    for fmt in formats:

        try:

            return datetime.strptime(
                text,
                fmt,
            ).date()

        except ValueError:

            continue

    raise ValueError(
        f"Unable to parse Prudential date: {value!r}"
    )


def parse_pruaccess_date(
    value: Any,
) -> date:

    if isinstance(
        value,
        datetime,
    ):
        return value.date()

    if isinstance(
        value,
        date,
    ):
        return value

    text = clean_text(
        value
    )

    formats = [
        "%d-%b-%Y",
        "%d-%B-%Y",
        "%d/%b/%Y",
        "%d/%B/%Y",
        "%d/%m/%Y",
        "%Y-%m-%d",
    ]

    for fmt in formats:

        try:

            return datetime.strptime(
                text,
                fmt,
            ).date()

        except ValueError:

            continue

    raise ValueError(
        f"Unable to parse PruAccess date: {value!r}"
    )


def format_pruaccess_date(
    value: date,
) -> str:

    return value.strftime(
        "%d-%b-%Y"
    )


def singapore_today() -> date:

    return datetime.now(
        SINGAPORE_TZ
    ).date()


def singapore_today_pruaccess() -> str:

    return format_pruaccess_date(
        singapore_today()
    )


def utc_now_iso() -> str:

    return datetime.now(
        ZoneInfo("UTC")
    ).isoformat()


def safe_filename(
    value: Any,
) -> str:

    text = clean_text(
        value
    )

    text = re.sub(
        r'[<>:"/\\|?*]',
        "_",
        text,
    )

    text = re.sub(
        r"\s+",
        " ",
        text,
    ).strip()

    return text or "unknown"


# ============================================================
# JSON SERIALIZATION
# ============================================================

def make_json_safe(
    value: Any,
) -> Any:
    """
    Recursively convert Python date/datetime objects into
    JSON-safe ISO strings.

    This is intentionally recursive because date objects can
    exist deep inside nested dictionaries/lists.
    """

    if isinstance(
        value,
        datetime,
    ):

        return value.isoformat()

    if isinstance(
        value,
        date,
    ):

        return value.isoformat()

    if isinstance(
        value,
        dict,
    ):

        return {
            str(key):
                make_json_safe(item)
            for key, item
            in value.items()
        }

    if isinstance(
        value,
        list,
    ):

        return [
            make_json_safe(item)
            for item
            in value
        ]

    if isinstance(
        value,
        tuple,
    ):

        return [
            make_json_safe(item)
            for item
            in value
        ]

    return value


def save_json(
    path: Path,
    data: Any,
) -> None:

    path.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    safe_data = make_json_safe(
        data
    )

    path.write_text(
        json.dumps(
            safe_data,
            indent=2,
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )


# ============================================================
# DATE WINDOW HELPERS
# ============================================================

def add_calendar_years(
    value: date,
    years: int,
) -> date:

    try:

        return value.replace(
            year=value.year + years
        )

    except ValueError:

        # Handles 29-Feb when the destination year
        # is not a leap year.
        return value.replace(
            year=value.year + years,
            day=28,
        )


def build_date_windows(
    inception_date: date,
    end_date: date,
) -> list[dict[str, Any]]:

    if end_date < inception_date:

        raise ValueError(
            "Historical end date is earlier than "
            "the Prudential inception date."
        )

    windows = []

    current_start = inception_date

    window_number = 1

    while current_start <= end_date:

        ten_year_boundary = (
            add_calendar_years(
                current_start,
                MAX_WINDOW_YEARS,
            )
            - timedelta(days=1)
        )

        current_end = min(
            ten_year_boundary,
            end_date,
        )

        if current_end < current_start:

            raise ValueError(
                "Generated invalid historical date window."
            )

        windows.append(
            {
                "windowNumber":
                    window_number,

                "startDate":
                    current_start,

                "endDate":
                    current_end,

                "startDateText":
                    format_pruaccess_date(
                        current_start
                    ),

                "endDateText":
                    format_pruaccess_date(
                        current_end
                    ),

                "calendarDays":
                    (
                        current_end
                        - current_start
                    ).days
                    + 1,
            }
        )

        current_start = (
            current_end
            + timedelta(days=1)
        )

        window_number += 1

    if not windows:

        raise ValueError(
            "No historical date windows were generated."
        )

    if (
        windows[0]["startDate"]
        != inception_date
    ):

        raise ValueError(
            "First historical window does not start "
            "at the official inception date."
        )

    if (
        windows[-1]["endDate"]
        != end_date
    ):

        raise ValueError(
            "Final historical window does not end "
            "at the requested current date."
        )

    # Verify no gaps or overlaps between windows.

    for index in range(
        1,
        len(windows),
    ):

        previous = windows[
            index - 1
        ]

        current = windows[
            index
        ]

        expected_start = (
            previous["endDate"]
            + timedelta(days=1)
        )

        if (
            current["startDate"]
            != expected_start
        ):

            raise ValueError(
                "Generated historical windows contain "
                "a gap or overlap."
            )

    return windows


# ============================================================
# EXCEL MASTER UNIVERSE
# ============================================================

def read_excel_funds() -> list[dict[str, Any]]:

    if not EXCEL_FILE.exists():

        raise FileNotFoundError(
            f"Excel master file not found: "
            f"{EXCEL_FILE}"
        )

    workbook = load_workbook(
        EXCEL_FILE,
        read_only=True,
        keep_vba=True,
        data_only=True,
    )

    try:

        worksheet = workbook[
            workbook.sheetnames[0]
        ]

        funds = []

        for row_number in range(
            2,
            worksheet.max_row + 1,
        ):

            prudential_url = clean_text(
                worksheet.cell(
                    row=row_number,
                    column=1,
                ).value
            )

            pruaccess_name = clean_text(
                worksheet.cell(
                    row=row_number,
                    column=2,
                ).value
            )

            if not prudential_url:

                continue

            funds.append(
                {
                    "excelRow":
                        row_number,

                    "prudentialUrl":
                        prudential_url,

                    "pruAccessName":
                        pruaccess_name,
                }
            )

    finally:

        workbook.close()

    if not funds:

        raise ValueError(
            "No populated Prudential fund URLs were found "
            "in Column A of Funds Links.xlsm."
        )

    print(
        "\n============================================================"
    )

    print(
        "EXCEL MASTER FUND UNIVERSE"
    )

    print(
        "============================================================"
    )

    print(
        f"Funds found: {len(funds)}"
    )

    for fund in funds:

        print(
            f"Row {fund['excelRow']}: "
            f"{fund['pruAccessName']} "
            f"| "
            f"{fund['prudentialUrl']}"
        )

    return funds


# ============================================================
# PRUDENTIAL API
# ============================================================

async def get_prudential_fund_data(
    page: Page,
    prudential_url: str,
) -> dict[str, Any]:

    captured_api_url = None

    async def response_listener(
        response,
    ):

        nonlocal captured_api_url

        url = response.url

        if (
            "ilpfunds.json"
            in url.lower()
        ):

            captured_api_url = url

    page.on(
        "response",
        response_listener,
    )

    try:

        await page.goto(
            prudential_url,
            wait_until="domcontentloaded",
            timeout=INITIAL_PAGE_TIMEOUT_MS,
        )

        await page.wait_for_timeout(
            PRUDENTIAL_WAIT_MS
        )

        if not captured_api_url:

            content = await page.content()

            match = re.search(
                r'https?[^"\']*ilpfunds\.json[^"\']*',
                content,
                re.IGNORECASE,
            )

            if match:

                captured_api_url = (
                    match.group(0)
                )

        if not captured_api_url:

            raise RuntimeError(
                "Could not identify Prudential ilpfunds.json API."
            )

        response = await page.request.get(
            captured_api_url,
            timeout=PAGE_TIMEOUT_MS,
        )

        if response.status != 200:

            raise RuntimeError(
                "Prudential fund API returned "
                f"HTTP {response.status}."
            )

        payload = await response.json()

        if isinstance(
            payload,
            dict,
        ):

            if isinstance(
                payload.get("data"),
                list,
            ):

                records = payload["data"]

            elif isinstance(
                payload.get("funds"),
                list,
            ):

                records = payload["funds"]

            else:

                records = [
                    payload
                ]

        elif isinstance(
            payload,
            list,
        ):

            records = payload

        else:

            raise RuntimeError(
                "Unexpected Prudential fund API response format."
            )

        if not records:

            raise RuntimeError(
                "Prudential fund API returned no fund records."
            )

        fund = None

        for item in records:

            if isinstance(
                item,
                dict,
            ):

                fund = item

                break

        if not fund:

            raise RuntimeError(
                "No Prudential fund object was found in API response."
            )

        def first_value(
            *keys,
        ):

            for key in keys:

                if key in fund:

                    value = fund[key]

                    if (
                        value is not None
                        and clean_text(value)
                    ):

                        return value

            return None

        fund_identifier = first_value(
            "fundIdentifier",
            "citicode",
            "citiCode",
            "fundId",
        )

        fund_name = first_value(
            "fundName",
            "name",
            "title",
        )

        inception_raw = first_value(
            "inceptionDate",
            "fundInceptionDate",
            "launchDate",
        )

        if not fund_identifier:

            raise RuntimeError(
                "Prudential API did not provide fundIdentifier."
            )

        if not fund_name:

            raise RuntimeError(
                "Prudential API did not provide fundName."
            )

        if not inception_raw:

            raise RuntimeError(
                "Prudential API did not provide inceptionDate."
            )

        inception_date = parse_prudential_date(
            inception_raw
        )

        result = {
            "fundIdentifier":
                clean_text(
                    fund_identifier
                ),

            "fundName":
                clean_text(
                    fund_name
                ),

            "fundCode":
                clean_text(
                    first_value(
                        "fundCode",
                        "code",
                    )
                ),

            "fundCurrency":
                clean_text(
                    first_value(
                        "fundCurrency",
                        "currency",
                    )
                ),

            "unitCurrency":
                clean_text(
                    first_value(
                        "unitCurrency",
                    )
                ),

            "assetClass":
                clean_text(
                    first_value(
                        "assetClass",
                    )
                ),

            "assetSubClass":
                clean_text(
                    first_value(
                        "assetSubClass",
                    )
                ),

            "riskClassification":
                clean_text(
                    first_value(
                        "riskClassification",
                        "risk",
                    )
                ),

            "bidPrice":
                first_value(
                    "bidPrice",
                    "bid",
                ),

            "offerPrice":
                first_value(
                    "offerPrice",
                    "offer",
                ),

            "valuationDate":
                first_value(
                    "valuationDate",
                ),

            "inceptionDate":
                format_pruaccess_date(
                    inception_date
                ),

            "cumulativeYtd":
                first_value(
                    "cumulativeYtd",
                ),

            "cumulative1m":
                first_value(
                    "cumulative1m",
                ),

            "cumulative3m":
                first_value(
                    "cumulative3m",
                ),

            "cumulative6m":
                first_value(
                    "cumulative6m",
                ),

            "cumulative1y":
                first_value(
                    "cumulative1y",
                ),

            "cumulative3y":
                first_value(
                    "cumulative3y",
                ),

            "cumulative5y":
                first_value(
                    "cumulative5y",
                ),

            "annualised3y":
                first_value(
                    "annualised3y",
                ),

            "annualised5y":
                first_value(
                    "annualised5y",
                ),

            "annualised10y":
                first_value(
                    "annualised10y",
                ),

            "annualisedSinceLaunch":
                first_value(
                    "annualisedSinceLaunch",
                ),

            "factsheetUrl":
                first_value(
                    "factsheetUrl",
                ),

            "prospectusUrl":
                first_value(
                    "prospectusUrl",
                ),

            "productHighlightSheetUrl":
                first_value(
                    "productHighlightSheetUrl",
                ),

            "annualReportUrl":
                first_value(
                    "annualReportUrl",
                ),

            "fundObjective":
                first_value(
                    "fundObjective",
                ),

            "investmentManager":
                first_value(
                    "investmentManager",
                ),

            "hasDividend":
                first_value(
                    "hasDividend",
                ),

            "dividendRate":
                first_value(
                    "dividendRate",
                ),

            "raw":
                fund,
        }

        return result

    finally:

        try:

            page.remove_listener(
                "response",
                response_listener,
            )

        except Exception:

            pass


# ============================================================
# PRUACCESS FUND OPTIONS
# ============================================================

async def get_fund_options(
    page: Page,
) -> list[dict[str, str]]:

    options = []

    option_elements = await page.locator(
        "#fundName option"
    ).all()

    for option in option_elements:

        text = clean_text(
            await option.inner_text()
        )

        value = clean_text(
            await option.get_attribute(
                "value"
            )
        )

        if not text and not value:

            continue

        options.append(
            {
                "text":
                    text,

                "value":
                    value,
            }
        )

    if not options:

        raise RuntimeError(
            "No PruAccess fund options were found."
        )

    return options


def find_exact_pruaccess_match(
    options: list[dict[str, str]],
    target_name: str,
) -> dict[str, str]:

    target_normalized = normalize_name(
        target_name
    )

    matches = [
        option
        for option
        in options
        if normalize_name(
            option["text"]
        )
        == target_normalized
    ]

    if not matches:

        raise ValueError(
            "No exact PruAccess fund-name match found for: "
            f"{target_name!r}"
        )

    if len(matches) > 1:

        raise ValueError(
            "Multiple exact PruAccess fund-name matches found for: "
            f"{target_name!r}"
        )

    return matches[0]


# ============================================================
# READONLY DATE INPUT
# ============================================================

async def set_readonly_input_value(
    page: Page,
    selector: str,
    value: str,
) -> None:

    locator = page.locator(
        selector
    )

    count = await locator.count()

    if count == 0:

        raise RuntimeError(
            f"Readonly date input not found: {selector}"
        )

    if count != 1:

        raise RuntimeError(
            f"Readonly date selector is not unique: "
            f"{selector!r} matched {count} elements."
        )

    await locator.evaluate(
        """
        (element, value) => {
            element.value = value;

            element.dispatchEvent(
                new Event("input", { bubbles: true })
            );

            element.dispatchEvent(
                new Event("change", { bubbles: true })
            );
        }
        """,
        value,
    )

    actual_value = clean_text(
        await locator.input_value()
    )

    if actual_value != value:

        raise RuntimeError(
            f"Failed to set {selector}. "
            f"Expected {value!r}, got {actual_value!r}."
        )


# ============================================================
# HISTORICAL HTML PARSER
# ============================================================

def extract_price_rows_from_html(
    html: str,
) -> list[dict[str, Any]]:

    soup = BeautifulSoup(
        html,
        "html.parser",
    )

    observations = []

    for table in soup.find_all(
        "table"
    ):

        for row in table.find_all(
            "tr"
        ):

            cells = row.find_all(
                ["td", "th"]
            )

            if len(cells) < 2:

                continue

            first = clean_text(
                cells[0].get_text(
                    " ",
                    strip=True,
                )
            )

            second = clean_text(
                cells[1].get_text(
                    " ",
                    strip=True,
                )
            )

            if not re.fullmatch(
                r"\d{1,2}-[A-Za-z]{3}-\d{4}",
                first,
            ):

                continue

            if not re.fullmatch(
                r"\d+(?:\.\d+)?",
                second,
            ):

                continue

            try:

                date_value = parse_pruaccess_date(
                    first
                )

            except ValueError:

                continue

            try:

                bid_value = float(
                    second
                )

            except ValueError:

                continue

            observations.append(
                {
                    "date":
                        date_value,

                    "bidPrice":
                        bid_value,
                }
            )

    return observations


# ============================================================
# PAGINATION DISCOVERY
# ============================================================

def discover_max_page(
    html: str,
) -> int | None:

    soup = BeautifulSoup(
        html,
        "html.parser",
    )

    maximum = None

    for anchor in soup.find_all(
        "a",
        href=True,
    ):

        href = clean_text(
            anchor.get(
                "href"
            )
        )

        match = re.search(
            r"page\.page=(\d+)",
            href,
            re.IGNORECASE,
        )

        if match:

            page_number = int(
                match.group(1)
            )

            if (
                maximum is None
                or page_number > maximum
            ):

                maximum = page_number

    text = clean_text(
        soup.get_text(
            " ",
            strip=True,
        )
    )

    patterns = [
        r"Page\s+\d+\s+of\s+(\d+)",
        r"Page\s+\d+\s*/\s*(\d+)",
    ]

    for pattern in patterns:

        for match in re.finditer(
            pattern,
            text,
            re.IGNORECASE,
        ):

            page_number = int(
                match.group(1)
            )

            if (
                maximum is None
                or page_number > maximum
            ):

                maximum = page_number

    return maximum


# ============================================================
# PAGE URL
# ============================================================

def build_page_url(
    base_url: str,
    fund_id: str,
    view_type: str,
    start_date: str,
    end_date: str,
    page_number: int,
) -> str:

    separator = (
        "&"
        if "?" in base_url
        else "?"
    )

    return (
        f"{base_url}"
        f"{separator}"
        f"fundId={fund_id}"
        f"&viewType={view_type}"
        f"&startDate={start_date}"
        f"&endDate={end_date}"
        f"&page.page={page_number}"
        f"&page.size={PAGE_SIZE}"
    )


# ============================================================
# PAGINATION HTML VALIDATION
# ============================================================

def validate_pagination_html(
    html: str,
    page_number: int,
) -> None:

    if not clean_text(
        html
    ):

        raise RuntimeError(
            f"Pagination page {page_number} returned empty HTML."
        )

    if (
        "PruLink Fund Price Table"
        not in html
        and
        "Fund Performance"
        not in html
    ):

        raise RuntimeError(
            f"Pagination page {page_number} does not contain "
            "the expected PruAccess fund-performance table."
        )


# ============================================================
# ASYNC HTTP PAGE RETRIEVAL
# ============================================================

async def request_page_html(
    request_context,
    url: str,
    page_number: int,
) -> tuple[str, int]:

    last_error = None

    for attempt in range(
        1,
        PAGE_RETRY_COUNT + 1,
    ):

        try:

            response = (
                await request_context.get(
                    url,
                    timeout=PAGE_TIMEOUT_MS,
                    fail_on_status_code=False,
                )
            )

            status = response.status

            if status != 200:

                raise RuntimeError(
                    f"HTTP {status}"
                )

            html = await response.text()

            validate_pagination_html(
                html,
                page_number,
            )

            return html, status

        except Exception as error:

            last_error = error

            if attempt < PAGE_RETRY_COUNT:

                print(
                    f"Page {page_number} "
                    f"attempt {attempt}/"
                    f"{PAGE_RETRY_COUNT} failed: "
                    f"{error}"
                )

                await asyncio.sleep(
                    PAGE_RETRY_DELAY_SECONDS
                )

    raise RuntimeError(
        f"Page {page_number} failed after "
        f"{PAGE_RETRY_COUNT} attempts: "
        f"{last_error}"
    )


# ============================================================
# PARALLEL PAGE RETRIEVAL
# ============================================================

async def fetch_pages_parallel(
    request_context,
    page_jobs: list[tuple[int, str]],
) -> dict[int, dict[str, Any]]:

    semaphore = asyncio.Semaphore(
        PARALLEL_PAGE_WORKERS
    )

    async def fetch_one(
        page_number: int,
        url: str,
    ):

        async with semaphore:

            html, status = (
                await request_page_html(
                    request_context,
                    url,
                    page_number,
                )
            )

            rows = (
                extract_price_rows_from_html(
                    html
                )
            )

            if not rows:

                raise RuntimeError(
                    f"Page {page_number} returned zero "
                    "historical BID observations."
                )

            return (
                page_number,
                {
                    "html":
                        html,

                    "httpStatus":
                        status,

                    "rows":
                        rows,

                    "url":
                        url,
                },
            )

    tasks = [
        asyncio.create_task(
            fetch_one(
                page_number,
                url,
            )
        )
        for page_number, url
        in page_jobs
    ]

    try:

        results = await asyncio.gather(
            *tasks
        )

    except Exception:

        for task in tasks:

            if not task.done():

                task.cancel()

        await asyncio.gather(
            *tasks,
            return_exceptions=True,
        )

        raise

    return {
        page_number:
            result
        for page_number, result
        in results
    }


# ============================================================
# EXTRACT ALL PAGINATION PAGES
# ============================================================

async def extract_all_pages(
    request_context,
    fund_id: str,
    start_date: str,
    end_date: str,
) -> tuple[
    list[dict[str, Any]],
    list[dict[str, Any]],
]:

    print(
        "\nFetching pagination:"
    )

    print(
        f"Fund ID: {fund_id}"
    )

    print(
        f"Date range: {start_date} -> {end_date}"
    )

    page_one_url = build_page_url(
        PRUACCESS_URL,
        fund_id,
        "TBL",
        start_date,
        end_date,
        1,
    )

    page_one_html, page_one_status = (
        await request_page_html(
            request_context,
            page_one_url,
            1,
        )
    )

    page_one_rows = (
        extract_price_rows_from_html(
            page_one_html
        )
    )

    if not page_one_rows:

        raise RuntimeError(
            "PruAccess page 1 returned zero observations."
        )

    page_diagnostics = [
        {
            "page":
                1,

            "url":
                page_one_url,

            "httpStatus":
                page_one_status,

            "rowCount":
                len(page_one_rows),

            "firstDate":
                min(
                    row["date"]
                    for row
                    in page_one_rows
                ),

            "lastDate":
                max(
                    row["date"]
                    for row
                    in page_one_rows
                ),

            "status":
                "success",
        }
    ]

    max_page = discover_max_page(
        page_one_html
    )

    all_page_rows = {
        1:
            page_one_rows
    }

    # --------------------------------------------------------
    # Explicit pagination count available.
    # --------------------------------------------------------

    if max_page is not None:

        if max_page < 1:

            raise RuntimeError(
                f"Invalid PruAccess maximum page: {max_page}"
            )

        if max_page > MAX_PAGES:

            raise RuntimeError(
                f"PruAccess reported {max_page} pages, "
                f"exceeding MAX_PAGES={MAX_PAGES}."
            )

        if max_page > 1:

            page_jobs = []

            for page_number in range(
                2,
                max_page + 1,
            ):

                url = build_page_url(
                    PRUACCESS_URL,
                    fund_id,
                    "TBL",
                    start_date,
                    end_date,
                    page_number,
                )

                page_jobs.append(
                    (
                        page_number,
                        url,
                    )
                )

            parallel_results = (
                await fetch_pages_parallel(
                    request_context,
                    page_jobs,
                )
            )

            for page_number in range(
                2,
                max_page + 1,
            ):

                if page_number not in parallel_results:

                    raise RuntimeError(
                        f"Expected pagination page "
                        f"{page_number} was not retrieved."
                    )

                result = parallel_results[
                    page_number
                ]

                rows = result[
                    "rows"
                ]

                all_page_rows[
                    page_number
                ] = rows

                page_diagnostics.append(
                    {
                        "page":
                            page_number,

                        "url":
                            result["url"],

                        "httpStatus":
                            result["httpStatus"],

                        "rowCount":
                            len(rows),

                        "firstDate":
                            min(
                                row["date"]
                                for row
                                in rows
                            ),

                        "lastDate":
                            max(
                                row["date"]
                                for row
                                in rows
                            ),

                        "status":
                            "success",
                    }
                )

    # --------------------------------------------------------
    # No explicit pagination count.
    # --------------------------------------------------------

    else:

        page_number = 1

        while True:

            current_rows = all_page_rows[
                page_number
            ]

            if len(current_rows) < PAGE_SIZE:

                break

            next_page = (
                page_number + 1
            )

            if next_page > MAX_PAGES:

                raise RuntimeError(
                    "PruAccess pagination exceeded "
                    f"MAX_PAGES={MAX_PAGES}."
                )

            next_url = build_page_url(
                PRUACCESS_URL,
                fund_id,
                "TBL",
                start_date,
                end_date,
                next_page,
            )

            next_html, next_status = (
                await request_page_html(
                    request_context,
                    next_url,
                    next_page,
                )
            )

            next_rows = (
                extract_price_rows_from_html(
                    next_html
                )
            )

            if not next_rows:

                raise RuntimeError(
                    f"Pagination page {next_page} "
                    "returned zero observations."
                )

            all_page_rows[
                next_page
            ] = next_rows

            page_diagnostics.append(
                {
                    "page":
                        next_page,

                    "url":
                        next_url,

                    "httpStatus":
                        next_status,

                    "rowCount":
                        len(next_rows),

                    "firstDate":
                        min(
                            row["date"]
                            for row
                            in next_rows
                        ),

                    "lastDate":
                        max(
                            row["date"]
                            for row
                            in next_rows
                        ),

                    "status":
                        "success",
                }
            )

            page_number = next_page

    # --------------------------------------------------------
    # Every expected page must exist.
    # --------------------------------------------------------

    expected_pages = list(
        range(
            1,
            max(all_page_rows.keys()) + 1,
        )
    )

    actual_pages = sorted(
        all_page_rows.keys()
    )

    if actual_pages != expected_pages:

        raise RuntimeError(
            "Pagination sequence is incomplete. "
            f"Expected {expected_pages}, "
            f"received {actual_pages}."
        )

    page_diagnostics.sort(
        key=lambda item:
            item["page"]
    )

    if [
        item["page"]
        for item
        in page_diagnostics
    ] != expected_pages:

        raise RuntimeError(
            "Pagination diagnostics are incomplete."
        )

    # --------------------------------------------------------
    # Reconstruct observations WITHOUT deduplication.
    # --------------------------------------------------------

    final_rows = []

    for page_number in expected_pages:

        final_rows.extend(
            all_page_rows[
                page_number
            ]
        )

    final_rows.sort(
        key=lambda row:
            parse_pruaccess_date(
                row["date"]
            ),
        reverse=True,
    )

    return (
        final_rows,
        page_diagnostics,
    )


# ============================================================
# PRUACCESS SEARCH FORM
# ============================================================

async def submit_pruaccess_search(
    page: Page,
    fund_id: str,
    start_date: str,
    end_date: str,
) -> dict[str, Any]:

    await page.goto(
        PRUACCESS_URL,
        wait_until="domcontentloaded",
        timeout=INITIAL_PAGE_TIMEOUT_MS,
    )

    await page.wait_for_timeout(
        PRUACCESS_INITIAL_WAIT_MS
    )

    fund_locator = page.locator(
        "#fundName"
    )

    if await fund_locator.count() == 0:

        raise RuntimeError(
            "PruAccess fundName selector not found."
        )

    await fund_locator.select_option(
        fund_id
    )

    view_type_locator = page.locator(
        "#viewType"
    )

    if await view_type_locator.count() == 0:

        raise RuntimeError(
            "PruAccess viewType selector not found."
        )

    await view_type_locator.select_option(
        "TBL"
    )

    # --------------------------------------------------------
    # IMPORTANT:
    #
    # Do NOT change the selected Fund Price Type.
    # We only inspect and record it.
    # --------------------------------------------------------

    fund_price_type = None

    fund_price_type_label = None

    fund_price_type_locator = page.locator(
        "#fundPriceType"
    )

    if (
        await fund_price_type_locator.count()
        > 0
    ):

        fund_price_type = clean_text(
            await fund_price_type_locator.input_value()
        )

        selected_option = (
            fund_price_type_locator.locator(
                "option:checked"
            )
        )

        if await selected_option.count() > 0:

            fund_price_type_label = clean_text(
                await selected_option.inner_text()
            )

    # --------------------------------------------------------
    # Set dates.
    #
    # IMPORTANT:
    #
    # PruAccess has duplicate IDs for the date controls:
    #
    #   div#startDate
    #   input#startDate[name="startDate"]
    #
    #   div#endDate
    #   input#endDate[name="endDate"]
    #
    # Therefore we explicitly target the input elements
    # by their name attribute.
    # --------------------------------------------------------

    await set_readonly_input_value(
        page,
        PRUACCESS_START_DATE_INPUT,
        start_date,
    )

    await set_readonly_input_value(
        page,
        PRUACCESS_END_DATE_INPUT,
        end_date,
    )

    # --------------------------------------------------------
    # CSRF.
    # --------------------------------------------------------

    csrf_locator = page.locator(
        "input[name='_csrf']"
    )

    if await csrf_locator.count() == 0:

        raise RuntimeError(
            "PruAccess CSRF input was not found."
        )

    csrf_value = clean_text(
        await csrf_locator.input_value()
    )

    if not csrf_value:

        raise RuntimeError(
            "PruAccess CSRF value is empty."
        )

    # --------------------------------------------------------
    # Verify requested dates.
    #
    # IMPORTANT:
    # Use the same unique input[name=...] selectors.
    # --------------------------------------------------------

    start_date_locator = page.locator(
        PRUACCESS_START_DATE_INPUT
    )

    end_date_locator = page.locator(
        PRUACCESS_END_DATE_INPUT
    )

    if await start_date_locator.count() != 1:

        raise RuntimeError(
            "PruAccess start-date input selector is not unique."
        )

    if await end_date_locator.count() != 1:

        raise RuntimeError(
            "PruAccess end-date input selector is not unique."
        )

    actual_start = clean_text(
        await start_date_locator.input_value()
    )

    actual_end = clean_text(
        await end_date_locator.input_value()
    )

    if actual_start != start_date:

        raise RuntimeError(
            "PruAccess start date verification failed. "
            f"Expected {start_date!r}, "
            f"got {actual_start!r}."
        )

    if actual_end != end_date:

        raise RuntimeError(
            "PruAccess end date verification failed. "
            f"Expected {end_date!r}, "
            f"got {actual_end!r}."
        )

    # --------------------------------------------------------
    # Verify the requested window does not exceed 10 years.
    # --------------------------------------------------------

    start_value = parse_pruaccess_date(
        start_date
    )

    end_value = parse_pruaccess_date(
        end_date
    )

    ten_year_limit = (
        add_calendar_years(
            start_value,
            MAX_WINDOW_YEARS,
        )
    )

    if end_value >= ten_year_limit:

        raise RuntimeError(
            "PruAccess historical window exceeds "
            f"the maximum {MAX_WINDOW_YEARS}-year limit."
        )

    # --------------------------------------------------------
    # Submit form.
    # --------------------------------------------------------

    form_locator = page.locator(
        "#fundForm"
    )

    if await form_locator.count() == 0:

        raise RuntimeError(
            "PruAccess fund form was not found."
        )

    await form_locator.evaluate(
        """
        form => form.submit()
        """
    )

    await page.wait_for_load_state(
        "domcontentloaded",
        timeout=INITIAL_PAGE_TIMEOUT_MS,
    )

    await page.wait_for_timeout(
        PRUACCESS_RESULT_WAIT_MS
    )

    body_text = clean_text(
        await page.locator(
            "body"
        ).inner_text()
    )

    if (
        "PruLink Fund Price Table"
        not in body_text
    ):

        raise RuntimeError(
            "PruAccess search did not return the expected "
            "Fund Price Table."
        )

    return {
        "fundPriceType":
            fund_price_type,

        "fundPriceTypeLabel":
            fund_price_type_label,

        "csrfPresent":
            bool(
                csrf_value
            ),

        "startDate":
            start_date,

        "endDate":
            end_date,

        "viewType":
            "TBL",

        "fundId":
            fund_id,
    }


# ============================================================
# DUPLICATE VALIDATION
# ============================================================

def find_duplicate_dates(
    observations: list[dict[str, Any]],
) -> list[dict[str, Any]]:

    grouped = {}

    for observation in observations:

        observation_date = parse_pruaccess_date(
            observation["date"]
        )

        key = observation_date.isoformat()

        grouped.setdefault(
            key,
            [],
        ).append(
            observation
        )

    duplicates = []

    for key, rows in grouped.items():

        if len(rows) > 1:

            duplicates.append(
                {
                    "date":
                        parse_pruaccess_date(
                            key
                        ),

                    "count":
                        len(rows),

                    "observations":
                        rows,
                }
            )

    duplicates.sort(
        key=lambda item:
            parse_pruaccess_date(
                item["date"]
            )
    )

    return duplicates


# ============================================================
# CHRONOLOGICAL RECONSTRUCTION
# ============================================================

def reconstruct_history(
    window_results: list[dict[str, Any]],
) -> tuple[
    list[dict[str, Any]],
    list[dict[str, Any]],
]:

    observations = []

    for window_result in window_results:

        observations.extend(
            window_result[
                "observations"
            ]
        )

    if not observations:

        raise ValueError(
            "No historical observations were extracted."
        )

    duplicate_dates = (
        find_duplicate_dates(
            observations
        )
    )

    chronological = sorted(
        observations,
        key=lambda row:
            parse_pruaccess_date(
                row["date"]
            ),
    )

    return (
        chronological,
        duplicate_dates,
    )


# ============================================================
# FULL HISTORY VALIDATION
# ============================================================

def validate_full_history_to_inception(
    historical_rows: list[dict[str, Any]],
    inception_date: date,
    end_date: date,
    duplicate_dates: list[dict[str, Any]],
) -> dict[str, Any]:

    if not historical_rows:

        raise ValueError(
            "Historical BID history is empty."
        )

    if duplicate_dates:

        raise ValueError(
            "Duplicate historical BID dates detected: "
            f"{len(duplicate_dates)} duplicate date(s)."
        )

    parsed_dates = [
        parse_pruaccess_date(
            row["date"]
        )
        for row
        in historical_rows
    ]

    oldest_date = min(
        parsed_dates
    )

    newest_date = max(
        parsed_dates
    )

    # --------------------------------------------------------
    # HARD INCEPTION RULE
    #
    # Oldest observation must be:
    #
    # inception date
    # OR
    # within +7 calendar days.
    # --------------------------------------------------------

    latest_allowed_oldest_date = (
        inception_date
        + timedelta(
            days=MAX_INCEPTION_DELAY_DAYS
        )
    )

    if oldest_date < inception_date:

        raise ValueError(
            "Historical PruAccess data contains an observation "
            "before the official Prudential inception date. "
            f"Inception={format_pruaccess_date(inception_date)}, "
            f"oldest={format_pruaccess_date(oldest_date)}"
        )

    if oldest_date > latest_allowed_oldest_date:

        raise ValueError(
            "Historical PruAccess data does not reach the "
            "Prudential inception date within the permitted "
            f"+{MAX_INCEPTION_DELAY_DAYS} calendar days. "
            f"Inception={format_pruaccess_date(inception_date)}, "
            f"oldest={format_pruaccess_date(oldest_date)}"
        )

    if newest_date > end_date:

        raise ValueError(
            "Historical PruAccess data contains an observation "
            "after the requested historical end date. "
            f"End={format_pruaccess_date(end_date)}, "
            f"newest={format_pruaccess_date(newest_date)}"
        )

    # --------------------------------------------------------
    # Chronological ordering validation.
    # --------------------------------------------------------

    for index in range(
        1,
        len(parsed_dates),
    ):

        if (
            parsed_dates[index]
            < parsed_dates[index - 1]
        ):

            raise ValueError(
                "Historical BID observations are not "
                "chronologically ordered."
            )

    # --------------------------------------------------------
    # Date gaps are allowed.
    #
    # We record them for diagnostics but NEVER fill them.
    # --------------------------------------------------------

    date_gaps = []

    for index in range(
        1,
        len(parsed_dates),
    ):

        previous_date = (
            parsed_dates[index - 1]
        )

        current_date = (
            parsed_dates[index]
        )

        gap_days = (
            current_date
            - previous_date
        ).days

        if gap_days > 1:

            date_gaps.append(
                {
                    "fromDate":
                        previous_date,

                    "toDate":
                        current_date,

                    "missingCalendarDays":
                        gap_days - 1,
                }
            )

    inception_delay_days = (
        oldest_date
        - inception_date
    ).days

    if inception_delay_days == 0:

        inception_coverage = "exact"

    else:

        inception_coverage = (
            f"inception+"
            f"{inception_delay_days}"
            f"-days"
        )

    return {
        "status":
            "passed",

        "inceptionDate":
            format_pruaccess_date(
                inception_date
            ),

        "oldestHistoricalDate":
            format_pruaccess_date(
                oldest_date
            ),

        "newestHistoricalDate":
            format_pruaccess_date(
                newest_date
            ),

        "allowedLatestOldestDate":
            format_pruaccess_date(
                latest_allowed_oldest_date
            ),

        "maximumAllowedDaysAfterInception":
            MAX_INCEPTION_DELAY_DAYS,

        "inceptionCoverage":
            inception_coverage,

        "inceptionDelayDays":
            inception_delay_days,

        "observationCount":
            len(
                historical_rows
            ),

        "duplicateDateCount":
            len(
                duplicate_dates
            ),

        "dateGapCount":
            len(
                date_gaps
            ),

        "dateGaps":
            date_gaps,
    }


# ============================================================
# WINDOW VALIDATION
# ============================================================

def validate_window_coverage(
    window: dict[str, Any],
    observations: list[dict[str, Any]],
) -> dict[str, Any]:

    if not observations:

        raise ValueError(
            f"Historical window "
            f"{window['windowNumber']} returned no observations."
        )

    window_start = (
        window["startDate"]
    )

    window_end = (
        window["endDate"]
    )

    parsed_dates = []

    for observation in observations:

        observation_date = parse_pruaccess_date(
            observation["date"]
        )

        if (
            observation_date
            < window_start
        ):

            raise ValueError(
                f"Observation date "
                f"{format_pruaccess_date(observation_date)} "
                f"is before requested window start "
                f"{format_pruaccess_date(window_start)}."
            )

        if (
            observation_date
            > window_end
        ):

            raise ValueError(
                f"Observation date "
                f"{format_pruaccess_date(observation_date)} "
                f"is after requested window end "
                f"{format_pruaccess_date(window_end)}."
            )

        parsed_dates.append(
            observation_date
        )

    duplicate_dates = (
        find_duplicate_dates(
            observations
        )
    )

    if duplicate_dates:

        raise ValueError(
            f"Duplicate dates found inside window "
            f"{window['windowNumber']}: "
            f"{len(duplicate_dates)}."
        )

    return {
        "status":
            "passed",

        "requestedStartDate":
            format_pruaccess_date(
                window_start
            ),

        "requestedEndDate":
            format_pruaccess_date(
                window_end
            ),

        "oldestObservation":
            format_pruaccess_date(
                min(
                    parsed_dates
                )
            ),

        "newestObservation":
            format_pruaccess_date(
                max(
                    parsed_dates
                )
            ),

        "observationCount":
            len(
                observations
            ),

        "duplicateDateCount":
            0,
    }


# ============================================================
# SINGLE FUND EXTRACTION
# ============================================================

async def extract_single_fund(
    context: BrowserContext,
    page: Page,
    excel_fund: dict[str, Any],
    pruaccess_options: list[dict[str, str]],
) -> dict[str, Any]:

    excel_row = (
        excel_fund[
            "excelRow"
        ]
    )

    prudential_url = (
        excel_fund[
            "prudentialUrl"
        ]
    )

    excel_pruaccess_name = (
        excel_fund[
            "pruAccessName"
        ]
    )

    # ========================================================
    # PRUDENTIAL FUND DATA
    # ========================================================

    prudential = (
        await get_prudential_fund_data(
            page,
            prudential_url,
        )
    )

    print(
        f"\nPrudential fund: "
        f"{prudential['fundName']}"
    )

    print(
        f"Fund identifier: "
        f"{prudential['fundIdentifier']}"
    )

    print(
        f"Fund code: "
        f"{prudential.get('fundCode')}"
    )

    print(
        f"Current BID: "
        f"{prudential.get('bidPrice')}"
    )

    print(
        f"Inception: "
        f"{prudential['inceptionDate']}"
    )

    inception_date = parse_prudential_date(
        prudential[
            "inceptionDate"
        ]
    )

    end_date = singapore_today()

    # ========================================================
    # DATE WINDOWS
    # ========================================================

    windows = build_date_windows(
        inception_date,
        end_date,
    )

    print(
        "\nHistorical windows:"
    )

    for window in windows:

        print(
            f"Window "
            f"{window['windowNumber']}: "
            f"{window['startDateText']} -> "
            f"{window['endDateText']} "
            f"({window['calendarDays']} calendar days)"
        )

    # ========================================================
    # EXACT PRUACCESS MATCH
    # ========================================================

    selected = find_exact_pruaccess_match(
        pruaccess_options,
        excel_pruaccess_name,
    )

    fund_id = clean_text(
        selected[
            "value"
        ]
    )

    if not fund_id:

        raise ValueError(
            "Exact PruAccess fund match was found, "
            "but its fund ID/value is empty."
        )

    print(
        "\nExact PruAccess match:"
    )

    print(
        f"Name: {selected['text']}"
    )

    print(
        f"Fund ID: {fund_id}"
    )

    # ========================================================
    # WINDOW EXTRACTION
    # ========================================================

    window_results = []

    window_diagnostics = []

    request_context = (
        context.request
    )

    for window in windows:

        window_number = (
            window[
                "windowNumber"
            ]
        )

        start_date = (
            window[
                "startDateText"
            ]
        )

        end_date_text = (
            window[
                "endDateText"
            ]
        )

        print(
            "\n============================================================"
        )

        print(
            f"WINDOW {window_number} / "
            f"{len(windows)}"
        )

        print(
            f"{start_date} -> {end_date_text}"
        )

        print(
            "------------------------------------------------------------"
        )

        # ----------------------------------------------------
        # Establish PruAccess session for this window.
        # ----------------------------------------------------

        pre_submit = (
            await submit_pruaccess_search(
                page,
                fund_id,
                start_date,
                end_date_text,
            )
        )

        # ----------------------------------------------------
        # Direct HTTP pagination using the same browser
        # context session cookies.
        # ----------------------------------------------------

        observations, page_diagnostics = (
            await extract_all_pages(
                request_context,
                fund_id,
                start_date,
                end_date_text,
            )
        )

        # ----------------------------------------------------
        # Window validation.
        # ----------------------------------------------------

        window_validation = (
            validate_window_coverage(
                window,
                observations,
            )
        )

        window_result = {
            "windowNumber":
                window_number,

            "startDate":
                start_date,

            "endDate":
                end_date_text,

            "fundId":
                fund_id,

            "fundPriceType":
                pre_submit.get(
                    "fundPriceType"
                ),

            "fundPriceTypeLabel":
                pre_submit.get(
                    "fundPriceTypeLabel"
                ),

            "historicalPriceType":
                "BID",

            "pageSize":
                PAGE_SIZE,

            "pagesExtracted":
                len(
                    page_diagnostics
                ),

            "observationCount":
                len(
                    observations
                ),

            "observations":
                observations,

            "pagination":
                page_diagnostics,

            "validation":
                window_validation,
        }

        window_results.append(
            window_result
        )

        window_diagnostics.append(
            {
                "windowNumber":
                    window_number,

                "startDate":
                    start_date,

                "endDate":
                    end_date_text,

                "pagesExtracted":
                    len(
                        page_diagnostics
                    ),

                "observationCount":
                    len(
                        observations
                    ),

                "oldestObservation":
                    window_validation[
                        "oldestObservation"
                    ],

                "newestObservation":
                    window_validation[
                        "newestObservation"
                    ],

                "status":
                    "success",
            }
        )

        print(
            f"\nWINDOW {window_number} SUCCESS"
        )

        print(
            f"Pages: "
            f"{len(page_diagnostics)}"
        )

        print(
            f"Observations: "
            f"{len(observations)}"
        )

    # ========================================================
    # CHRONOLOGICAL RECONSTRUCTION
    # ========================================================

    (
        chronological_history,
        duplicate_dates,
    ) = reconstruct_history(
        window_results
    )

    print(
        "\n============================================================"
    )

    print(
        "CHRONOLOGICAL RECONSTRUCTION"
    )

    print(
        "============================================================"
    )

    print(
        f"Total observations: "
        f"{len(chronological_history)}"
    )

    print(
        f"Duplicate dates: "
        f"{len(duplicate_dates)}"
    )

    # ========================================================
    # FULL HISTORY VALIDATION
    # ========================================================

    history_validation = (
        validate_full_history_to_inception(
            chronological_history,
            inception_date,
            end_date,
            duplicate_dates,
        )
    )

    print(
        "\n============================================================"
    )

    print(
        "FULL-HISTORY VALIDATION PASSED"
    )

    print(
        "============================================================"
    )

    print(
        f"Inception: "
        f"{history_validation['inceptionDate']}"
    )

    print(
        f"Oldest BID: "
        f"{history_validation['oldestHistoricalDate']}"
    )

    print(
        f"Newest BID: "
        f"{history_validation['newestHistoricalDate']}"
    )

    print(
        f"Inception coverage: "
        f"{history_validation['inceptionCoverage']}"
    )

    print(
        f"Date gaps: "
        f"{history_validation['dateGapCount']}"
    )

    # ========================================================
    # FINAL BID HISTORY
    # ========================================================

    bid_history = {
        "source":
            "PruAccess",

        "fundName":
            prudential["fundName"],

        "pruAccessFundName":
            selected["text"],

        "fundId":
            fund_id,

        "fundIdentifier":
            prudential["fundIdentifier"],

        "fundCode":
            prudential.get(
                "fundCode"
            ),

        "currency":
            prudential.get(
                "fundCurrency"
            ),

        "startDate":
            format_pruaccess_date(
                inception_date
            ),

        "endDate":
            format_pruaccess_date(
                end_date
            ),

        "priceType":
            "BID",

        "pageSize":
            PAGE_SIZE,

        "windowCount":
            len(
                windows
            ),

        "observationCount":
            len(
                chronological_history
            ),

        "observations":
            chronological_history,
    }

    # ========================================================
    # SUMMARY
    # ========================================================

    summary = {
        "status":
            "success",

        "excelRow":
            excel_row,

        "prudentialUrl":
            prudential_url,

        "prudentialFundName":
            prudential["fundName"],

        "excelPruAccessName":
            excel_pruaccess_name,

        "matchedPruAccessName":
            selected["text"],

        "pruAccessFundId":
            fund_id,

        "fundIdentifier":
            prudential["fundIdentifier"],

        "fundCode":
            prudential.get(
                "fundCode"
            ),

        "inceptionDate":
            prudential["inceptionDate"],

        "startDate":
            format_pruaccess_date(
                inception_date
            ),

        "endDate":
            format_pruaccess_date(
                end_date
            ),

        "currentBidPrice":
            prudential.get(
                "bidPrice"
            ),

        "currentOfferPrice":
            prudential.get(
                "offerPrice"
            ),

        "historicalPriceType":
            "BID",

        "pageSize":
            PAGE_SIZE,

        "windowCount":
            len(
                windows
            ),

        "pagesExtracted":
            sum(
                item["pagesExtracted"]
                for item
                in window_diagnostics
            ),

        "historicalObservationCount":
            len(
                chronological_history
            ),

        "newestObservation":
            chronological_history[-1],

        "oldestObservation":
            chronological_history[0],

        "inceptionValidation":
            history_validation,

        "windows":
            window_diagnostics,
    }

    # ========================================================
    # FUND RECORD
    # ========================================================

    return {
        "status":
            "success",

        "excelRow":
            excel_row,

        "prudentialUrl":
            prudential_url,

        "excelPruAccessName":
            excel_pruaccess_name,

        "prudential":
            prudential,

        "pruAccess": {
            "fundName":
                selected["text"],

            "fundId":
                fund_id,

            "startDate":
                format_pruaccess_date(
                    inception_date
                ),

            "endDate":
                format_pruaccess_date(
                    end_date
                ),

            "priceType":
                "BID",

            "pageSize":
                PAGE_SIZE,

            "windowCount":
                len(
                    windows
                ),

            "observationCount":
                len(
                    chronological_history
                ),
        },

        "dateWindows":
            windows,

        "windowResults":
            window_results,

        "windowDiagnostics":
            window_diagnostics,

        "inceptionValidation":
            history_validation,

        "summary":
            summary,

        "bidHistory":
            bid_history,
    }


# ============================================================
# SAVE SUCCESSFUL FUND
# ============================================================

def save_successful_fund(
    result: dict,
) -> None:

    prudential = result[
        "prudential"
    ]

    excel_row = result[
        "excelRow"
    ]

    fund_identifier = clean_text(
        prudential.get(
            "fundIdentifier"
        )
    )

    fund_code = clean_text(
        prudential.get(
            "fundCode"
        )
    )

    fund_name = clean_text(
        prudential.get(
            "fundName"
        )
    )

    safe_fund_name = safe_filename(
        fund_identifier
        or fund_code
        or fund_name
    )

    directory_name = (
        f"{excel_row}_"
        f"{safe_fund_name}"
    )

    fund_output_dir = (
        FUNDS_OUTPUT_DIR
        / directory_name
    )

    fund_output_dir.mkdir(
        parents=True,
        exist_ok=True,
    )

    save_json(
        fund_output_dir
        / "prudential_fund.json",
        prudential,
    )

    save_json(
        fund_output_dir
        / "date_windows.json",
        result["dateWindows"],
    )

    save_json(
        fund_output_dir
        / "window_diagnostics.json",
        result["windowDiagnostics"],
    )

    save_json(
        fund_output_dir
        / "inception_validation.json",
        result["inceptionValidation"],
    )

    save_json(
        fund_output_dir
        / "window_results.json",
        result["windowResults"],
    )

    save_json(
        fund_output_dir
        / "bid_history.json",
        result["bidHistory"],
    )

    save_json(
        fund_output_dir
        / "summary.json",
        result["summary"],
    )


# ============================================================
# MAIN
# ============================================================

async def main():

    OUTPUT_DIR.mkdir(
        parents=True,
        exist_ok=True,
    )

    FUNDS_OUTPUT_DIR.mkdir(
        parents=True,
        exist_ok=True,
    )

    print(
        "============================================================"
    )

    print(
        "PRUACCESS HISTORICAL BID PRICE EXTRACTION"
    )

    print(
        "10-YEAR WINDOW + PARALLEL PAGINATION VERSION"
    )

    print(
        "============================================================"
    )

    # ========================================================
    # READ MASTER EXCEL
    # ========================================================

    funds = read_excel_funds()

    # ========================================================
    # RUN METADATA
    # ========================================================

    run_started = utc_now_iso()

    successful_funds = []

    failed_funds = []

    all_bid_history = {}

    # ========================================================
    # PLAYWRIGHT
    # ========================================================

    async with async_playwright() as playwright:

        browser: Browser = (
            await playwright.chromium.launch(
                headless=BROWSER_HEADLESS
            )
        )

        context: BrowserContext = (
            await browser.new_context(
                viewport={
                    "width": 1440,
                    "height": 1000,
                }
            )
        )

        page = await context.new_page()

        # ====================================================
        # LOAD PRUACCESS OPTIONS ONCE
        # ====================================================

        print(
            "\n============================================================"
        )

        print(
            "LOADING PRUACCESS FUND OPTIONS"
        )

        print(
            "============================================================"
        )

        await page.goto(
            PRUACCESS_URL,
            wait_until="domcontentloaded",
            timeout=INITIAL_PAGE_TIMEOUT_MS,
        )

        await page.wait_for_timeout(
            PRUACCESS_INITIAL_WAIT_MS
        )

        pruaccess_options = (
            await get_fund_options(
                page
            )
        )

        save_json(
            OUTPUT_DIR
            / "fund_options.json",
            {
                "source":
                    PRUACCESS_URL,

                "retrievedAtUtc":
                    utc_now_iso(),

                "count":
                    len(
                        pruaccess_options
                    ),

                "options":
                    pruaccess_options,
            },
        )

        # ====================================================
        # PROCESS EVERY EXCEL FUND
        # ====================================================

        total_funds = len(
            funds
        )

        for fund_index, excel_fund in enumerate(
            funds,
            start=1,
        ):

            print(
                "\n\n"
                "############################################################"
            )

            print(
                f"FUND {fund_index} / {total_funds}"
            )

            print(
                f"EXCEL ROW {excel_fund['excelRow']}"
            )

            print(
                "############################################################"
            )

            try:

                result = (
                    await extract_single_fund(
                        context,
                        page,
                        excel_fund,
                        pruaccess_options,
                    )
                )

                # ------------------------------------------------
                # DO NOT save the fund until every window and
                # full-history validation has passed.
                # ------------------------------------------------

                save_successful_fund(
                    result
                )

                successful_funds.append(
                    result
                )

                fund_identifier = clean_text(
                    result["prudential"].get(
                        "fundIdentifier"
                    )
                )

                fund_code = clean_text(
                    result["prudential"].get(
                        "fundCode"
                    )
                )

                history_key = (
                    fund_identifier
                    or fund_code
                    or str(
                        excel_fund[
                            "excelRow"
                        ]
                    )
                )

                all_bid_history[
                    history_key
                ] = result[
                    "bidHistory"
                ]

                print(
                    "\n============================================================"
                )

                print(
                    "FUND SUCCESS"
                )

                print(
                    "============================================================"
                )

                print(
                    f"Fund: "
                    f"{result['prudential']['fundName']}"
                )

                print(
                    f"Identifier: "
                    f"{fund_identifier}"
                )

                print(
                    f"PruAccess ID: "
                    f"{result['pruAccess']['fundId']}"
                )

                print(
                    f"Windows: "
                    f"{result['pruAccess']['windowCount']}"
                )

                print(
                    f"Pages: "
                    f"{result['summary']['pagesExtracted']}"
                )

                print(
                    f"Historical BID observations: "
                    f"{result['pruAccess']['observationCount']}"
                )

                print(
                    f"Oldest BID: "
                    f"{result['summary']['oldestObservation']}"
                )

                print(
                    f"Newest BID: "
                    f"{result['summary']['newestObservation']}"
                )

            except Exception as error:

                error_text = clean_text(
                    str(error)
                )

                print(
                    "\n============================================================"
                )

                print(
                    "FUND FAILED"
                )

                print(
                    "============================================================"
                )

                print(
                    f"Excel row: "
                    f"{excel_fund['excelRow']}"
                )

                print(
                    f"Prudential URL: "
                    f"{excel_fund['prudentialUrl']}"
                )

                print(
                    f"PruAccess name: "
                    f"{excel_fund['pruAccessName']}"
                )

                print(
                    f"Error: "
                    f"{error_text}"
                )

                failure = {
                    "status":
                        "failed",

                    "excelRow":
                        excel_fund[
                            "excelRow"
                        ],

                    "prudentialUrl":
                        excel_fund[
                            "prudentialUrl"
                        ],

                    "excelPruAccessName":
                        excel_fund[
                            "pruAccessName"
                        ],

                    "error":
                        error_text,
                }

                failed_funds.append(
                    failure
                )

                failure_dir_name = (
                    f"{excel_fund['excelRow']}_failed"
                )

                failure_output_dir = (
                    FUNDS_OUTPUT_DIR
                    / failure_dir_name
                )

                failure_output_dir.mkdir(
                    parents=True,
                    exist_ok=True,
                )

                save_json(
                    failure_output_dir
                    / "failure.json",
                    failure,
                )

                continue

        await browser.close()

    # ========================================================
    # CONSOLIDATED DATA
    # ========================================================

    run_finished = utc_now_iso()

    total_observations = sum(
        fund["pruAccess"][
            "observationCount"
        ]
        for fund in successful_funds
    )

    total_windows = sum(
        fund["pruAccess"][
            "windowCount"
        ]
        for fund in successful_funds
    )

    total_pages = sum(
        fund["summary"][
            "pagesExtracted"
        ]
        for fund in successful_funds
    )

    # ========================================================
    # ALL FUNDS
    # ========================================================

    all_funds = {
        "source":
            "Funds Links.xlsm",

        "generatedAtUtc":
            run_finished,

        "fundUniverseCount":
            len(
                funds
            ),

        "successfulFundCount":
            len(
                successful_funds
            ),

        "failedFundCount":
            len(
                failed_funds
            ),

        "funds":
            successful_funds,

        "failed":
            failed_funds,

        "failedFunds":
            failed_funds,
    }

    save_json(
        OUTPUT_DIR
        / "all_funds.json",
        all_funds,
    )

    # ========================================================
    # ALL BID HISTORY
    # ========================================================

    consolidated_bid_history = {
        "source":
            "PruAccess",

        "priceType":
            "BID",

        "generatedAtUtc":
            run_finished,

        "fundCount":
            len(
                all_bid_history
            ),

        "totalWindows":
            total_windows,

        "totalPages":
            total_pages,

        "totalHistoricalBidObservations":
            total_observations,

        "funds":
            all_bid_history,
    }

    save_json(
        OUTPUT_DIR
        / "all_bid_history.json",
        consolidated_bid_history,
    )

    # ========================================================
    # RUN SUMMARY
    #
    # IMPORTANT:
    #
    # Any failed fund means the complete run is NOT a
    # successful production extraction.
    # ========================================================

    run_status = (
        "success"
        if (
            not failed_funds
            and len(successful_funds)
            == len(funds)
        )
        else "failed"
    )

    run_summary = {
        "status":
            run_status,

        "startedAtUtc":
            run_started,

        "completedAtUtc":
            run_finished,

        "excelFile":
            str(
                EXCEL_FILE
            ),

        "fundUniverseCount":
            len(
                funds
            ),

        "successfulFundCount":
            len(
                successful_funds
            ),

        "failedFundCount":
            len(
                failed_funds
            ),

        "totalWindows":
            total_windows,

        "totalPages":
            total_pages,

        "totalHistoricalBidObservations":
            total_observations,

        "historicalPriceType":
            "BID",

        "maximumWindowYears":
            MAX_WINDOW_YEARS,

        "pageSize":
            PAGE_SIZE,

        "parallelPageWorkers":
            PARALLEL_PAGE_WORKERS,

        "pageRetryCount":
            PAGE_RETRY_COUNT,

        "pageRetryDelaySeconds":
            PAGE_RETRY_DELAY_SECONDS,

        "pageTimeoutMs":
            PAGE_TIMEOUT_MS,

        "inceptionCoverageRule":
            {
                "oldestObservationMustBeOnOrAfterInception":
                    True,

                "maximumAllowedDaysAfterInception":
                    MAX_INCEPTION_DELAY_DAYS,

                "dateGapsAllowed":
                    True,

                "duplicateDatesAllowed":
                    False,

                "syntheticDataAllowed":
                    False,

                "interpolationAllowed":
                    False,

                "estimationAllowed":
                    False,
            },

        "successfulFunds":
            [
                {
                    "excelRow":
                        fund["excelRow"],

                    "fundName":
                        fund[
                            "prudential"
                        ].get(
                            "fundName"
                        ),

                    "fundIdentifier":
                        fund[
                            "prudential"
                        ].get(
                            "fundIdentifier"
                        ),

                    "fundCode":
                        fund[
                            "prudential"
                        ].get(
                            "fundCode"
                        ),

                    "pruAccessFundName":
                        fund[
                            "pruAccess"
                        ].get(
                            "fundName"
                        ),

                    "pruAccessFundId":
                        fund[
                            "pruAccess"
                        ].get(
                            "fundId"
                        ),

                    "observationCount":
                        fund[
                            "pruAccess"
                        ].get(
                            "observationCount"
                        ),

                    "windowCount":
                        fund[
                            "pruAccess"
                        ].get(
                            "windowCount"
                        ),

                    "pagesExtracted":
                        fund[
                            "summary"
                        ].get(
                            "pagesExtracted"
                        ),

                    "newestObservation":
                        fund[
                            "summary"
                        ].get(
                            "newestObservation"
                        ),

                    "oldestObservation":
                        fund[
                            "summary"
                        ].get(
                            "oldestObservation"
                        ),

                    "inceptionValidation":
                        fund.get(
                            "inceptionValidation"
                        ),
                }

                for fund
                in successful_funds
            ],

        "failedFunds":
            failed_funds,
    }

    save_json(
        OUTPUT_DIR
        / "run_summary.json",
        run_summary,
    )

    # ========================================================
    # VALIDATION SUMMARY
    # ========================================================

    validation_summary = {
        "status":
            run_status,

        "generatedAtUtc":
            run_finished,

        "excelFundUniverse":
            len(
                funds
            ),

        "successfulFunds":
            len(
                successful_funds
            ),

        "failedFunds":
            len(
                failed_funds
            ),

        "totalHistoricalBidObservations":
            total_observations,

        "totalWindows":
            total_windows,

        "totalPages":
            total_pages,

        "rules": {
            "automaticTenYearWindows":
                True,

            "parallelPageRetrieval":
                True,

            "retryHandling":
                True,

            "completePaginationValidation":
                True,

            "noSyntheticData":
                True,

            "chronologicalReconstruction":
                True,

            "fullHistoryToInceptionValidation":
                True,

            "maximumInceptionDelayDays":
                MAX_INCEPTION_DELAY_DAYS,

            "firstObservationMayBeUpToSevenDaysAfterInception":
                True,

            "dateGapsAllowed":
                True,

            "duplicateDatesCauseFailure":
                True,

            "exactPruAccessNameMatch":
                True,

            "excelMasterUniverse":
                True,

            "hardcodedFundCount":
                False,
        },

        "failedFunds":
            failed_funds,
    }

    save_json(
        OUTPUT_DIR
        / "validation.json",
        validation_summary,
    )

    # ========================================================
    # CONSOLE SUMMARY
    # ========================================================

    print(
        "\n\n"
        "============================================================"
    )

    print(
        "ALL-FUND EXTRACTION COMPLETE"
    )

    print(
        "============================================================"
    )

    print(
        f"Excel fund universe: "
        f"{len(funds)}"
    )

    print(
        f"Successful funds: "
        f"{len(successful_funds)}"
    )

    print(
        f"Failed funds: "
        f"{len(failed_funds)}"
    )

    print(
        f"Total windows: "
        f"{total_windows}"
    )

    print(
        f"Total pages: "
        f"{total_pages}"
    )

    print(
        f"Total historical BID observations: "
        f"{total_observations}"
    )

    print(
        "\nInception rule:"
    )

    print(
        " - Exact inception date: PASS"
    )

    print(
        " - Up to +7 calendar days after inception: PASS"
    )

    print(
        " - Date gaps: ALLOWED"
    )

    print(
        " - Synthetic data: FORBIDDEN"
    )

    print(
        " - Interpolation: FORBIDDEN"
    )

    print(
        " - Duplicate dates: FAIL"
    )

    if failed_funds:

        print(
            "\nFailed funds:"
        )

        for failure in failed_funds:

            print(
                f" - Row "
                f"{failure['excelRow']}: "
                f"{failure['excelPruAccessName']} "
                f"| "
                f"{failure['error']}"
            )

    print(
        "\nFiles created:"
    )

    for file in sorted(
        OUTPUT_DIR.iterdir()
    ):

        if file.is_file():

            print(
                f" - {file}"
            )

    print(
        "\nIndividual fund results:"
    )

    print(
        f" - {FUNDS_OUTPUT_DIR}"
    )

    # ========================================================
    # PRODUCTION GATE
    # ========================================================

    if failed_funds:

        print(
            "\n============================================================"
        )

        print(
            "VALIDATION FAILED"
        )

        print(
            "============================================================"
        )

        print(
            "One or more Excel-master funds failed."
        )

        print(
            "The run must NOT be treated as a complete "
            "production dataset."
        )

        raise SystemExit(1)

    if len(successful_funds) != len(funds):

        print(
            "\n============================================================"
        )

        print(
            "VALIDATION FAILED"
        )

        print(
            "============================================================"
        )

        print(
            "Successful fund count does not equal "
            "the Excel master-universe count."
        )

        raise SystemExit(1)

    print(
        "\n============================================================"
    )

    print(
        "VALIDATION PASSED"
    )

    print(
        "============================================================"
    )

    print(
        f"{len(successful_funds)} / "
        f"{len(funds)} Excel-master funds passed."
    )

    print(
        "All historical BID windows passed."
    )

    print(
        "All pagination passed."
    )

    print(
        "All inception-coverage checks passed."
    )

    print(
        "No synthetic data was generated."
    )

    print(
        "Date gaps were allowed without filling."
    )

    print(
        "\nDone."
    )


# ============================================================
# ENTRY POINT
# ============================================================

if __name__ == "__main__":

    asyncio.run(
        main()
    )
