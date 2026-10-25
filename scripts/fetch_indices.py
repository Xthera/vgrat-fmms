#!/usr/bin/env python3
"""
VGrat FMS - FETCH INDICES

Reads the index list from "Indices.xlsx" (sheet "Indices") and writes
data/indices.json with each index's daily closing levels and returns.

Indices.xlsx columns (row 1 = headings):

    A  Name            e.g. Straits Times Index
    B  Yahoo Symbol    e.g. ^STI   (required)
    C  Code            short code shown on the site, e.g. STI (required)
    D  Region          e.g. Singapore
    E  Currency        e.g. SGD (optional: taken from Yahoo if blank)
    F  Stooq Symbol    e.g. ^spx (optional fallback source)
    G  Description     optional text shown in the index profile

Source: Yahoo Finance daily closes for each index's full history (as far
back as Yahoo has it, e.g. S&P 500 from 1927), with Stooq as a
fallback when a symbol is given. Both are unofficial public sources.

Failed index: the last good record in data/indices.json is kept and
marked "stale"; an index that has never been fetched is left out.
Either way it is listed in the log and the GitHub Actions run summary.
The script exits with an error only if no index could be published.
"""

from __future__ import annotations

import csv
import io
import json
import os
import sys
import time
from datetime import date, datetime, timedelta, timezone
from pathlib import Path

import requests
from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parent.parent
CONFIG_FILE = ROOT / "Indices.xlsx"
CONFIG_SHEET = "Indices"
OUTPUT_FILE = ROOT / "data" / "indices.json"

HISTORY_START = date(1900, 1, 1)     # full history: everything the source has
MIN_OBSERVATIONS = 50
MAX_STALE_DAYS = 14          # latest close older than this = failed fetch
RETRIES = 3
TIMEOUT = 30

HEADERS = {
    "User-Agent": (
        "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
        "(KHTML, like Gecko) Chrome/124.0 Safari/537.36"
    ),
    "Accept": "application/json,text/csv,*/*",
}


# ---------------------------------------------------------------------------
# helpers
# ---------------------------------------------------------------------------

def clean(value) -> str:
    return " ".join(str(value).split()) if value is not None else ""


def utc_now() -> str:
    return datetime.now(timezone.utc).replace(microsecond=0).isoformat()


def iso(value: date) -> str:
    return value.isoformat()


def months_back(day: date, months: int) -> date:
    year = day.year
    month = day.month - months

    while month <= 0:
        month += 12
        year -= 1

    for last in (31, 30, 29, 28):
        try:
            return date(year, month, min(day.day, last))
        except ValueError:
            continue

    raise ValueError("invalid date")


def read_config() -> list[dict]:
    if not CONFIG_FILE.exists():
        raise FileNotFoundError(f"{CONFIG_FILE.name} not found")

    workbook = load_workbook(CONFIG_FILE, read_only=True, data_only=True)

    try:
        sheet = (
            workbook[CONFIG_SHEET]
            if CONFIG_SHEET in workbook.sheetnames
            else workbook[workbook.sheetnames[0]]
        )

        rows = []

        for number, row in enumerate(
            sheet.iter_rows(min_row=2, max_col=7, values_only=True), start=2
        ):
            values = list(row) + [None] * (7 - len(row))
            name, symbol, code, region, currency, stooq, description = (
                clean(v) for v in values
            )

            if not symbol or not code:
                continue

            rows.append({
                "row": number,
                "name": name or code,
                "symbol": symbol,
                "code": code.upper(),
                "region": region,
                "currency": currency.upper(),
                "stooq": stooq,
                "description": description,
            })
    finally:
        workbook.close()

    codes = [row["code"] for row in rows]
    duplicates = {code for code in codes if codes.count(code) > 1}

    if duplicates:
        raise ValueError(f"Duplicate codes in {CONFIG_FILE.name}: {sorted(duplicates)}")

    return rows


def get(url: str, params: dict | None = None) -> requests.Response:
    last_error = None

    for attempt in range(1, RETRIES + 1):
        try:
            response = requests.get(url, params=params, headers=HEADERS, timeout=TIMEOUT)

            if response.status_code == 200:
                return response

            last_error = RuntimeError(f"HTTP {response.status_code}")
        except requests.RequestException as error:
            last_error = error

        time.sleep(3 * attempt)

    raise RuntimeError(f"{url}: {last_error}")


# ---------------------------------------------------------------------------
# sources
# ---------------------------------------------------------------------------

def fetch_yahoo(symbol: str) -> tuple[list[tuple[str, float]], dict]:
    start = int(datetime(HISTORY_START.year, HISTORY_START.month, HISTORY_START.day, tzinfo=timezone.utc).timestamp())
    end = int(time.time()) + 86400

    response = get(
        f"https://query1.finance.yahoo.com/v8/finance/chart/{requests.utils.quote(symbol, safe='')}",
        {"period1": start, "period2": end, "interval": "1d", "events": "history"},
    )

    chart = response.json().get("chart") or {}
    result = (chart.get("result") or [None])[0]

    if not result:
        raise RuntimeError(f"Yahoo returned no data ({chart.get('error')})")

    meta = result.get("meta") or {}
    offset = int(meta.get("gmtoffset") or 0)
    stamps = result.get("timestamp") or []
    closes = ((result.get("indicators") or {}).get("quote") or [{}])[0].get("close") or []

    by_date: dict[str, float] = {}

    for stamp, close in zip(stamps, closes):
        if close is None:
            continue

        day = datetime.fromtimestamp(stamp + offset, tz=timezone.utc).date()
        by_date[iso(day)] = round(float(close), 2)

    observations = sorted(by_date.items())

    return observations, {
        "currency": clean(meta.get("currency")).upper(),
        "exchange": clean(meta.get("fullExchangeName") or meta.get("exchangeName")),
        "sourceName": clean(meta.get("longName") or meta.get("shortName")),
        "source": "Yahoo Finance",
    }


def fetch_stooq(symbol: str) -> tuple[list[tuple[str, float]], dict]:
    response = get("https://stooq.com/q/d/l/", {"s": symbol.lower(), "i": "d"})
    text = response.text.strip()

    if not text or "Date" not in text.splitlines()[0]:
        raise RuntimeError("Stooq returned no CSV data")

    observations = []

    for row in csv.DictReader(io.StringIO(text)):
        day = clean(row.get("Date"))
        close = clean(row.get("Close"))

        if not day or not close or day < iso(HISTORY_START):
            continue

        observations.append((day, round(float(close), 2)))

    return sorted(observations), {"source": "Stooq"}


# ---------------------------------------------------------------------------
# returns
# ---------------------------------------------------------------------------

def close_on_or_before(observations: list[tuple[str, float]], day: str) -> tuple[str, float] | None:
    low, high, found = 0, len(observations) - 1, None

    while low <= high:
        middle = (low + high) // 2

        if observations[middle][0] <= day:
            found = observations[middle]
            low = middle + 1
        else:
            high = middle - 1

    return found


def returns_for(observations: list[tuple[str, float]]) -> dict:
    latest_day, latest = observations[-1]
    latest_date = date.fromisoformat(latest_day)
    first_day, first = observations[0]

    def cumulative(target: date):
        point = close_on_or_before(observations, iso(target))

        if not point or point[0] < first_day or point[0] >= latest_day or point[1] <= 0:
            return None

        # Need history reaching back to the target date
        if iso(target) < first_day:
            return None

        return round((latest / point[1] - 1) * 100, 2)

    def annualised(years: int):
        target = months_back(latest_date, 12 * years)
        point = close_on_or_before(observations, iso(target))

        if not point or iso(target) < first_day or point[1] <= 0:
            return None

        return round(((latest / point[1]) ** (1 / years) - 1) * 100, 2)

    span = (latest_date - date.fromisoformat(first_day)).days / 365.25

    return {
        "ytd": cumulative(date(latest_date.year - 1, 12, 31)),
        "m1": cumulative(months_back(latest_date, 1)),
        "m3": cumulative(months_back(latest_date, 3)),
        "m6": cumulative(months_back(latest_date, 6)),
        "y1": cumulative(months_back(latest_date, 12)),
        "y3": cumulative(months_back(latest_date, 36)),
        "y5": cumulative(months_back(latest_date, 60)),
        "a3": annualised(3),
        "a5": annualised(5),
        "a10": annualised(10),
        "sinceStart": round(((latest / first) ** (1 / span) - 1) * 100, 2) if span >= 1 and first > 0 else None,
    }


# ---------------------------------------------------------------------------
# main
# ---------------------------------------------------------------------------

def load_previous() -> dict:
    try:
        data = json.loads(OUTPUT_FILE.read_text(encoding="utf-8"))
        return {item["code"]: item for item in data.get("indices", []) if isinstance(item, dict) and item.get("code")}
    except Exception:
        return {}


def fetch_one(config: dict) -> dict:
    errors = []
    observations, meta = [], {}

    for name, fetcher, symbol in (
        ("Yahoo", fetch_yahoo, config["symbol"]),
        ("Stooq", fetch_stooq, config["stooq"]),
    ):
        if not symbol:
            continue

        try:
            observations, meta = fetcher(symbol)
        except Exception as error:
            errors.append(f"{name}: {error}")
            continue

        if len(observations) < MIN_OBSERVATIONS:
            errors.append(f"{name}: only {len(observations)} prices")
            observations = []
            continue

        age = (date.today() - date.fromisoformat(observations[-1][0])).days

        if age > MAX_STALE_DAYS:
            errors.append(f"{name}: latest price is {observations[-1][0]} ({age} days old)")
            observations = []
            continue

        break

    if not observations:
        raise RuntimeError("; ".join(errors) or "no source available")

    return {
        "id": f"IDX-{config['code']}",
        "code": config["code"],
        "name": config["name"],
        "symbol": config["symbol"],
        "region": config["region"],
        "currency": config["currency"] or meta.get("currency") or "",
        "description": config["description"],
        "exchange": meta.get("exchange", ""),
        "source": meta.get("source", ""),
        "firstDate": observations[0][0],
        "latestDate": observations[-1][0],
        "latestClose": observations[-1][1],
        "returns": returns_for(observations),
        "observationCount": len(observations),
        "observations": [[day, close] for day, close in observations],
    }


def write_summary(lines: list[str]) -> None:
    path = os.environ.get("GITHUB_STEP_SUMMARY")

    if path:
        try:
            with open(path, "a", encoding="utf-8") as handle:
                handle.write("\n".join(lines) + "\n")
        except OSError as error:
            print(f"Could not write the run summary: {error}")


def main() -> int:
    generated = utc_now()
    config = read_config()
    previous = load_previous()

    published, report = [], []

    for item in config:
        print(f"{item['code']:<8} {item['symbol']:<12} ", end="", flush=True)

        try:
            record = fetch_one(item)
            record["updateStatus"] = {"status": "updated", "lastSuccessfulUpdateUtc": generated}
            published.append(record)
            print(f"ok   {record['observationCount']} prices to {record['latestDate']} ({record['source']})")
        except Exception as error:
            old = previous.get(item["code"])

            if old:
                last_good = (old.get("updateStatus") or {}).get("lastSuccessfulUpdateUtc")
                old = {**old, "name": item["name"], "region": item["region"], "description": item["description"]}
                old["updateStatus"] = {
                    "status": "stale",
                    "lastSuccessfulUpdateUtc": last_good,
                    "failedUpdateUtc": generated,
                    "reasons": [str(error)],
                }
                published.append(old)
                action = "Kept last good data"
            else:
                action = "Left out (never fetched)"

            report.append({"code": item["code"], "name": item["name"], "reason": str(error)[:240], "action": action})
            print(f"FAILED {error} -> {action}")

        time.sleep(1)

    lines = [
        "## Indices",
        "",
        f"- Indices in {CONFIG_FILE.name}: **{len(config)}**",
        f"- Updated: **{len(published) - sum(1 for r in report if r['action'].startswith('Kept'))}**",
        f"- With a failure: **{len(report)}**",
        "",
    ]

    if report:
        lines += ["| Code | Index | Reason | What happened |", "| --- | --- | --- | --- |"]
        lines += [
            f"| {r['code']} | {r['name']} | {r['reason'].replace('|', '/')} | {r['action']} |"
            for r in report
        ]
    else:
        lines.append("All indices updated.")

    write_summary(lines)

    if not published:
        print("No index could be published; data/indices.json is left unchanged.")
        return 1

    output = {
        "schemaVersion": 1,
        "generatedAtUtc": generated,
        "source": "Yahoo Finance (Stooq fallback)",
        "sourceNote": "Daily closing levels from public sources, for information only.",
        "indexCount": len(published),
        "summary": {"configured": len(config), "failed": report},
        "indices": published,
    }

    OUTPUT_FILE.parent.mkdir(parents=True, exist_ok=True)
    temporary = OUTPUT_FILE.with_suffix(".tmp")
    temporary.write_text(json.dumps(output, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
    temporary.replace(OUTPUT_FILE)

    print(f"Wrote {OUTPUT_FILE.relative_to(ROOT)} ({len(published)} indices)")
    return 0


if __name__ == "__main__":
    sys.exit(main())
