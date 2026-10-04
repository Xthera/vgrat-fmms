#!/usr/bin/env python3

"""
VGrat FMS - MARKET NEWS AI ANALYZER
===================================

Purpose
-------
Analyze pending market-news articles using Gemini.

IMPORTANT DATA RULE
-------------------
Research Funds.xlsx is the authoritative vocabulary for:

    Geography
    Sector

Gemini may initially return a value that does not exactly match the
Excel master. In that case, the script enters CORRECTION MODE and asks
Gemini to remap invalid values to the closest valid Excel values.

Python performs the final validation.

Nothing is written to analysis/current.json unless:

    1. Gemini analysis succeeds
    2. Every article has exactly one analysis
    3. Every article ID matches the requested batch
    4. Every geography exactly exists in the Excel Geography Master
    5. Every sector exactly exists in the Excel Sector Master
    6. Required fields are present
    7. Corrected output passes the same validation

INPUTS
------
    data/market_news/current.json
    Research Funds.xlsx

OUTPUT
------
    data/market_news/analysis/current.json

The production input files are READ ONLY.

MODEL
-----
    Gemini

DEFAULT MODEL:
    gemini-3.8-flash

BATCH
-----
    BATCH_SIZE = 2

This is intentionally kept at 2 because the 5-article batch previously
experienced HTTP 503 responses while 2-article batches succeeded.

PROCESSING
----------
    Newest pending articles first.

CORRECTION MODE
---------------
    Initial Gemini response is validated.

    If Geography or Sector contains invalid values:

        Initial response
              ↓
        Identify invalid values
              ↓
        Gemini correction request
              ↓
        Validate corrected response
              ↓
        Accept only if everything is valid

The original response and correction response are both stored.

FAILURE SAFETY
--------------
    If analysis or correction fails validation:

        analysis/current.json is NOT modified.

Only a completely successful batch is appended and written atomically.
"""

from __future__ import annotations

import hashlib
import json
import os
import sys
import time
import urllib.error
import urllib.request
from datetime import datetime, timezone
from pathlib import Path
from typing import Any


# ============================================================
# PATHS
# ============================================================

ROOT = Path(__file__).resolve().parent.parent

CURRENT_JSON = ROOT / "data" / "market_news" / "current.json"
RESEARCH_FUNDS = ROOT / "Research Funds.xlsx"

ANALYSIS_DIR = ROOT / "data" / "market_news" / "analysis"
ANALYSIS_CURRENT = ANALYSIS_DIR / "current.json"


# ============================================================
# CONFIGURATION
# ============================================================

BATCH_SIZE = 2
MAX_BATCHES = 1
PROCESS_NEWEST_FIRST = True

DEFAULT_MODEL = "gemini-3.8-flash"
DEFAULT_TIMEOUT = 180
DEFAULT_MAX_OUTPUT_TOKENS = 16384

GEMINI_API_BASE = (
    "https://generativelanguage.googleapis.com/v1beta/models"
)

# Maximum number of correction attempts for one batch.
MAX_CORRECTION_ATTEMPTS = 2


# ============================================================
# LOGGING
# ============================================================

def log(message: str) -> None:
    timestamp = datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M:%S")
    print(f"[{timestamp} UTC] {message}", flush=True)


# ============================================================
# TIME
# ============================================================

def utc_now_iso() -> str:
    return datetime.now(timezone.utc).isoformat()


# ============================================================
# JSON
# ============================================================

def load_json(path: Path) -> Any:
    with path.open("r", encoding="utf-8") as f:
        return json.load(f)


def save_json_atomic(path: Path, data: Any) -> None:
    """
    Atomically replace a JSON file.

    The existing file is untouched if serialization or writing fails.
    """

    path.parent.mkdir(parents=True, exist_ok=True)

    temp_path = path.with_suffix(path.suffix + ".tmp")

    with temp_path.open("w", encoding="utf-8") as f:
        json.dump(
            data,
            f,
            ensure_ascii=False,
            indent=2,
        )
        f.write("\n")

    temp_path.replace(path)


# ============================================================
# ARTICLE EXTRACTION
# ============================================================

def extract_articles(data: Any) -> list[dict[str, Any]]:
    """
    Support several possible wrappers used by the market-news collector.
    """

    if isinstance(data, list):
        return [x for x in data if isinstance(x, dict)]

    if not isinstance(data, dict):
        raise ValueError("current.json must contain a JSON object or array.")

    for key in (
        "articles",
        "news",
        "items",
        "data",
        "results",
    ):
        value = data.get(key)

        if isinstance(value, list):
            return [x for x in value if isinstance(x, dict)]

    raise ValueError(
        "Could not find article list in current.json. "
        "Expected one of: articles, news, items, data, results."
    )


# ============================================================
# ARTICLE ID
# ============================================================

def article_id(article: dict[str, Any]) -> str:
    """
    Prefer the collector's existing ID.

    Fall back to URL hash, then canonical article hash.
    """

    for key in (
        "id",
        "articleId",
        "article_id",
        "newsId",
        "news_id",
    ):
        value = article.get(key)

        if value is not None and str(value).strip():
            return str(value).strip()

    url = article.get("url")

    if url:
        return hashlib.sha256(
            str(url).encode("utf-8")
        ).hexdigest()

    canonical = json.dumps(
        article,
        ensure_ascii=False,
        sort_keys=True,
        separators=(",", ":"),
    )

    return hashlib.sha256(
        canonical.encode("utf-8")
    ).hexdigest()


# ============================================================
# ARTICLE NORMALIZATION
# ============================================================

def normalize_article(article: dict[str, Any]) -> dict[str, Any]:
    return {
        "articleId": article_id(article),
        "title": article.get("title"),
        "description": article.get("description"),
        "url": article.get("url"),
        "published": article.get("published"),
        "source": article.get("source"),
    }


# ============================================================
# DATE SORTING
# ============================================================

def publication_sort_key(article: dict[str, Any]) -> tuple[int, str]:
    value = article.get("published")

    if value is None:
        return (0, "")

    text_value = str(value).strip()

    if not text_value:
        return (0, "")

    return (1, text_value)


# ============================================================
# RESEARCH FUNDS MASTER LOADING
# ============================================================

def clean_excel_value(value: Any) -> str | None:
    if value is None:
        return None

    text_value = str(value).strip()

    if not text_value:
        return None

    return text_value


def load_master_values(ws) -> list[str]:
    """
    Extract all non-empty cell values from a worksheet.

    Duplicate values are removed while preserving their original order.
    """

    values: list[str] = []
    seen: set[str] = set()

    for row in ws.iter_rows():
        for cell in row:
            value = clean_excel_value(cell.value)

            if value is None:
                continue

            if value not in seen:
                seen.add(value)
                values.append(value)

    return values


def load_research_masters() -> tuple[list[str], list[str]]:
    """
    Research Funds.xlsx structure:

        Sheet 0:
            Fund Research

        Sheet 1:
            Geography master

        Sheet 2:
            Sector master
    """

    if not RESEARCH_FUNDS.exists():
        raise FileNotFoundError(
            f"Research Funds.xlsx not found: {RESEARCH_FUNDS}"
        )

    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError(
            "openpyxl is required. Install it with: pip install openpyxl"
        ) from exc

    log(f"Loading Excel masters: {RESEARCH_FUNDS}")

    workbook = load_workbook(
        RESEARCH_FUNDS,
        read_only=True,
        data_only=True,
    )

    try:
        worksheets = workbook.worksheets

        if len(worksheets) < 3:
            raise ValueError(
                "Research Funds.xlsx must contain at least 3 worksheets: "
                "Fund Research, Geography master, Sector master."
            )

        geography_values = load_master_values(worksheets[1])
        sector_values = load_master_values(worksheets[2])

    finally:
        workbook.close()

    if not geography_values:
        raise ValueError(
            "Geography master in Research Funds.xlsx is empty."
        )

    if not sector_values:
        raise ValueError(
            "Sector master in Research Funds.xlsx is empty."
        )

    log(
        f"Geography master values: {len(geography_values)}"
    )

    log(
        f"Sector master values: {len(sector_values)}"
    )

    return geography_values, sector_values


# ============================================================
# ANALYSIS FILE
# ============================================================

def new_analysis_document() -> dict[str, Any]:
    now = utc_now_iso()

    return {
        "version": 1,
        "createdAt": now,
        "updatedAt": now,
        "totalSuccessfulBatches": 0,
        "totalSuccessfulArticles": 0,
        "batches": [],
    }


def load_existing_analysis() -> dict[str, Any]:
    if not ANALYSIS_CURRENT.exists():
        log("analysis/current.json does not exist. A new file will be created.")
        return new_analysis_document()

    data = load_json(ANALYSIS_CURRENT)

    if not isinstance(data, dict):
        raise ValueError(
            "analysis/current.json must contain a JSON object."
        )

    if "batches" not in data or not isinstance(data["batches"], list):
        data["batches"] = []

    data.setdefault("version", 1)
    data.setdefault("createdAt", utc_now_iso())
    data.setdefault("updatedAt", utc_now_iso())
    data.setdefault("totalSuccessfulBatches", len(data["batches"]))
    data.setdefault("totalSuccessfulArticles", 0)

    return data


# ============================================================
# COMPLETED ARTICLE IDS
# ============================================================

def completed_article_ids(
    analysis: dict[str, Any],
) -> set[str]:

    completed: set[str] = set()

    for batch in analysis.get("batches", []):
        if not isinstance(batch, dict):
            continue

        for value in batch.get("articleIds", []):
            if value is not None:
                completed.add(str(value))

        for item in batch.get("analyses", []):
            if not isinstance(item, dict):
                continue

            value = item.get("articleId")

            if value is not None:
                completed.add(str(value))

    return completed


# ============================================================
# GEMINI API
# ============================================================

def gemini_model() -> str:
    return os.environ.get(
        "GEMINI_MODEL",
        DEFAULT_MODEL,
    )


def gemini_timeout() -> int:
    raw = os.environ.get(
        "GEMINI_TIMEOUT_SECONDS",
        str(DEFAULT_TIMEOUT),
    )

    try:
        return int(raw)
    except ValueError:
        return DEFAULT_TIMEOUT


def gemini_max_output_tokens() -> int:
    raw = os.environ.get(
        "GEMINI_MAX_OUTPUT_TOKENS",
        str(DEFAULT_MAX_OUTPUT_TOKENS),
    )

    try:
        return int(raw)
    except ValueError:
        return DEFAULT_MAX_OUTPUT_TOKENS


def call_gemini(
    *,
    model: str,
    api_key: str,
    prompt: str,
) -> dict[str, Any]:

    url = (
        f"{GEMINI_API_BASE}/"
        f"{model}:generateContent"
        f"?key={api_key}"
    )

    request_body = {
        "contents": [
            {
                "parts": [
                    {
                        "text": prompt,
                    }
                ]
            }
        ],
        "generationConfig": {
            "temperature": 0,
            "maxOutputTokens": gemini_max_output_tokens(),
            "responseMimeType": "application/json",
        },
    }

    encoded = json.dumps(
        request_body,
        ensure_ascii=False,
    ).encode("utf-8")

    request = urllib.request.Request(
        url,
        data=encoded,
        headers={
            "Content-Type": "application/json",
        },
        method="POST",
    )

    started = time.monotonic()

    try:
        with urllib.request.urlopen(
            request,
            timeout=gemini_timeout(),
        ) as response:

            response_seconds = time.monotonic() - started

            status = response.status

            raw_bytes = response.read()

            raw_text = raw_bytes.decode(
                "utf-8",
                errors="replace",
            )

            parsed = json.loads(raw_text)

            return {
                "ok": True,
                "httpStatus": status,
                "responseSeconds": round(
                    response_seconds,
                    3,
                ),
                "raw": parsed,
                "requestBody": request_body,
            }

    except urllib.error.HTTPError as exc:

        response_seconds = time.monotonic() - started

        try:
            error_body = exc.read().decode(
                "utf-8",
                errors="replace",
            )
        except Exception:
            error_body = ""

        try:
            error_json = json.loads(error_body)
        except Exception:
            error_json = error_body

        return {
            "ok": False,
            "httpStatus": exc.code,
            "responseSeconds": round(
                response_seconds,
                3,
            ),
            "error": error_json,
            "requestBody": request_body,
        }

    except Exception as exc:

        response_seconds = time.monotonic() - started

        return {
            "ok": False,
            "httpStatus": None,
            "responseSeconds": round(
                response_seconds,
                3,
            ),
            "error": str(exc),
            "requestBody": request_body,
        }


# ============================================================
# GEMINI RESPONSE EXTRACTION
# ============================================================

def extract_response_text(response: dict[str, Any]) -> str:
    raw = response.get("raw")

    if not isinstance(raw, dict):
        raise ValueError("Gemini response is not a JSON object.")

    candidates = raw.get("candidates")

    if not isinstance(candidates, list) or not candidates:
        raise ValueError(
            "Gemini response contains no candidates."
        )

    candidate = candidates[0]

    if not isinstance(candidate, dict):
        raise ValueError("Gemini candidate is invalid.")

    content = candidate.get("content")

    if not isinstance(content, dict):
        raise ValueError(
            "Gemini candidate contains no content."
        )

    parts = content.get("parts")

    if not isinstance(parts, list):
        raise ValueError(
            "Gemini content contains no parts."
        )

    texts: list[str] = []

    for part in parts:
        if not isinstance(part, dict):
            continue

        text_value = part.get("text")

        if isinstance(text_value, str):
            texts.append(text_value)

    if not texts:
        raise ValueError(
            "Gemini response contains no text."
        )

    return "".join(texts).strip()


def parse_generated_json(text_value: str) -> Any:
    """
    Parse JSON.

    Also handles accidental markdown fences.
    """

    cleaned = text_value.strip()

    if cleaned.startswith("```"):
        lines = cleaned.splitlines()

        if lines and lines[0].startswith("```"):
            lines = lines[1:]

        if lines and lines[-1].strip() == "```":
            lines = lines[:-1]

        cleaned = "\n".join(lines).strip()

        if cleaned.lower().startswith("json"):
            cleaned = cleaned[4:].strip()

    return json.loads(cleaned)


def normalize_analysis_list(value: Any) -> list[dict[str, Any]]:
    """
    Accept:

        [...]
        {"analyses": [...]}
        {"results": [...]}
        {single analysis object}
    """

    if isinstance(value, list):
        return [
            item
            for item in value
            if isinstance(item, dict)
        ]

    if isinstance(value, dict):

        for key in (
            "analyses",
            "results",
        ):
            nested = value.get(key)

            if isinstance(nested, list):
                return [
                    item
                    for item in nested
                    if isinstance(item, dict)
                ]

        if value.get("articleId") is not None:
            return [value]

    raise ValueError(
        "Generated JSON does not contain a valid analysis list."
    )


# ============================================================
# PROMPT HELPERS
# ============================================================

def numbered_values(values: list[str]) -> str:
    return "\n".join(
        f"{index}. {value}"
        for index, value in enumerate(values, start=1)
    )


def article_prompt_json(
    articles: list[dict[str, Any]],
) -> str:

    return json.dumps(
        articles,
        ensure_ascii=False,
        indent=2,
    )


# ============================================================
# INITIAL ANALYSIS PROMPT
# ============================================================

def build_analysis_prompt(
    articles: list[dict[str, Any]],
    geography_master: list[str],
    sector_master: list[str],
) -> str:

    geography_text = numbered_values(
        geography_master
    )

    sector_text = numbered_values(
        sector_master
    )

    articles_text = article_prompt_json(
        articles
    )

    return f"""
You are the market-news intelligence analyst for VGrat FMS.

Analyze EVERY article supplied below.

IMPORTANT
=========
You must return exactly ONE analysis object for every article.

Do not omit articles.

Do not create articles.

Do not merge articles.

Do not duplicate articles.

ARTICLE ID
==========
The articleId supplied in the input is authoritative.

Return exactly the same articleId.

GEOGRAPHY MASTER
================
The following list comes directly from Research Funds.xlsx.

These are the ONLY valid Geography values.

You MUST select geography values EXACTLY as written below.

Do not create synonyms.

Do not abbreviate.

Do not invent alternative names.

Do not use "US" if the master contains "United States".

Do not use "USA" if the master contains "United States".

Do not use "Americas" unless it is actually present below.

VALID GEOGRAPHY VALUES:

{geography_text}

SECTOR MASTER
=============
The following list comes directly from Research Funds.xlsx.

These are the ONLY valid Sector values.

You MUST select sector values EXACTLY as written below.

Do not create synonyms.

Do not abbreviate.

Do not invent alternative names.

For example, if the master contains "Information Technology",
do not return "Technology" unless "Technology" itself appears in
the master.

VALID SECTOR VALUES:

{sector_text}

ASSET CLASSES
=============
Asset classes are NOT restricted by the Geography or Sector masters.

Use sensible asset-class descriptions.

OUTPUT SCHEMA
=============
Return JSON only.

Return a JSON array.

Each object MUST contain exactly these fields:

{{
  "articleId": "string",
  "relevant": true,
  "category": "MARKET",
  "sentiment": "positive",
  "importance": "medium",
  "assetClasses": ["Equities"],
  "geographies": ["United States"],
  "sectors": ["Information Technology"],
  "investorImpact": "string",
  "reasoning": "string",
  "fundMonitoringRelevant": true
}}

FIELD RULES
===========
relevant:
    boolean

category:
    one of:
    MARKET
    ECONOMIC
    TECHNOLOGY
    GEOPOLITICAL
    REJECT

sentiment:
    positive
    negative
    mixed
    neutral

importance:
    low
    medium
    high

assetClasses:
    JSON array of strings

geographies:
    JSON array of strings
    EVERY VALUE MUST EXACTLY MATCH THE GEOGRAPHY MASTER

sectors:
    JSON array of strings
    EVERY VALUE MUST EXACTLY MATCH THE SECTOR MASTER

investorImpact:
    concise explanation of the investment implications

reasoning:
    concise explanation of why the article was classified this way

fundMonitoringRelevant:
    boolean

IMPORTANT FINAL CHECK
=====================
Before returning your answer, compare every geography and sector
against the supplied master lists.

Every value must be an EXACT string match.

Return JSON only.

ARTICLES
========

{articles_text}
""".strip()


# ============================================================
# INITIAL VALIDATION
# ============================================================

REQUIRED_ANALYSIS_FIELDS = {
    "articleId",
    "relevant",
    "category",
    "sentiment",
    "importance",
    "assetClasses",
    "geographies",
    "sectors",
    "investorImpact",
    "reasoning",
    "fundMonitoringRelevant",
}


def validate_analysis(
    analyses: list[dict[str, Any]],
    articles: list[dict[str, Any]],
    geography_master: set[str],
    sector_master: set[str],
) -> dict[str, Any]:

    expected_ids = [
        str(article["articleId"])
        for article in articles
    ]

    expected_set = set(expected_ids)

    received_ids = [
        str(item.get("articleId"))
        for item in analyses
    ]

    received_set = set(received_ids)

    errors: list[str] = []

    # --------------------------------------------------------
    # Count
    # --------------------------------------------------------

    if len(analyses) != len(articles):
        errors.append(
            f"Expected {len(articles)} analyses, "
            f"received {len(analyses)}."
        )

    # --------------------------------------------------------
    # Duplicate IDs
    # --------------------------------------------------------

    if len(received_ids) != len(received_set):
        errors.append(
            "Duplicate article IDs detected."
        )

    # --------------------------------------------------------
    # Exact article coverage
    # --------------------------------------------------------

    missing_ids = sorted(
        expected_set - received_set
    )

    unexpected_ids = sorted(
        received_set - expected_set
    )

    if missing_ids:
        errors.append(
            "Missing article IDs: "
            + ", ".join(missing_ids)
        )

    if unexpected_ids:
        errors.append(
            "Unexpected article IDs: "
            + ", ".join(unexpected_ids)
        )

    invalid_geographies: dict[str, list[str]] = {}
    invalid_sectors: dict[str, list[str]] = {}

    # --------------------------------------------------------
    # Per-analysis validation
    # --------------------------------------------------------

    for item in analyses:

        current_id = str(
            item.get("articleId")
        )

        missing_fields = sorted(
            REQUIRED_ANALYSIS_FIELDS
            - set(item.keys())
        )

        if missing_fields:
            errors.append(
                f"{current_id}: missing fields: "
                + ", ".join(missing_fields)
            )

        # ----------------------------------------------------
        # Basic types
        # ----------------------------------------------------

        if not isinstance(
            item.get("relevant"),
            bool,
        ):
            errors.append(
                f"{current_id}: relevant must be boolean."
            )

        if not isinstance(
            item.get("fundMonitoringRelevant"),
            bool,
        ):
            errors.append(
                f"{current_id}: fundMonitoringRelevant "
                "must be boolean."
            )

        # ----------------------------------------------------
        # Arrays
        # ----------------------------------------------------

        for field in (
            "assetClasses",
            "geographies",
            "sectors",
        ):

            value = item.get(field)

            if not isinstance(value, list):
                errors.append(
                    f"{current_id}: {field} must be a JSON array."
                )
                continue

            if not all(
                isinstance(x, str)
                for x in value
            ):
                errors.append(
                    f"{current_id}: {field} must contain "
                    "only strings."
                )

        # ----------------------------------------------------
        # Geography master validation
        # ----------------------------------------------------

        geographies = item.get(
            "geographies"
        )

        if isinstance(geographies, list):

            invalid = [
                str(value)
                for value in geographies
                if str(value) not in geography_master
            ]

            if invalid:
                invalid_geographies[
                    current_id
                ] = invalid

                errors.append(
                    f"{current_id}: invalid geography values: "
                    + ", ".join(invalid)
                )

        # ----------------------------------------------------
        # Sector master validation
        # ----------------------------------------------------

        sectors = item.get("sectors")

        if isinstance(sectors, list):

            invalid = [
                str(value)
                for value in sectors
                if str(value) not in sector_master
            ]

            if invalid:
                invalid_sectors[
                    current_id
                ] = invalid

                errors.append(
                    f"{current_id}: invalid sector values: "
                    + ", ".join(invalid)
                )

    return {
        "passed": not errors,
        "errors": errors,
        "expectedArticleCount": len(articles),
        "receivedArticleCount": len(analyses),
        "expectedArticleIds": expected_ids,
        "receivedArticleIds": received_ids,
        "invalidGeographies": invalid_geographies,
        "invalidSectors": invalid_sectors,
    }


# ============================================================
# CORRECTION PROMPT
# ============================================================

def build_correction_prompt(
    *,
    articles: list[dict[str, Any]],
    analyses: list[dict[str, Any]],
    validation: dict[str, Any],
    geography_master: list[str],
    sector_master: list[str],
) -> str:

    geography_text = numbered_values(
        geography_master
    )

    sector_text = numbered_values(
        sector_master
    )

    invalid_geographies = (
        validation.get("invalidGeographies", {})
    )

    invalid_sectors = (
        validation.get("invalidSectors", {})
    )

    invalid_summary = {
        "invalidGeographies": invalid_geographies,
        "invalidSectors": invalid_sectors,
    }

    return f"""
You are correcting a market-news classification produced by another
AI analyst.

The original analysis is mostly complete, but some Geography and/or
Sector values are INVALID.

Your job is to correct ONLY the invalid classification values while
preserving the rest of the analysis.

AUTHORITATIVE RULE
==================
Research Funds.xlsx is the source of truth.

A Geography is valid ONLY if it exactly matches one of the values in
the Geography Master below.

A Sector is valid ONLY if it exactly matches one of the values in the
Sector Master below.

Do NOT invent values.

Do NOT use synonyms.

Do NOT abbreviate.

Do NOT return values that are not present in the masters.

GEOGRAPHY MASTER
================

{geography_text}

SECTOR MASTER
=============

{sector_text}

INVALID VALUES FOUND
====================

{json.dumps(
    invalid_summary,
    ensure_ascii=False,
    indent=2,
)}

ORIGINAL ARTICLES
=================

{json.dumps(
    articles,
    ensure_ascii=False,
    indent=2,
)}

ORIGINAL ANALYSES
=================

{json.dumps(
    analyses,
    ensure_ascii=False,
    indent=2,
)}

CORRECTION RULES
================

1. Return exactly one analysis per article.

2. Preserve every articleId exactly.

3. Preserve the meaning of the original analysis.

4. Correct invalid Geography values by selecting the most appropriate
   EXACT value from the Geography Master.

5. Correct invalid Sector values by selecting the most appropriate
   EXACT value from the Sector Master.

6. Do not modify valid Geography values unless necessary to make the
   classification internally consistent.

7. Do not modify valid Sector values unless necessary to make the
   classification internally consistent.

8. Do not modify unrelated fields.

9. If a value has no valid match, select the closest applicable
   authoritative master value.

10. Never return a value outside the supplied masters.

OUTPUT
======

Return JSON only.

Return a JSON array using this schema:

{{
  "articleId": "string",
  "relevant": true,
  "category": "MARKET",
  "sentiment": "positive",
  "importance": "medium",
  "assetClasses": ["Equities"],
  "geographies": ["United States"],
  "sectors": ["Information Technology"],
  "investorImpact": "string",
  "reasoning": "string",
  "fundMonitoringRelevant": true
}}
""".strip()


# ============================================================
# BATCH ANALYSIS
# ============================================================

def analyze_batch(
    *,
    articles: list[dict[str, Any]],
    geography_master: list[str],
    sector_master: list[str],
    api_key: str,
    model: str,
) -> dict[str, Any]:

    log(
        f"Starting Gemini analysis for "
        f"{len(articles)} article(s)."
    )

    prompt = build_analysis_prompt(
        articles,
        geography_master,
        sector_master,
    )

    log(
        f"Initial prompt characters: {len(prompt)}"
    )

    initial_response = call_gemini(
        model=model,
        api_key=api_key,
        prompt=prompt,
    )

    if not initial_response["ok"]:
        return {
            "success": False,
            "stage": "initial_analysis",
            "prompt": prompt,
            "response": initial_response,
            "error": (
                "Initial Gemini request failed."
            ),
        }

    try:
        raw_text = extract_response_text(
            initial_response
        )

        parsed = parse_generated_json(
            raw_text
        )

        analyses = normalize_analysis_list(
            parsed
        )

    except Exception as exc:

        return {
            "success": False,
            "stage": "initial_parse",
            "prompt": prompt,
            "response": initial_response,
            "error": str(exc),
        }

    validation = validate_analysis(
        analyses,
        articles,
        set(geography_master),
        set(sector_master),
    )

    # --------------------------------------------------------
    # Initial validation passed
    # --------------------------------------------------------

    if validation["passed"]:

        log(
            "Initial analysis passed validation. "
            "No correction required."
        )

        return {
            "success": True,
            "mode": "initial",
            "prompt": prompt,
            "response": initial_response,
            "rawText": raw_text,
            "parsedResponse": parsed,
            "analyses": analyses,
            "validation": validation,
            "correctionAttempts": [],
        }

    # --------------------------------------------------------
    # Correction required
    # --------------------------------------------------------

    log(
        "Initial analysis failed validation."
    )

    log(
        "Invalid Geography values: "
        + json.dumps(
            validation.get(
                "invalidGeographies",
                {},
            ),
            ensure_ascii=False,
        )
    )

    log(
        "Invalid Sector values: "
        + json.dumps(
            validation.get(
                "invalidSectors",
                {},
            ),
            ensure_ascii=False,
        )
    )

    correction_attempts: list[dict[str, Any]] = []

    current_analyses = analyses
    current_validation = validation

    for attempt_number in range(
        1,
        MAX_CORRECTION_ATTEMPTS + 1,
    ):

        log(
            f"Starting correction attempt "
            f"{attempt_number}/{MAX_CORRECTION_ATTEMPTS}."
        )

        correction_prompt = build_correction_prompt(
            articles=articles,
            analyses=current_analyses,
            validation=current_validation,
            geography_master=geography_master,
            sector_master=sector_master,
        )

        correction_response = call_gemini(
            model=model,
            api_key=api_key,
            prompt=correction_prompt,
        )

        correction_record: dict[str, Any] = {
            "attempt": attempt_number,
            "prompt": correction_prompt,
            "response": correction_response,
        }

        if not correction_response["ok"]:

            correction_record["success"] = False
            correction_record["error"] = (
                "Correction Gemini request failed."
            )

            correction_attempts.append(
                correction_record
            )

            log(
                f"Correction attempt {attempt_number} "
                "failed at API level."
            )

            continue

        try:

            correction_raw_text = extract_response_text(
                correction_response
            )

            correction_parsed = parse_generated_json(
                correction_raw_text
            )

            corrected_analyses = normalize_analysis_list(
                correction_parsed
            )

        except Exception as exc:

            correction_record["success"] = False
            correction_record["error"] = str(exc)

            correction_attempts.append(
                correction_record
            )

            log(
                f"Correction attempt {attempt_number} "
                f"failed during parsing: {exc}"
            )

            continue

        corrected_validation = validate_analysis(
            corrected_analyses,
            articles,
            set(geography_master),
            set(sector_master),
        )

        correction_record.update(
            {
                "rawText": correction_raw_text,
                "parsedResponse": correction_parsed,
                "analyses": corrected_analyses,
                "validation": corrected_validation,
                "success": corrected_validation["passed"],
            }
        )

        correction_attempts.append(
            correction_record
        )

        if corrected_validation["passed"]:

            log(
                f"Correction attempt {attempt_number} "
                "passed validation."
            )

            return {
                "success": True,
                "mode": "corrected",
                "prompt": prompt,
                "response": initial_response,
                "rawText": raw_text,
                "parsedResponse": parsed,
                "analyses": corrected_analyses,
                "validation": corrected_validation,
                "initialValidation": validation,
                "initialAnalyses": analyses,
                "correctionAttempts": correction_attempts,
            }

        current_analyses = corrected_analyses
        current_validation = corrected_validation

        log(
            f"Correction attempt {attempt_number} "
            "still contains invalid values."
        )

    # --------------------------------------------------------
    # Correction failed
    # --------------------------------------------------------

    return {
        "success": False,
        "stage": "correction_validation",
        "prompt": prompt,
        "response": initial_response,
        "rawText": raw_text,
        "parsedResponse": parsed,
        "analyses": analyses,
        "validation": validation,
        "correctionAttempts": correction_attempts,
        "error": (
            "Initial analysis contained invalid Geography/Sector "
            "values and correction did not produce a fully valid "
            "result."
        ),
    }


# ============================================================
# BATCH RECORD
# ============================================================

def create_batch_record(
    *,
    batch_number: int,
    articles: list[dict[str, Any]],
    result: dict[str, Any],
    model: str,
) -> dict[str, Any]:

    analyses = result.get(
        "analyses",
        [],
    )

    initial_response = result.get(
        "response",
        {},
    )

    usage = {}

    if isinstance(
        initial_response,
        dict,
    ):

        raw = initial_response.get(
            "raw"
        )

        if isinstance(raw, dict):

            usage_metadata = raw.get(
                "usageMetadata"
            )

            if isinstance(
                usage_metadata,
                dict,
            ):
                usage = usage_metadata

    request_data = {
        "articleCount": len(articles),
        "articleIds": [
            article["articleId"]
            for article in articles
        ],
        "promptCharacters": len(
            result.get("prompt", "")
        ),
        "responseSeconds": initial_response.get(
            "responseSeconds"
        ),
        "maxOutputTokens": gemini_max_output_tokens(),
        "responseMimeType": "application/json",
        "correctionMode": True,
        "correctionAttempts": len(
            result.get(
                "correctionAttempts",
                [],
            )
        ),
    }

    record = {
        "batchNumber": batch_number,
        "storedAt": utc_now_iso(),
        "model": model,
        "articleIds": [
            article["articleId"]
            for article in articles
        ],
        "request": request_data,
        "validation": result.get(
            "validation",
            {},
        ),
        "analysisMode": result.get(
            "mode"
        ),
        "prompt": result.get(
            "prompt"
        ),
        "requestBody": initial_response.get(
            "requestBody"
        ),
        "parsedResponse": result.get(
            "parsedResponse"
        ),
        "analyses": analyses,
        "rawGeminiResponse": initial_response.get(
            "raw"
        ),
        "inputArticles": articles,
        "geminiUsage": usage,
    }

    # --------------------------------------------------------
    # Preserve initial response when corrected
    # --------------------------------------------------------

    if result.get("mode") == "corrected":

        record["initialValidation"] = result.get(
            "initialValidation"
        )

        record["initialAnalyses"] = result.get(
            "initialAnalyses"
        )

        record["correctionAttempts"] = result.get(
            "correctionAttempts"
        )

    return record


# ============================================================
# MAIN
# ============================================================

def main() -> int:

    log("=" * 70)
    log("VGrat FMS - MARKET NEWS AI ANALYZER")
    log("=" * 70)

    # --------------------------------------------------------
    # API KEY
    # --------------------------------------------------------

    api_key = os.environ.get(
        "GEMINI_API_KEY"
    )

    if not api_key:
        log(
            "ERROR: GEMINI_API_KEY environment variable is missing."
        )
        return 1

    model = gemini_model()

    log(f"Gemini model: {model}")
    log(f"Batch size: {BATCH_SIZE}")
    log(f"Max batches: {MAX_BATCHES}")
    log(
        f"Max correction attempts: "
        f"{MAX_CORRECTION_ATTEMPTS}"
    )

    # --------------------------------------------------------
    # Verify inputs
    # --------------------------------------------------------

    if not CURRENT_JSON.exists():
        log(
            f"ERROR: Missing market news file: {CURRENT_JSON}"
        )
        return 1

    if not RESEARCH_FUNDS.exists():
        log(
            f"ERROR: Missing Research Funds.xlsx: "
            f"{RESEARCH_FUNDS}"
        )
        return 1

    # --------------------------------------------------------
    # Load current news
    # --------------------------------------------------------

    log(
        f"Loading production market news: "
        f"{CURRENT_JSON}"
    )

    current_data = load_json(
        CURRENT_JSON
    )

    raw_articles = extract_articles(
        current_data
    )

    log(
        f"Articles found in current.json: "
        f"{len(raw_articles)}"
    )

    normalized_articles = [
        normalize_article(article)
        for article in raw_articles
    ]

    # --------------------------------------------------------
    # Load Excel masters
    # --------------------------------------------------------

    geography_master, sector_master = (
        load_research_masters()
    )

    geography_set = set(
        geography_master
    )

    sector_set = set(
        sector_master
    )

    # --------------------------------------------------------
    # Load existing analysis
    # --------------------------------------------------------

    analysis = load_existing_analysis()

    completed_ids = completed_article_ids(
        analysis
    )

    log(
        f"Previously analyzed article IDs: "
        f"{len(completed_ids)}"
    )

    # --------------------------------------------------------
    # Pending articles
    # --------------------------------------------------------

    pending = [
        article
        for article in normalized_articles
        if article["articleId"]
        not in completed_ids
    ]

    if PROCESS_NEWEST_FIRST:
        pending.sort(
            key=publication_sort_key,
            reverse=True,
        )

    log(
        f"Pending articles: {len(pending)}"
    )

    if not pending:
        log(
            "No pending articles. Nothing to do."
        )
        return 0

    # --------------------------------------------------------
    # Process batches
    # --------------------------------------------------------

    successful_batches = 0

    for batch_index in range(
        MAX_BATCHES
    ):

        start = batch_index * BATCH_SIZE
        end = start + BATCH_SIZE

        batch_articles = pending[
            start:end
        ]

        if not batch_articles:
            break

        batch_number = (
            len(analysis.get("batches", []))
            + 1
        )

        log("")
        log("=" * 70)
        log(
            f"BATCH {batch_number}"
        )
        log("=" * 70)

        log(
            f"Articles in batch: "
            f"{len(batch_articles)}"
        )

        for index, article in enumerate(
            batch_articles,
            start=1,
        ):
            log(
                f"{index}. "
                f"{article['articleId']} | "
                f"{article.get('title')}"
            )

        # ----------------------------------------------------
        # Analyze
        # ----------------------------------------------------

        result = analyze_batch(
            articles=batch_articles,
            geography_master=geography_master,
            sector_master=sector_master,
            api_key=api_key,
            model=model,
        )

        # ----------------------------------------------------
        # Failure
        # ----------------------------------------------------

        if not result["success"]:

            log("")
            log("=" * 70)
            log("BATCH FAILED")
            log("=" * 70)

            log(
                str(
                    result.get(
                        "error",
                        "Unknown error.",
                    )
                )
            )

            validation = result.get(
                "validation"
            )

            if validation:
                log(
                    "Validation errors:"
                )

                for error in validation.get(
                    "errors",
                    [],
                ):
                    log(
                        f"  - {error}"
                    )

            log(
                "No changes were written to "
                "analysis/current.json."
            )

            return 1

        # ----------------------------------------------------
        # Create record
        # ----------------------------------------------------

        batch_record = create_batch_record(
            batch_number=batch_number,
            articles=batch_articles,
            result=result,
            model=model,
        )

        # ----------------------------------------------------
        # Update in memory
        # ----------------------------------------------------

        analysis["batches"].append(
            batch_record
        )

        analysis[
            "totalSuccessfulBatches"
        ] = len(
            analysis["batches"]
        )

        analysis[
            "totalSuccessfulArticles"
        ] = sum(
            len(
                batch.get(
                    "articleIds",
                    [],
                )
            )
            for batch in analysis["batches"]
        )

        analysis[
            "updatedAt"
        ] = utc_now_iso()

        successful_batches += 1

        log("")
        log(
            f"BATCH {batch_number} SUCCESS"
        )

        log(
            f"Mode: "
            f"{result.get('mode')}"
        )

        log(
            f"Articles accepted: "
            f"{len(batch_articles)}"
        )

        log(
            f"Correction attempts: "
            f"{len(result.get('correctionAttempts', []))}"
        )

        # ----------------------------------------------------
        # Only now write to disk
        # ----------------------------------------------------

        save_json_atomic(
            ANALYSIS_CURRENT,
            analysis,
        )

        log(
            f"Written: {ANALYSIS_CURRENT}"
        )

    # --------------------------------------------------------
    # Summary
    # --------------------------------------------------------

    log("")
    log("=" * 70)
    log("ANALYSIS COMPLETE")
    log("=" * 70)

    log(
        f"Successful batches this run: "
        f"{successful_batches}"
    )

    log(
        f"Total successful batches: "
        f"{analysis['totalSuccessfulBatches']}"
    )

    log(
        f"Total successful articles: "
        f"{analysis['totalSuccessfulArticles']}"
    )

    log(
        f"Output: {ANALYSIS_CURRENT}"
    )

    return 0


if __name__ == "__main__":
    sys.exit(main())
