#!/usr/bin/env python3

"""
VGrat FMS - Prudential Top Holdings FAILED-FUND RECOVERY 2

PURPOSE
=======

Recover ONLY funds that still failed during Recovery 1:

    scripts/test_prudential_holdings.py

Recovery 1 remains untouched.

RECOVERY SOURCE
===============

1. output_holdings_recovery/run_summary.json
       |
       v
   stillFailedFundsDetail
       |
       v
2. Funds Links.xlsm
       |
       v
   authoritative Excel URL + PruAccess name
       |
       v
3. Official Prudential Singapore fund page
       |
       v
4. Official Prudential factsheet PDF
       |
       v
5. Normal PDF extraction
       |
       v
6. Primary holdings parser
       |
       v
7. Fallback holdings parser
       |
       v
8. Spatial PDF recovery
       |
       v
9. Exact official-PDF re-download
       |
       v
10. Exact extraction/signature verification


IMPORTANT
=========

This script is a SEPARATE recovery script.

It does NOT modify:

    scripts/test_prudential_holdings.py

It does NOT expand the universe.

It does NOT use PruAccess to discover holdings.

It does NOT use third-party holdings sources.

It does NOT fabricate holdings.

It does NOT estimate percentages.

It does NOT use fuzzy matching.

It does NOT use proximity matching between arbitrary names and
percentages.

It does NOT silently accept a different result during final verification.

Recovery 3 is intentionally NOT included in this script.

Recovery 2 ends after its own:
    1. primary parser
    2. fallback parser
    3. spatial recovery + fallback parser
    4. exact final verification

Any fund still failing those stages remains FAILED for Recovery 2.


HARD RULES
==========

- Recovery 1 stillFailedFundsDetail controls the Recovery 2 universe.
- Funds Links.xlsm controls the authoritative URL and Excel row.
- Only official Prudential Singapore URLs are accepted.
- Only official Prudential Singapore factsheets are accepted.
- No third-party holdings sources.
- No inferred holdings.
- No fabricated holdings.
- No fabricated percentages.
- No forced ten holdings.
- Fewer than ten published holdings is valid.
- Duplicate holding names are allowed.
- Duplicate holding percentages are allowed.
- Published order must be preserved.
- Fixed-income coupon percentages may occur inside holding names.
- The last percentage on a logical fixed-income line is the portfolio
  weight when fallback parsing is required.
- A hyphenated PDF line break is rejoined before percentage parsing.
- Spatial recovery MUST still pass the reconstructed section through
  the fallback parser.
- Spatial recovery does NOT independently pair names and percentages
  using coordinate proximity.
- Final verification downloads the official PDF again.
- Final verification runs the extraction engine again.
- Final verification requires the exact same:
      rank
      holding name
      weight
  signature.
- If exact verification fails, the fund remains FAILED.
- A recovery result never overwrites the baseline result.
- Recovery 3 is a separate script and is NOT executed here.
"""


from __future__ import annotations

import json
import re
import sys
import time
from io import BytesIO
from datetime import datetime, timezone
from pathlib import Path
from urllib.parse import urljoin, urlparse

from openpyxl import load_workbook
from pypdf import PdfReader
from playwright.sync_api import (
    TimeoutError as PlaywrightTimeoutError,
    sync_playwright,
)


# =============================================================================
# CONFIGURATION
# =============================================================================

EXCEL_FILE = Path(
    "Funds Links.xlsm"
)

RECOVERY_1_RUN_SUMMARY_FILE = Path(
    "output_holdings_recovery/run_summary.json"
)

RECOVERY_OUTPUT_DIR = Path(
    "output_holdings_recovery_2"
)

RECOVERY_FUNDS_OUTPUT_DIR = (
    RECOVERY_OUTPUT_DIR / "funds"
)

RECOVERY_RUN_SUMMARY_FILE = (
    RECOVERY_OUTPUT_DIR / "run_summary.json"
)


# =============================================================================
# BROWSER
# =============================================================================

BROWSER_HEADLESS = True

PAGE_TIMEOUT_MS = 120000

FACTSHEET_DOWNLOAD_TIMEOUT_MS = 120000

POST_PAGE_WAIT_MS = 1500

RETRY_COUNT = 3

RETRY_DELAY_SECONDS = 3.0


# =============================================================================
# HOLDINGS
# =============================================================================

MAX_HOLDINGS = 10


# =============================================================================
# OFFICIAL PRUDENTIAL HOSTS
# =============================================================================

PRUDENTIAL_HOSTS = {
    "prudential.com.sg",
    "www.prudential.com.sg",
}


# =============================================================================
# SPATIAL RECOVERY
# =============================================================================

SPATIAL_MIN_NAME_LENGTH = 3

SPATIAL_MAX_LINES_AFTER_HEADING = 80


# =============================================================================
# GENERAL HELPERS
# =============================================================================

def clean_text(
    value,
) -> str:

    if value is None:
        return ""

    text = str(value)

    text = text.replace(
        "\xa0",
        " ",
    )

    text = re.sub(
        r"\s+",
        " ",
        text,
    )

    return text.strip()


def normalize_text(
    value,
) -> str:

    text = clean_text(
        value
    )

    text = text.casefold()

    text = text.replace(
        "–",
        "-",
    )

    text = text.replace(
        "—",
        "-",
    )

    return re.sub(
        r"\s+",
        " ",
        text,
    ).strip()


def utc_now_iso() -> str:

    return (
        datetime.now(
            timezone.utc
        )
        .replace(
            microsecond=0
        )
        .isoformat()
    )


def save_json(
    path: Path,
    data,
) -> None:

    path.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    path.write_text(
        json.dumps(
            data,
            indent=2,
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )


def safe_filename(
    value,
) -> str:

    text = clean_text(
        value
    )

    text = re.sub(
        r"[^A-Za-z0-9._-]+",
        "_",
        text,
    )

    text = text.strip(
        "._"
    )

    if not text:
        text = "fund"

    return text[:120]


def is_prudential_url(
    url: str,
) -> bool:

    try:

        parsed = urlparse(
            url
        )

        hostname = (
            parsed.hostname
            or ""
        ).lower()

        return (
            parsed.scheme.lower()
            in {
                "http",
                "https",
            }
            and hostname
            in PRUDENTIAL_HOSTS
        )

    except Exception:

        return False


def ensure_prudential_url(
    url: str,
) -> str:

    url = clean_text(
        url
    )

    if not url:

        raise RuntimeError(
            "URL is empty."
        )

    if not is_prudential_url(
        url
    ):

        raise RuntimeError(
            "Non-Prudential URL rejected: "
            f"{url}"
        )

    return url


class HoldingsParseFailure(RuntimeError):

    def __init__(
        self,
        message: str,
        section_text: str = "",
        full_text: str = "",
        diagnostics: dict | None = None,
    ) -> None:

        super().__init__(
            message
        )

        self.section_text = section_text

        self.full_text = full_text

        self.diagnostics = (
            diagnostics
            or {}
        )


# =============================================================================
# BASELINE
# =============================================================================

def load_recovery_1_run_summary() -> dict:

    if not RECOVERY_1_RUN_SUMMARY_FILE.exists():

        raise FileNotFoundError(
            "Recovery 1 source run summary not found: "
            f"{RECOVERY_1_RUN_SUMMARY_FILE}"
        )

    data = json.loads(
        RECOVERY_1_RUN_SUMMARY_FILE.read_text(
            encoding="utf-8"
        )
    )

    if not isinstance(
        data,
        dict,
    ):

        raise RuntimeError(
            "Baseline run summary is not a JSON object."
        )

    failed = data.get(
        "stillFailedFundsDetail"
    )

    if not isinstance(
        failed,
        list,
    ):

        raise RuntimeError(
            "Baseline run summary does not contain "
            "a valid stillFailedFundsDetail list."
        )

    return data


# =============================================================================
# EXCEL MASTER UNIVERSE
# =============================================================================

def read_excel_funds() -> dict[int, dict]:

    if not EXCEL_FILE.exists():

        raise FileNotFoundError(
            f"Excel file not found: {EXCEL_FILE}"
        )

    workbook = load_workbook(
        EXCEL_FILE,
        read_only=True,
        keep_vba=True,
        data_only=True,
    )

    worksheet = workbook.active

    funds = {}

    try:

        for row_number in range(
            2,
            worksheet.max_row + 1,
        ):

            column_a = worksheet.cell(
                row=row_number,
                column=1,
            ).value

            column_b = worksheet.cell(
                row=row_number,
                column=2,
            ).value

            prudential_url = clean_text(
                column_a
            )

            pruaccess_name = clean_text(
                column_b
            )

            if not prudential_url:
                continue

            funds[row_number] = {
                "excelRow":
                    row_number,

                "prudentialUrl":
                    prudential_url,

                "pruAccessName":
                    pruaccess_name,
            }

    finally:

        workbook.close()

    if not funds:

        raise RuntimeError(
            "No populated Prudential URLs were found "
            "in Excel Column A."
        )

    return funds


def build_recovery_universe(
    recovery_1: dict,
    excel_funds: dict[int, dict],
) -> list[dict]:

    failed_entries = recovery_1.get(
        "stillFailedFundsDetail"
    )

    if not isinstance(
        failed_entries,
        list,
    ):

        raise RuntimeError(
            "Recovery 1 stillFailedFundsDetail is invalid."
        )

    recovery = []

    seen_rows = set()

    for failure in failed_entries:

        if not isinstance(
            failure,
            dict,
        ):

            raise RuntimeError(
                "A stillFailedFundsDetail entry is not an object."
            )

        if "excelRow" not in failure:

            raise RuntimeError(
                "A Recovery 1 failure entry has no excelRow."
            )

        try:

            excel_row = int(
                failure[
                    "excelRow"
                ]
            )

        except Exception as error:

            raise RuntimeError(
                "Recovery 1 failure contains an invalid excelRow: "
                f"{failure.get('excelRow')}"
            ) from error

        if excel_row in seen_rows:

            raise RuntimeError(
                f"Duplicate Recovery 1 failed excelRow: {excel_row}"
            )

        seen_rows.add(
            excel_row
        )

        if excel_row not in excel_funds:

            raise RuntimeError(
                "Recovery 1 failed row "
                f"{excel_row} does not exist in Excel."
            )

        excel_fund = excel_funds[
            excel_row
        ]

        baseline_url = clean_text(
            failure.get(
                "prudentialUrl"
            )
        )

        baseline_name = clean_text(
            failure.get(
                "pruAccessName"
            )
            or
            failure.get(
                "excelPruAccessName"
            )
        )

        excel_url = clean_text(
            excel_fund[
                "prudentialUrl"
            ]
        )

        excel_name = clean_text(
            excel_fund.get(
                "pruAccessName"
            )
        )

        if (
            baseline_url
            and
            baseline_url
            !=
            excel_url
        ):

            raise RuntimeError(
                "Recovery 1 URL does not match Excel for "
                f"row {excel_row}."
            )

        if (
            baseline_name
            and
            excel_name
            and
            baseline_name
            !=
            excel_name
        ):

            raise RuntimeError(
                "Recovery 1 PruAccess name does not match Excel for "
                f"row {excel_row}."
            )

        recovery.append(
            {
                "excelRow":
                    excel_row,

                "prudentialUrl":
                    excel_url,

                "pruAccessName":
                    excel_name,

                "baselineFailure":
                    failure,
            }
        )

    return recovery


# =============================================================================
# FACTSHEET DISCOVERY
# =============================================================================

def find_factsheet_url(
    page,
) -> str:

    anchors = page.locator(
        "a"
    )

    count = anchors.count()

    candidates = []

    for index in range(
        count
    ):

        anchor = anchors.nth(
            index
        )

        try:

            href = clean_text(
                anchor.get_attribute(
                    "href"
                )
            )

            if not href:
                continue

            absolute_url = urljoin(
                page.url,
                href,
            )

            if not is_prudential_url(
                absolute_url
            ):
                continue

            anchor_text = clean_text(
                anchor.inner_text()
            )

        except Exception:

            continue

        href_lower = (
            absolute_url.lower()
        )

        text_lower = (
            anchor_text.lower()
        )

        score = 0

        if "fund factsheet" in text_lower:
            score += 200

        elif "factsheet" in text_lower:
            score += 150

        if "factsheet" in href_lower:
            score += 100

        if href_lower.endswith(
            ".pdf"
        ):
            score += 50

        if score > 0:

            candidates.append(
                {
                    "score":
                        score,

                    "url":
                        absolute_url,

                    "text":
                        anchor_text,
                }
            )

    if not candidates:

        raise RuntimeError(
            "Could not find an official Prudential factsheet link."
        )

    candidates.sort(
        key=lambda item: (
            -item["score"],
            item["url"],
        )
    )

    return ensure_prudential_url(
        candidates[0]["url"]
    )


# =============================================================================
# PDF DOWNLOAD
# =============================================================================

def download_factsheet(
    page,
    factsheet_url: str,
) -> bytes:

    factsheet_url = ensure_prudential_url(
        factsheet_url
    )

    response = page.request.get(
        factsheet_url,
        timeout=FACTSHEET_DOWNLOAD_TIMEOUT_MS,
    )

    if response.status != 200:

        raise RuntimeError(
            "Factsheet HTTP status was "
            f"{response.status}."
        )

    body = response.body()

    if not body:

        raise RuntimeError(
            "Factsheet response contained zero bytes."
        )

    if not body.startswith(
        b"%PDF"
    ):

        content_type = clean_text(
            response.headers.get(
                "content-type"
            )
        )

        raise RuntimeError(
            "Factsheet response was not a PDF. "
            f"Content-Type={content_type}"
        )

    return body


# =============================================================================
# PDF TEXT
# =============================================================================

def extract_pdf_text(
    pdf_bytes: bytes,
) -> tuple[str, int]:

    reader = PdfReader(
        BytesIO(
            pdf_bytes
        )
    )

    page_count = len(
        reader.pages
    )

    if page_count == 0:

        raise RuntimeError(
            "Factsheet PDF has zero pages."
        )

    pages = []

    for page_number, pdf_page in enumerate(
        reader.pages,
        start=1,
    ):

        try:

            text = (
                pdf_page.extract_text()
                or ""
            )

        except Exception as error:

            raise RuntimeError(
                "Failed to extract PDF text from page "
                f"{page_number}: {error}"
            ) from error

        pages.append(
            text
        )

    full_text = "\n".join(
        pages
    )

    if not clean_text(
        full_text
    ):

        raise RuntimeError(
            "Factsheet PDF contains no extractable text."
        )

    return (
        full_text,
        page_count,
    )


# =============================================================================
# PDF LINES
# =============================================================================

def pdf_lines(
    text: str,
) -> list[str]:

    result = []

    for raw_line in text.splitlines():

        line = clean_text(
            raw_line
        )

        if line:
            result.append(
                line
            )

    return result


# =============================================================================
# FACTSHEET DATES
# =============================================================================

def extract_data_as_at(
    text: str,
) -> str | None:

    patterns = [

        re.compile(
            r"\bdata\s+as\s+at\s*"
            r"([0-9]{1,2}\s+[A-Za-z]{3,9}\s+[0-9]{4})\b",
            re.IGNORECASE,
        ),

        re.compile(
            r"\ball\s+data\s+as\s+at\s*"
            r"([0-9]{1,2}\s+[A-Za-z]{3,9}\s+[0-9]{4})\b",
            re.IGNORECASE,
        ),
    ]

    for pattern in patterns:

        match = pattern.search(
            text
        )

        if match:

            return clean_text(
                match.group(1)
            )

    return None


def extract_document_date(
    text: str,
) -> str | None:

    pattern = re.compile(
        r"\b("
        r"January|February|March|April|May|June|July|"
        r"August|September|October|November|December"
        r")\s+([0-9]{4})\b",
        re.IGNORECASE,
    )

    for line in pdf_lines(
        text
    )[:50]:

        match = pattern.search(
            line
        )

        if match:

            return clean_text(
                match.group(0)
            )

    return None


# =============================================================================
# HOLDINGS SECTION
# =============================================================================

def find_holdings_start(
    lines: list[str],
) -> int | None:

    for index, line in enumerate(
        lines
    ):

        normalized = normalize_text(
            line
        )

        if (
            "top 10 holdings"
            in normalized
            or
            "top ten holdings"
            in normalized
        ):

            return index

    return None


def is_holdings_end(
    line: str,
) -> bool:

    normalized = normalize_text(
        line
    )

    if normalized in {
        "source",
        "source:",
        "inception date",
        "important information",
        "important information:",
        "disclaimer",
        "past performance",
        "portfolio characteristics",
        "asset allocation",
    }:

        return True

    if normalized.startswith(
        "source:"
    ):

        return True

    if normalized.startswith(
        "inception date:"
    ):

        return True

    if normalized.startswith(
        "important information"
    ):

        return True

    if re.fullmatch(
        r"page\s+\d+(\s+of\s+\d+)?",
        normalized,
    ):

        return True

    return False


def extract_holdings_section(
    text: str,
) -> tuple[str, str]:

    lines = pdf_lines(
        text
    )

    start = find_holdings_start(
        lines
    )

    if start is None:

        return (
            "",
            "not_published",
        )

    section = []

    for line in lines[
        start + 1:
    ]:

        if is_holdings_end(
            line
        ):

            break

        section.append(
            line
        )

    return (
        "\n".join(
            section
        ),
        "published",
    )


# =============================================================================
# NAME CLEANING
# =============================================================================

def clean_holding_name(
    value: str,
) -> str:

    name = clean_text(
        value
    )

    if not name:
        return ""

    name = re.sub(
        r"^[•·▪■□*]+",
        "",
        name,
    )

    name = re.sub(
        r"^\d{1,2}[.)]\s+",
        "",
        name,
    )

    name = name.replace(
        "|",
        " ",
    )

    name = re.sub(
        r"\bNone\b",
        " ",
        name,
        flags=re.IGNORECASE,
    )

    name = clean_text(
        name
    )

    name = name.strip(
        " -|"
    )

    if normalize_text(
        name
    ) in {
        "",
        "none",
        "null",
        "-",
        "—",
    }:

        return ""

    return name


def clean_holding_fragment(
    value: str,
) -> str:

    fragment = clean_text(
        value
    )

    if not fragment:
        return ""

    fragment = fragment.replace(
        "|",
        " ",
    )

    fragment = re.sub(
        r"\bNone\b",
        " ",
        fragment,
        flags=re.IGNORECASE,
    )

    fragment = re.sub(
        r"^[•·▪■□*]+",
        "",
        fragment,
    )

    return clean_text(
        fragment
    )


def combine_holding_name_fragments(
    fragments: list[str],
) -> str:

    cleaned = []

    for fragment in fragments:

        fragment = clean_holding_fragment(
            fragment
        )

        if not fragment:
            continue

        if re.fullmatch(
            r"\d{1,2}",
            fragment,
        ):

            continue

        cleaned.append(
            fragment
        )

    if not cleaned:
        return ""

    combined = " ".join(
        cleaned
    )

    combined = re.sub(
        r"\s+",
        " ",
        combined,
    ).strip()

    return clean_holding_name(
        combined
    )


# =============================================================================
# PERCENTAGES
# =============================================================================

def percentage_matches(
    line: str,
) -> list[re.Match]:

    return list(
        re.finditer(
            r"(?<![\d.])"
            r"([0-9]+(?:\.[0-9]+)?)"
            r"\s*%",
            line,
        )
    )


def find_percentage_in_line(
    line: str,
) -> tuple[float, str, int, int] | None:

    matches = percentage_matches(
        line
    )

    if not matches:
        return None

    match = matches[0]

    value = float(
        match.group(1)
    )

    if not 0 <= value <= 100:
        return None

    return (
        value,
        clean_text(
            match.group(0)
        ),
        match.start(),
        match.end(),
    )


def find_last_percentage_in_line(
    line: str,
) -> tuple[float, str, int, int] | None:

    matches = percentage_matches(
        line
    )

    if not matches:
        return None

    match = matches[-1]

    value = float(
        match.group(1)
    )

    if not 0 <= value <= 100:
        return None

    return (
        value,
        clean_text(
            match.group(0)
        ),
        match.start(),
        match.end(),
    )


# =============================================================================
# RANKS
# =============================================================================

def extract_leading_rank(
    line: str,
) -> tuple[int | None, str]:

    text = clean_text(
        line
    )

    if not text:
        return (
            None,
            "",
        )

    match = re.match(
        r"^\s*(\d{1,2})[.)\-:]\s*(.*)$",
        text,
    )

    if match:

        rank = int(
            match.group(1)
        )

        if 1 <= rank <= 99:

            return (
                rank,
                clean_text(
                    match.group(2)
                ),
            )

    match = re.match(
        r"^\s*(\d{1,2})\s+(.+)$",
        text,
    )

    if match:

        rank = int(
            match.group(1)
        )

        if 1 <= rank <= 99:

            return (
                rank,
                clean_text(
                    match.group(2)
                ),
            )

    return (
        None,
        text,
    )


# =============================================================================
# NOISE
# =============================================================================

def is_holding_header_or_noise(
    line: str,
) -> bool:

    normalized = normalize_text(
        line
    )

    if not normalized:
        return True

    if normalized in {
        "holding",
        "holdings",
        "name",
        "names",
        "weight",
        "weights",
        "allocation",
        "allocations",
        "%",
        "portfolio holdings",
    }:

        return True

    if normalized.startswith(
        "top 10 holdings"
    ):

        return True

    if normalized.startswith(
        "top ten holdings"
    ):

        return True

    return False


# =============================================================================
# HYPHENATED LINE RECONSTRUCTION
# =============================================================================

def merge_hyphenated_line_breaks(
    lines: list[str],
) -> list[str]:

    merged = []

    index = 0

    while index < len(lines):

        current = clean_text(
            lines[index]
        )

        while (
            current.endswith("-")
            and
            index + 1 < len(lines)
        ):

            index += 1

            next_line = clean_text(
                lines[index]
            )

            current = (
                current
                +
                next_line
            )

        merged.append(
            current
        )

        index += 1

    return merged


# =============================================================================
# LOGICAL LINES
# =============================================================================

def build_logical_holding_lines(
    section_text: str,
) -> list[str]:

    raw_lines = pdf_lines(
        section_text
    )

    if not raw_lines:
        return []

    return merge_hyphenated_line_breaks(
        raw_lines
    )


# =============================================================================
# VALIDATION
# =============================================================================

def validate_holdings(
    holdings: list[dict],
    parser_name: str,
) -> list[dict]:

    if not holdings:

        raise RuntimeError(
            f"{parser_name} parser produced no holdings."
        )

    if len(holdings) > MAX_HOLDINGS:

        raise RuntimeError(
            f"{parser_name} parser produced more than "
            f"{MAX_HOLDINGS} holdings."
        )

    expected = list(
        range(
            1,
            len(holdings) + 1,
        )
    )

    ranks = [
        item[
            "rank"
        ]
        for item in holdings
    ]

    if ranks != expected:

        raise RuntimeError(
            f"{parser_name} parser ranks invalid. "
            f"Parsed={ranks}; Expected={expected}"
        )

    for holding in holdings:

        name = clean_holding_name(
            holding.get(
                "name"
            )
        )

        if not name:

            raise RuntimeError(
                f"{parser_name} parser produced an empty holding name."
            )

        weight = holding.get(
            "weightPercent"
        )

        if not isinstance(
            weight,
            (int, float),
        ):

            raise RuntimeError(
                f"{parser_name} parser produced an invalid weight."
            )

        if not 0 <= weight <= 100:

            raise RuntimeError(
                f"{parser_name} parser produced weight outside 0-100."
            )

        holding[
            "name"
        ] = name

    return holdings


# =============================================================================
# PRIMARY PARSER
# =============================================================================

def parse_holdings(
    section_text: str,
) -> list[dict]:

    lines = build_logical_holding_lines(
        section_text
    )

    if not lines:

        raise RuntimeError(
            "Holdings section is empty."
        )

    holdings = []

    pending_fragments = []

    pending_rank = None

    for raw_line in lines:

        line = clean_text(
            raw_line
        )

        if not line:
            continue

        if is_holding_header_or_noise(
            line
        ):

            continue

        detected_rank, remainder = (
            extract_leading_rank(
                line
            )
        )

        if detected_rank is not None:

            if pending_fragments:

                raise RuntimeError(
                    "Primary parser encountered a new rank "
                    "before the previous holding received a weight."
                )

            pending_rank = detected_rank

            line = remainder

            if not line:
                continue

        matches = percentage_matches(
            line
        )

        if len(matches) > 1:

            raise RuntimeError(
                "Primary parser found multiple percentages on "
                "one logical line."
            )

        percentage_info = (
            find_percentage_in_line(
                line
            )
        )

        if percentage_info:

            (
                weight,
                weight_text,
                start,
                _end,
            ) = percentage_info

            fragment = clean_text(
                line[:start]
            )

            if fragment:
                pending_fragments.append(
                    fragment
                )

            name = combine_holding_name_fragments(
                pending_fragments
            )

            if not name:

                raise RuntimeError(
                    "Primary parser found a percentage "
                    "without a holding name."
                )

            rank = (
                pending_rank
                if pending_rank is not None
                else len(holdings) + 1
            )

            holdings.append(
                {
                    "rank":
                        rank,

                    "name":
                        name,

                    "weightPercent":
                        weight,

                    "weightText":
                        weight_text,
                }
            )

            pending_fragments = []

            pending_rank = None

            if len(holdings) >= MAX_HOLDINGS:

                break

            continue

        fragment = clean_holding_fragment(
            line
        )

        if fragment:
            pending_fragments.append(
                fragment
            )

    return validate_holdings(
        holdings,
        "Primary",
    )


# =============================================================================
# FALLBACK PARSER
# =============================================================================

def parse_holdings_fallback(
    section_text: str,
) -> list[dict]:

    lines = build_logical_holding_lines(
        section_text
    )

    if not lines:

        raise RuntimeError(
            "Fallback parser received an empty section."
        )

    holdings = []

    pending_fragments = []

    pending_rank = None

    for raw_line in lines:

        line = clean_text(
            raw_line
        )

        if not line:
            continue

        if is_holding_header_or_noise(
            line
        ):

            continue

        detected_rank, remainder = (
            extract_leading_rank(
                line
            )
        )

        if detected_rank is not None:

            if pending_fragments:

                raise RuntimeError(
                    "Fallback parser encountered a new rank "
                    "before the previous holding received a weight."
                )

            pending_rank = detected_rank

            line = remainder

            if not line:
                continue

        percentage_info = (
            find_last_percentage_in_line(
                line
            )
        )

        if percentage_info:

            (
                weight,
                weight_text,
                start,
                _end,
            ) = percentage_info

            fragment = clean_text(
                line[:start]
            )

            if fragment:
                pending_fragments.append(
                    fragment
                )

            name = combine_holding_name_fragments(
                pending_fragments
            )

            if not name:

                raise RuntimeError(
                    "Fallback parser found a percentage "
                    "without a holding name."
                )

            rank = (
                pending_rank
                if pending_rank is not None
                else len(holdings) + 1
            )

            holdings.append(
                {
                    "rank":
                        rank,

                    "name":
                        name,

                    "weightPercent":
                        weight,

                    "weightText":
                        weight_text,
                }
            )

            pending_fragments = []

            pending_rank = None

            if len(holdings) >= MAX_HOLDINGS:

                break

            continue

        fragment = clean_holding_fragment(
            line
        )

        if fragment:

            pending_fragments.append(
                fragment
            )

    return validate_holdings(
        holdings,
        "Fallback",
    )


# =============================================================================
# HOLDING SIGNATURE
# =============================================================================

def holding_signature(
    holdings: list[dict],
) -> list[tuple]:

    signature = []

    for holding in holdings:

        signature.append(
            (
                int(
                    holding[
                        "rank"
                    ]
                ),

                normalize_text(
                    holding[
                        "name"
                    ]
                ),

                float(
                    holding[
                        "weightPercent"
                    ]
                ),
            )
        )

    return signature


# =============================================================================
# SPATIAL PDF EXTRACTION
# =============================================================================

def positioned_pdf_lines(
    pdf_bytes: bytes,
) -> list[dict]:

    reader = PdfReader(
        BytesIO(
            pdf_bytes
        )
    )

    all_lines = []

    for page_number, page in enumerate(
        reader.pages,
        start=1,
    ):

        fragments = []

        def visitor_text(
            text,
            cm,
            tm,
            font_dict,
            font_size,
        ):

            value = clean_text(
                text
            )

            if not value:
                return

            try:

                x = float(
                    tm[4]
                )

                y = float(
                    tm[5]
                )

            except Exception:

                return

            fragments.append(
                {
                    "text":
                        value,

                    "x":
                        x,

                    "y":
                        y,

                    "fontSize":
                        float(
                            font_size
                            or
                            0
                        ),
                }
            )

        try:

            page.extract_text(
                visitor_text=visitor_text
            )

        except Exception:

            continue

        fragments.sort(
            key=lambda item: (
                -item["y"],
                item["x"],
            )
        )

        page_lines = []

        tolerance = 3.0

        for fragment in fragments:

            if not page_lines:

                page_lines.append(
                    {
                        "page":
                            page_number,

                        "y":
                            fragment[
                                "y"
                            ],

                        "fragments":
                            [
                                fragment
                            ],
                    }
                )

                continue

            current = page_lines[
                -1
            ]

            if abs(
                current["y"]
                -
                fragment["y"]
            ) <= tolerance:

                current[
                    "fragments"
                ].append(
                    fragment
                )

            else:

                page_lines.append(
                    {
                        "page":
                            page_number,

                        "y":
                            fragment[
                                "y"
                            ],

                        "fragments":
                            [
                                fragment
                            ],
                    }
                )

        for line in page_lines:

            fragments_sorted = sorted(
                line[
                    "fragments"
                ],
                key=lambda item: item[
                    "x"
                ],
            )

            text = clean_text(
                " ".join(
                    item[
                        "text"
                    ]
                    for item
                    in fragments_sorted
                )
            )

            if not text:
                continue

            all_lines.append(
                {
                    "page":
                        line[
                            "page"
                        ],

                    "y":
                        line[
                            "y"
                        ],

                    "xMin":
                        min(
                            item[
                                "x"
                            ]
                            for item
                            in fragments_sorted
                        ),

                    "xMax":
                        max(
                            item[
                                "x"
                            ]
                            for item
                            in fragments_sorted
                        ),

                    "text":
                        text,
                }
            )

    return all_lines


def find_spatial_heading(
    lines: list[dict],
) -> int | None:

    for index, line in enumerate(
        lines
    ):

        normalized = normalize_text(
            line[
                "text"
            ]
        )

        if (
            "top 10 holdings"
            in normalized
            or
            "top ten holdings"
            in normalized
        ):

            return index

    return None


def spatial_line_has_weight(
    line: dict,
) -> bool:

    return bool(
        percentage_matches(
            line[
                "text"
            ]
        )
    )


def build_spatial_candidates(
    lines: list[dict],
    heading_index: int,
) -> list[dict]:

    heading = lines[
        heading_index
    ]

    heading_page = heading[
        "page"
    ]

    candidates = []

    for index in range(
        heading_index + 1,
        min(
            len(lines),
            heading_index
            +
            SPATIAL_MAX_LINES_AFTER_HEADING
            +
            1,
        ),
    ):

        line = lines[
            index
        ]

        if line[
            "page"
        ] != heading_page:

            break

        text = clean_text(
            line[
                "text"
            ]
        )

        if not text:
            continue

        if is_holdings_end(
            text
        ):

            break

        if normalize_text(
            text
        ) in {
            "name",
            "holding",
            "holdings",
            "weight",
            "weights",
            "allocation",
        }:

            continue

        candidates.append(
            line
        )

    return candidates


def choose_spatial_left_column(
    lines: list[dict],
) -> list[dict]:

    if not lines:
        return []

    weight_lines = [
        line
        for line
        in lines
        if spatial_line_has_weight(
            line
        )
    ]

    if not weight_lines:
        return []

    rightmost_weight_x = max(
        line[
            "xMax"
        ]
        for line
        in weight_lines
    )

    # The holdings name column normally occupies the area left of the
    # portfolio-weight column. We do not pair coordinates here.
    # We only reconstruct the textual left-side section.
    boundary = rightmost_weight_x

    selected = []

    for line in lines:

        text = line[
            "text"
        ]

        if not text:
            continue

        if line[
            "xMin"
        ] < boundary:

            selected.append(
                line
            )

    return selected


def spatial_lines_to_section(
    lines: list[dict],
) -> str:

    if not lines:
        return ""

    # Preserve PDF reading order.
    lines = sorted(
        lines,
        key=lambda item: (
            item[
                "page"
            ],
            -item[
                "y"
            ],
            item[
                "xMin"
            ],
        )
    )

    return "\n".join(
        clean_text(
            line[
                "text"
            ]
        )
        for line
        in lines
        if clean_text(
            line[
                "text"
            ]
        )
    )


def extract_spatial_fallback(
    pdf_bytes: bytes,
) -> tuple[
    list[dict],
    str,
    dict,
]:

    lines = positioned_pdf_lines(
        pdf_bytes
    )

    diagnostics = {
        "positionedLineCount":
            len(lines),

        "headingFound":
            False,

        "candidateLineCount":
            0,

        "selectedLineCount":
            0,

        "parser":
            "fallback",
    }

    heading_index = find_spatial_heading(
        lines
    )

    if heading_index is None:

        raise RuntimeError(
            "Spatial recovery could not locate "
            "Top 10 holdings heading."
        )

    diagnostics[
        "headingFound"
    ] = True

    candidate_lines = (
        build_spatial_candidates(
            lines,
            heading_index,
        )
    )

    diagnostics[
        "candidateLineCount"
    ] = len(
        candidate_lines
    )

    selected_lines = (
        choose_spatial_left_column(
            candidate_lines
        )
    )

    diagnostics[
        "selectedLineCount"
    ] = len(
        selected_lines
    )

    section = spatial_lines_to_section(
        selected_lines
    )

    if not section:

        raise RuntimeError(
            "Spatial recovery reconstructed an empty "
            "Top 10 holdings section."
        )

    # IMPORTANT:
    #
    # Spatial recovery is only allowed to reconstruct the section.
    # It MUST then use the normal fallback parser.
    holdings = parse_holdings_fallback(
        section
    )

    diagnostics[
        "parsedHoldingCount"
    ] = len(
        holdings
    )

    return (
        holdings,
        section,
        diagnostics,
    )


# =============================================================================
# EXTRACTION ENGINE
# =============================================================================

def extract_with_recovery_engine(
    pdf_bytes: bytes,
) -> dict:

    (
        full_text,
        page_count,
    ) = extract_pdf_text(
        pdf_bytes
    )

    (
        section_text,
        section_status,
    ) = extract_holdings_section(
        full_text
    )

    if section_status == "not_published":

        return {
            "status":
                "no_holdings_section",

            "holdings":
                [],

            "parser":
                None,

            "primaryParserError":
                None,

            "fallbackParserError":
                None,

            "spatialRecoveryError":
                None,

            "fullText":
                full_text,

            "sectionText":
                "",

            "pageCount":
                page_count,

            "spatialDiagnostics":
                None,
        }

    primary_error = None

    try:

        holdings = parse_holdings(
            section_text
        )

        return {
            "status":
                "success",

            "holdings":
                holdings,

            "parser":
                "primary",

            "primaryParserError":
                None,

            "fallbackParserError":
                None,

            "spatialRecoveryError":
                None,

            "fullText":
                full_text,

            "sectionText":
                section_text,

            "pageCount":
                page_count,

            "spatialDiagnostics":
                None,
        }

    except Exception as error:

        primary_error = clean_text(
            str(error)
        )

    fallback_error = None

    try:

        holdings = parse_holdings_fallback(
            section_text
        )

        return {
            "status":
                "success",

            "holdings":
                holdings,

            "parser":
                "fallback",

            "primaryParserError":
                primary_error,

            "fallbackParserError":
                None,

            "spatialRecoveryError":
                None,

            "fullText":
                full_text,

            "sectionText":
                section_text,

            "pageCount":
                page_count,

            "spatialDiagnostics":
                None,
        }

    except Exception as error:

        fallback_error = clean_text(
            str(error)
        )

    spatial_error = None

    try:

        (
            spatial_holdings,
            spatial_section,
            spatial_diagnostics,
        ) = extract_spatial_fallback(
            pdf_bytes
        )

        return {
            "status":
                "success",

            "holdings":
                spatial_holdings,

            "parser":
                "spatial_fallback",

            "primaryParserError":
                primary_error,

            "fallbackParserError":
                fallback_error,

            "spatialRecoveryError":
                None,

            "fullText":
                full_text,

            "sectionText":
                spatial_section,

            "pageCount":
                page_count,

            "spatialDiagnostics":
                spatial_diagnostics,
        }

    except Exception as error:

        spatial_error = clean_text(
            str(error)
        )

    raise HoldingsParseFailure(
        "All Recovery 2 extraction engines failed. "
        f"Primary={primary_error}; "
        f"Fallback={fallback_error}; "
        f"Spatial={spatial_error}",
        section_text=section_text,
        full_text=full_text,
        diagnostics={
            "primaryParserError":
                primary_error,

            "fallbackParserError":
                fallback_error,

            "spatialRecoveryError":
                spatial_error,
        },
    )


# =============================================================================
# FUND PAGE
# =============================================================================

def extract_fund_page_name(
    page,
) -> str:

    try:

        h1 = page.locator(
            "h1"
        ).first

        if h1.count():

            name = clean_text(
                h1.inner_text()
            )

            if name:
                return name

    except Exception:

        pass

    try:

        body = clean_text(
            page.locator(
                "body"
            ).inner_text()
        )

        match = re.search(
            r"\bPRU(?:Link|Prime)\s+[^\n]+",
            body,
            re.IGNORECASE,
        )

        if match:

            return clean_text(
                match.group(0)
            )

    except Exception:

        pass

    return ""


# =============================================================================
# SINGLE FUND RECOVERY
# =============================================================================

def recover_single_fund(
    page,
    excel_fund: dict,
) -> dict:

    excel_row = int(
        excel_fund[
            "excelRow"
        ]
    )

    prudential_url = ensure_prudential_url(
        excel_fund[
            "prudentialUrl"
        ]
    )

    print(
        "\n"
        + "=" * 72
    )

    print(
        f"RECOVERY 2 - FUND ROW {excel_row}"
    )

    print(
        "=" * 72
    )

    print(
        f"URL: {prudential_url}"
    )

    page.goto(
        prudential_url,
        wait_until="domcontentloaded",
        timeout=PAGE_TIMEOUT_MS,
    )

    try:

        page.wait_for_load_state(
            "networkidle",
            timeout=25000,
        )

    except PlaywrightTimeoutError:

        pass

    page.wait_for_timeout(
        POST_PAGE_WAIT_MS
    )

    final_url = clean_text(
        page.url
    )

    if not is_prudential_url(
        final_url
    ):

        raise RuntimeError(
            "Fund page redirected outside Prudential Singapore: "
            f"{final_url}"
        )

    page_title = clean_text(
        page.title()
    )

    fund_name = extract_fund_page_name(
        page
    )

    factsheet_url = find_factsheet_url(
        page
    )

    print(
        f"Fund: {fund_name or '-'}"
    )

    print(
        f"Factsheet: {factsheet_url}"
    )

    factsheet_bytes = download_factsheet(
        page,
        factsheet_url,
    )

    engine = extract_with_recovery_engine(
        factsheet_bytes
    )

    if engine[
        "status"
    ] == "no_holdings_section":

        return {
            "status":
                "no_holdings_section",

            "excelRow":
                excel_row,

            "prudentialUrl":
                prudential_url,

            "finalUrl":
                final_url,

            "pageTitle":
                page_title,

            "fundName":
                fund_name,

            "excelPruAccessName":
                excel_fund.get(
                    "pruAccessName"
                ),

            "factsheetUrl":
                factsheet_url,

            "factsheetPageCount":
                engine[
                    "pageCount"
                ],

            "factsheetDocumentDate":
                extract_document_date(
                    engine[
                        "fullText"
                    ]
                ),

            "factsheetDataAsAt":
                extract_data_as_at(
                    engine[
                        "fullText"
                    ]
                ),

            "topHoldingsCount":
                0,

            "topHoldings":
                [],

            "holdingsParser":
                None,

            "recoveryEngine":
                "no_holdings_section",

            "_factsheetBytes":
                factsheet_bytes,

            "_fullText":
                engine[
                    "fullText"
                ],

            "_sectionText":
                "",
        }

    holdings = engine[
        "holdings"
    ]

    result = {
        "status":
            "success",

        "excelRow":
            excel_row,

        "prudentialUrl":
            prudential_url,

        "finalUrl":
            final_url,

        "pageTitle":
            page_title,

        "fundName":
            fund_name,

        "excelPruAccessName":
            excel_fund.get(
                "pruAccessName"
            ),

        "factsheetUrl":
            factsheet_url,

        "factsheetDocumentDate":
            extract_document_date(
                engine[
                    "fullText"
                ]
            ),

        "factsheetDataAsAt":
            extract_data_as_at(
                engine[
                    "fullText"
                ]
            ),

        "factsheetPageCount":
            engine[
                "pageCount"
            ],

        "holdingsSectionStatus":
            "published",

        "topHoldingsCount":
            len(
                holdings
            ),

        "topHoldings":
            holdings,

        "holdingsParser":
            engine[
                "parser"
            ],

        "primaryParserError":
            engine[
                "primaryParserError"
            ],

        "fallbackParserError":
            engine[
                "fallbackParserError"
            ],

        "spatialRecoveryError":
            engine[
                "spatialRecoveryError"
            ],

        "spatialDiagnostics":
            engine[
                "spatialDiagnostics"
            ],

        "_factsheetBytes":
            factsheet_bytes,

        "_fullText":
            engine[
                "fullText"
            ],

        "_sectionText":
            engine[
                "sectionText"
            ],
    }

    return result


# =============================================================================
# EXACT FINAL VERIFICATION
# =============================================================================

def verify_exact_against_official_pdf(
    page,
    result: dict,
) -> dict:

    print(
        "FINAL VERIFICATION: downloading official PDF again..."
    )

    factsheet_url = ensure_prudential_url(
        result[
            "factsheetUrl"
        ]
    )

    response = page.request.get(
        factsheet_url,
        timeout=FACTSHEET_DOWNLOAD_TIMEOUT_MS,
    )

    if response.status != 200:

        raise RuntimeError(
            "Final verification factsheet download returned "
            f"HTTP {response.status}."
        )

    verification_bytes = response.body()

    if not verification_bytes.startswith(
        b"%PDF"
    ):

        raise RuntimeError(
            "Final verification response was not a PDF."
        )

    verification_engine = extract_with_recovery_engine(
        verification_bytes
    )

    if verification_engine[
        "status"
    ] != "success":

        raise RuntimeError(
            "Final verification could not extract published holdings."
        )

    original_holdings = result[
        "topHoldings"
    ]

    verified_holdings = (
        verification_engine[
            "holdings"
        ]
    )

    original_signature = holding_signature(
        original_holdings
    )

    verified_signature = holding_signature(
        verified_holdings
    )

    if (
        original_signature
        !=
        verified_signature
    ):

        diagnostics = {
            "originalSignature":
                original_signature,

            "verifiedSignature":
                verified_signature,

            "originalHoldings":
                original_holdings,

            "verifiedHoldings":
                verified_holdings,

            "verificationParser":
                verification_engine[
                    "parser"
                ],

            "verificationPrimaryParserError":
                verification_engine[
                    "primaryParserError"
                ],

            "verificationFallbackParserError":
                verification_engine[
                    "fallbackParserError"
                ],

            "verificationSpatialDiagnostics":
                verification_engine.get(
                    "spatialDiagnostics"
                ),
        }

        raise HoldingsParseFailure(
            "EXACT VERIFICATION FAILED: the re-downloaded "
            "official Prudential PDF produced a different "
            "rank/name/weight signature.",
            section_text=verification_engine[
                "sectionText"
            ],
            full_text=verification_engine[
                "fullText"
            ],
            diagnostics=diagnostics,
        )

    return {
        "verified":
            True,

        "verifiedParser":
            verification_engine[
                "parser"
            ],

        "verifiedHoldingCount":
            len(
                verified_holdings
            ),

        "verifiedSignature":
            verified_signature,

        "verificationFactsheetBytes":
            verification_bytes,

        "verificationFullText":
            verification_engine[
                "fullText"
            ],

        "verificationSectionText":
            verification_engine[
                "sectionText"
            ],
    }


# =============================================================================
# OUTPUT HELPERS
# =============================================================================

def recovery_directory(
    result: dict,
) -> Path:

    excel_row = int(
        result[
            "excelRow"
        ]
    )

    identifier = safe_filename(
        result.get(
            "fundName"
        )
        or
        urlparse(
            result.get(
                "finalUrl"
            )
            or
            result.get(
                "prudentialUrl"
            )
        ).path.rstrip(
            "/"
        ).split(
            "/"
        )[-1]
        or
        f"fund_{excel_row}"
    )

    return (
        RECOVERY_FUNDS_OUTPUT_DIR
        /
        f"{excel_row}_{identifier}"
    )


def save_success(
    result: dict,
    verification: dict,
) -> Path:

    directory = recovery_directory(
        result
    )

    directory.mkdir(
        parents=True,
        exist_ok=True,
    )

    verification_bytes = (
        verification[
            "verificationFactsheetBytes"
        ]
    )

    verification_full_text = (
        verification[
            "verificationFullText"
        ]
    )

    verification_section = (
        verification[
            "verificationSectionText"
        ]
    )

    (
        directory
        /
        "factsheet.pdf"
    ).write_bytes(
        verification_bytes
    )

    (
        directory
        /
        "factsheet_text.txt"
    ).write_text(
        verification_full_text,
        encoding="utf-8",
    )

    (
        directory
        /
        "top_holdings_section.txt"
    ).write_text(
        verification_section,
        encoding="utf-8",
    )

    clean_result = {
        key:
            value
        for key, value
        in result.items()
        if not key.startswith("_")
    }

    clean_result[
        "verification"
    ] = {
        "verified":
            True,

        "verifiedParser":
            verification[
                "verifiedParser"
            ],

        "verifiedHoldingCount":
            verification[
                "verifiedHoldingCount"
            ],

        "exactSignatureMatch":
            True,

        "verifiedAtUtc":
            utc_now_iso(),
    }

    save_json(
        directory
        /
        "top_holdings.json",
        clean_result,
    )

    save_json(
        directory
        /
        "recovery_diagnostics.json",
        {
            "recoveryEngine":
                result.get(
                    "recoveryEngine"
                ),

            "holdingsParser":
                result.get(
                    "holdingsParser"
                ),

            "primaryParserError":
                result.get(
                    "primaryParserError"
                ),

            "fallbackParserError":
                result.get(
                    "fallbackParserError"
                ),

            "spatialRecoveryError":
                result.get(
                    "spatialRecoveryError"
                ),

            "spatialDiagnostics":
                result.get(
                    "spatialDiagnostics"
                ),

            "verification":
                {
                    "verifiedParser":
                        verification[
                            "verifiedParser"
                        ],

                    "exactSignatureMatch":
                        True,
                },

            "savedAtUtc":
                utc_now_iso(),
        },
    )

    save_json(
        directory
        /
        "metadata.json",
        {
            "excelRow":
                result.get(
                    "excelRow"
                ),

            "fundName":
                result.get(
                    "fundName"
                ),

            "factsheetUrl":
                result.get(
                    "factsheetUrl"
                ),

            "factsheetDocumentDate":
                result.get(
                    "factsheetDocumentDate"
                ),

            "factsheetDataAsAt":
                result.get(
                    "factsheetDataAsAt"
                ),

            "topHoldingsCount":
                result.get(
                    "topHoldingsCount"
                ),

            "holdingsParser":
                result.get(
                    "holdingsParser"
                ),

            "verificationParser":
                verification.get(
                    "verifiedParser"
                ),

            "exactVerification":
                True,

            "savedAtUtc":
                utc_now_iso(),
        },
    )

    return directory


def save_no_holdings_section(
    result: dict,
) -> Path:

    directory = recovery_directory(
        result
    )

    directory.mkdir(
        parents=True,
        exist_ok=True,
    )

    clean_result = {
        key:
            value
        for key, value
        in result.items()
        if not key.startswith("_")
    }

    clean_result[
        "recoveryEngine"
    ] = "no_holdings_section"

    save_json(
        directory
        /
        "top_holdings.json",
        clean_result,
    )

    (
        directory
        /
        "factsheet_text.txt"
    ).write_text(
        result.get(
            "_fullText"
        )
        or
        "",
        encoding="utf-8",
    )

    save_json(
        directory
        /
        "metadata.json",
        {
            "excelRow":
                result.get(
                    "excelRow"
                ),

            "fundName":
                result.get(
                    "fundName"
                ),

            "factsheetUrl":
                result.get(
                    "factsheetUrl"
                ),

            "status":
                "no_holdings_section",

            "savedAtUtc":
                utc_now_iso(),
        },
    )

    return directory


def save_failure(
    excel_fund: dict,
    error_text: str,
    attempt_diagnostics: list[dict],
    section_text: str | None = None,
    full_text: str | None = None,
    diagnostics: dict | None = None,
) -> Path:

    excel_row = int(
        excel_fund[
            "excelRow"
        ]
    )

    directory = (
        RECOVERY_FUNDS_OUTPUT_DIR
        /
        f"{excel_row}_failed"
    )

    directory.mkdir(
        parents=True,
        exist_ok=True,
    )

    failure = {
        "status":
            "failed",

        "excelRow":
            excel_row,

        "prudentialUrl":
            excel_fund[
                "prudentialUrl"
            ],

        "excelPruAccessName":
            excel_fund.get(
                "pruAccessName"
            ),

        "error":
            clean_text(
                error_text
            ),

        "attempts":
            attempt_diagnostics,

        "failedAtUtc":
            utc_now_iso(),
    }

    save_json(
        directory
        /
        "failure.json",
        failure,
    )

    if section_text:

        (
            directory
            /
            "top_holdings_section.txt"
        ).write_text(
            section_text,
            encoding="utf-8",
        )

    if full_text:

        (
            directory
            /
            "factsheet_text.txt"
        ).write_text(
            full_text,
            encoding="utf-8",
        )

    if diagnostics:

        save_json(
            directory
            /
            "recovery_diagnostics.json",
            diagnostics,
        )

    return directory


# =============================================================================
# MAIN
# =============================================================================

def main() -> int:

    RECOVERY_OUTPUT_DIR.mkdir(
        parents=True,
        exist_ok=True,
    )

    RECOVERY_FUNDS_OUTPUT_DIR.mkdir(
        parents=True,
        exist_ok=True,
    )

    started_at = utc_now_iso()

    print(
        "\n"
        + "#" * 72
    )

    print(
        "VGRAT FMS - PRUDENTIAL TOP HOLDINGS RECOVERY 2"
    )

    print(
        "#" * 72
    )

    print(
        f"Started UTC: {started_at}"
    )

    # -------------------------------------------------------------------------
    # Load Recovery 1 source
    # -------------------------------------------------------------------------

    recovery_1 = load_recovery_1_run_summary()

    print(
        "\nRECOVERY 1 SOURCE"
    )

    print(
        f"Recovery 1 status: "
        f"{recovery_1.get('status')}"
    )

    print(
        f"Recovery 1 recorded universe: "
        f"{recovery_1.get('excelFundUniverse')}"
    )

    print(
        f"Recovery 1 still-failed funds: "
        f"{recovery_1.get('failedFunds')}"
    )

    # -------------------------------------------------------------------------
    # Load Excel
    # -------------------------------------------------------------------------

    excel_funds = read_excel_funds()

    recovery_universe = build_recovery_universe(
        recovery_1,
        excel_funds,
    )

    print(
        f"\nRecovery universe: "
        f"{len(recovery_universe)} failed fund(s)"
    )

    if not recovery_universe:

        completed_at = utc_now_iso()

        summary = {
            "status":
                "nothing_to_recover",

            "startedAtUtc":
                started_at,

            "completedAtUtc":
                completed_at,

            "recovery1RunSummary":
                str(
                    RECOVERY_1_RUN_SUMMARY_FILE
                ),

            "excelFile":
                str(
                    EXCEL_FILE
                ),

            "recovery1StillFailedFunds":
                int(
                    recovery_1.get(
                        "failedFunds",
                        0,
                    )
                    or
                    0
                ),

            "recoveryUniverse":
                0,

            "recoveredFunds":
                0,

            "noHoldingsSectionFunds":
                0,

            "failedFunds":
                0,

            "rules": {
                "recovery1StillFailedFundsOnly":
                    True,

                "officialPrudentialOnly":
                    True,

                "noThirdPartyHoldings":
                    True,

                "noInferredHoldings":
                    True,

                "noFabricatedPercentages":
                    True,

                "exactFinalVerification":
                    True,

                "recovery3Excluded":
                    True,
            },
        }

        save_json(
            RECOVERY_RUN_SUMMARY_FILE,
            summary,
        )

        print(
            "\nNothing to recover."
        )

        return 0

    # -------------------------------------------------------------------------
    # Recovery state
    # -------------------------------------------------------------------------

    recovered = []

    no_holdings_section = []

    failed = []

    fallback_recovered = []

    spatial_recovered = []

    for index, excel_fund in enumerate(
        recovery_universe,
        start=1,
    ):

        excel_row = int(
            excel_fund[
                "excelRow"
            ]
        )

        print(
            "\n"
            + "=" * 78
        )

        print(
            f"RECOVERY FUND {index}/{len(recovery_universe)}"
        )

        print(
            f"Excel row: {excel_row}"
        )

        print(
            "=" * 78
        )

        attempt_diagnostics = []

        final_result = None

        last_exception = None

        for attempt in range(
            1,
            RETRY_COUNT + 1,
        ):

            print(
                f"\nAttempt {attempt}/{RETRY_COUNT}"
            )

            attempt_record = {
                "attempt":
                    attempt,

                "startedAtUtc":
                    utc_now_iso(),
            }

            try:

                with sync_playwright() as playwright:

                    browser = playwright.chromium.launch(
                        headless=BROWSER_HEADLESS
                    )

                    context = browser.new_context(
                        viewport={
                            "width": 1440,
                            "height": 1000,
                        },

                        user_agent=(
                            "Mozilla/5.0 "
                            "(Windows NT 10.0; Win64; x64) "
                            "AppleWebKit/537.36 "
                            "(KHTML, like Gecko) "
                            "Chrome/153.0.0.0 Safari/537.36"
                        ),
                    )

                    page = context.new_page()

                    try:

                        result = recover_single_fund(
                            page,
                            excel_fund,
                        )

                        if (
                            result[
                                "status"
                            ]
                            ==
                            "no_holdings_section"
                        ):

                            directory = (
                                save_no_holdings_section(
                                    result
                                )
                            )

                            result[
                                "outputDirectory"
                            ] = str(
                                directory
                            )

                            final_result = result

                        else:

                            # -------------------------------------------------
                            # Exact final verification
                            # -------------------------------------------------

                            verification = (
                                verify_exact_against_official_pdf(
                                    page,
                                    result,
                                )
                            )

                            directory = save_success(
                                result,
                                verification,
                            )

                            result[
                                "outputDirectory"
                            ] = str(
                                directory
                            )

                            result[
                                "exactVerification"
                            ] = True

                            result[
                                "verificationParser"
                            ] = verification[
                                "verifiedParser"
                            ]

                            final_result = result

                    finally:

                        context.close()

                        browser.close()

                attempt_record[
                    "status"
                ] = "success"

                attempt_record[
                    "completedAtUtc"
                ] = utc_now_iso()

                if final_result:

                    attempt_record[
                        "parser"
                    ] = final_result.get(
                        "holdingsParser"
                    )

                    attempt_record[
                        "status"
                    ] = final_result.get(
                        "status"
                    )

                attempt_diagnostics.append(
                    attempt_record
                )

                last_exception = None

                break

            except Exception as error:

                last_exception = error

                error_text = clean_text(
                    str(error)
                )

                attempt_record[
                    "status"
                ] = "failed"

                attempt_record[
                    "error"
                ] = error_text

                attempt_record[
                    "completedAtUtc"
                ] = utc_now_iso()

                if isinstance(
                    error,
                    HoldingsParseFailure,
                ):

                    attempt_record[
                        "diagnostics"
                    ] = error.diagnostics

                attempt_diagnostics.append(
                    attempt_record
                )

                print(
                    f"Attempt failed: {error_text}"
                )

                if attempt < RETRY_COUNT:

                    print(
                        f"Retrying in "
                        f"{RETRY_DELAY_SECONDS} seconds..."
                    )

                    time.sleep(
                        RETRY_DELAY_SECONDS
                    )

        # ---------------------------------------------------------------------
        # No successful recovery
        # ---------------------------------------------------------------------

        if final_result is None:

            section_text = getattr(
                last_exception,
                "section_text",
                None,
            )

            full_text = getattr(
                last_exception,
                "full_text",
                None,
            )

            diagnostics = getattr(
                last_exception,
                "diagnostics",
                None,
            )

            directory = save_failure(
                excel_fund,
                clean_text(
                    str(
                        last_exception
                    )
                    if last_exception
                    else
                    "Unknown recovery failure."
                ),
                attempt_diagnostics,
                section_text=section_text,
                full_text=full_text,
                diagnostics=diagnostics,
            )

            failed.append(
                {
                    "excelRow":
                        excel_row,

                    "prudentialUrl":
                        excel_fund[
                            "prudentialUrl"
                        ],

                    "pruAccessName":
                        excel_fund.get(
                            "pruAccessName"
                        ),

                    "error":
                        clean_text(
                            str(
                                last_exception
                            )
                            if last_exception
                            else
                            "Unknown recovery failure."
                        ),

                    "outputDirectory":
                        str(
                            directory
                        ),

                    "attempts":
                        attempt_diagnostics,
                }
            )

            continue

        # ---------------------------------------------------------------------
        # No holdings section
        # ---------------------------------------------------------------------

        if (
            final_result[
                "status"
            ]
            ==
            "no_holdings_section"
        ):

            no_holdings_section.append(
                final_result
            )

            print(
                "\nRECOVERY RESULT: "
                "NO TOP HOLDINGS SECTION"
            )

            continue

        # ---------------------------------------------------------------------
        # Successful recovery
        # ---------------------------------------------------------------------

        recovered.append(
            final_result
        )

        parser = final_result.get(
            "holdingsParser"
        )

        if parser == "fallback":

            fallback_recovered.append(
                final_result
            )

        elif parser == "spatial_fallback":

            spatial_recovered.append(
                final_result
            )

        print(
            "\nRECOVERY SUCCESS"
        )

        print(
            f"Fund: "
            f"{final_result.get('fundName') or '-'}"
        )

        print(
            f"Parser: "
            f"{parser or '-'}"
        )

        print(
            f"Holdings: "
            f"{final_result.get('topHoldingsCount')}"
        )

        print(
            "Exact final verification: PASS"
        )

        for holding in final_result[
            "topHoldings"
        ]:

            print(
                f"  {holding['rank']}. "
                f"{holding['name']} - "
                f"{holding['weightText']}"
            )

    # =========================================================================
    # CONSOLIDATED SUMMARY
    # =========================================================================

    completed_at = utc_now_iso()

    total_holdings = sum(
        int(
            result.get(
                "topHoldingsCount",
                0,
            )
            or
            0
        )
        for result
        in recovered
    )

    summary = {
        "status":
            (
                "success"
                if not failed
                else
                "partial"
            ),

        "startedAtUtc":
            started_at,

        "completedAtUtc":
            completed_at,

        "recovery1RunSummary":
            str(
                RECOVERY_1_RUN_SUMMARY_FILE
            ),

        "excelFile":
            str(
                EXCEL_FILE
            ),

        "recovery1ExcelFundUniverse":
            recovery_1.get(
                "excelFundUniverse"
            ),

        "recovery1StillFailedFunds":
            recovery_1.get(
                "failedFunds"
            ),

        "recoveryUniverse":
            len(
                recovery_universe
            ),

        "recoveredFunds":
            len(
                recovered
            ),

        "noHoldingsSectionFunds":
            len(
                no_holdings_section
            ),

        "failedFunds":
            len(
                failed
            ),

        "fallbackRecoveredFunds":
            len(
                fallback_recovered
            ),

        "spatialRecoveredFunds":
            len(
                spatial_recovered
            ),

        "totalPublishedTopHoldingsRecovered":
            total_holdings,

        "recoveredFundsDetail":
            [
                {
                    "excelRow":
                        result.get(
                            "excelRow"
                        ),

                    "fundName":
                        result.get(
                            "fundName"
                        ),

                    "factsheetUrl":
                        result.get(
                            "factsheetUrl"
                        ),

                    "factsheetDocumentDate":
                        result.get(
                            "factsheetDocumentDate"
                        ),

                    "factsheetDataAsAt":
                        result.get(
                            "factsheetDataAsAt"
                        ),

                    "topHoldingsCount":
                        result.get(
                            "topHoldingsCount"
                        ),

                    "holdingsParser":
                        result.get(
                            "holdingsParser"
                        ),

                    "exactVerification":
                        result.get(
                            "exactVerification"
                        ),

                    "verificationParser":
                        result.get(
                            "verificationParser"
                        ),

                    "outputDirectory":
                        result.get(
                            "outputDirectory"
                        ),
                }

                for result
                in recovered
            ],

        "fallbackRecoveredFundsDetail":
            [
                {
                    "excelRow":
                        result.get(
                            "excelRow"
                        ),

                    "fundName":
                        result.get(
                            "fundName"
                        ),

                    "topHoldingsCount":
                        result.get(
                            "topHoldingsCount"
                        ),

                    "primaryParserError":
                        result.get(
                            "primaryParserError"
                        ),
                }

                for result
                in fallback_recovered
            ],

        "spatialRecoveredFundsDetail":
            [
                {
                    "excelRow":
                        result.get(
                            "excelRow"
                        ),

                    "fundName":
                        result.get(
                            "fundName"
                        ),

                    "topHoldingsCount":
                        result.get(
                            "topHoldingsCount"
                        ),

                    "fallbackParserError":
                        result.get(
                            "fallbackParserError"
                        ),

                    "spatialDiagnostics":
                        result.get(
                            "spatialDiagnostics"
                        ),
                }

                for result
                in spatial_recovered
            ],

        "noHoldingsSectionDetail":
            [
                {
                    "excelRow":
                        result.get(
                            "excelRow"
                        ),

                    "fundName":
                        result.get(
                            "fundName"
                        ),

                    "factsheetUrl":
                        result.get(
                            "factsheetUrl"
                        ),
                }

                for result
                in no_holdings_section
            ],

        "failedFundsDetail":
            failed,

        "recovery1FailureField":
            "stillFailedFundsDetail",

        "recovery1FailedRows":
            [
                int(item["excelRow"])
                for item in recovery_universe
            ],

        "rules":
            {
                "recovery1StillFailedFundsOnly":
                    True,

                "excelColumnAControlsMasterUniverseForSelectedRows":
                    True,

                "officialPrudentialSingaporeOnly":
                    True,

                "officialFactsheetOnly":
                    True,

                "thirdPartyHoldings":
                    False,

                "noInferredHoldings":
                    True,

                "noFabricatedHoldings":
                    True,

                "noFabricatedPercentages":
                    True,

                "noForcedTenEntries":
                    True,

                "fewerThanTenPublishedHoldingsAllowed":
                    True,

                "duplicateHoldingNamesAllowed":
                    True,

                "duplicateHoldingPercentagesAllowed":
                    True,

                "multilineHoldingNamesSupported":
                    True,

                "hyphenatedLineWrapReconstruction":
                    True,

                "fixedIncomeCouponPercentagesSupported":
                    True,

                "primaryParser":
                    True,

                "fallbackParser":
                    True,

                "spatialRecovery":
                    True,

                "spatialRecoveryUsesFallbackParser":
                    True,

                "noCoordinateProximityPairing":
                    True,

                "noFuzzyHoldingMatching":
                    True,

                "noSyntheticValues":
                    True,

                "exactFinalPdfVerification":
                    True,

                "exactRankNameWeightSignatureRequired":
                    True,

                "recovery3Excluded":
                    True,

                "recovery3RunsSeparately":
                    True,
            },
    }

    save_json(
        RECOVERY_RUN_SUMMARY_FILE,
        summary,
    )

    # =========================================================================
    # CONSOLE SUMMARY
    # =========================================================================

    print(
        "\n\n"
        + "=" * 78
    )

    print(
        "PRUDENTIAL TOP HOLDINGS RECOVERY 2 COMPLETE"
    )

    print(
        "=" * 78
    )

    print(
        f"Recovery 1 recorded universe: "
        f"{recovery_1.get('excelFundUniverse')}"
    )

    print(
        f"Recovery 1 still-failed funds: "
        f"{recovery_1.get('failedFunds')}"
    )

    print(
        f"Recovery universe: "
        f"{len(recovery_universe)}"
    )

    print(
        f"Recovered: "
        f"{len(recovered)}"
    )

    print(
        f"No Top Holdings section: "
        f"{len(no_holdings_section)}"
    )

    print(
        f"Still failed: "
        f"{len(failed)}"
    )

    print(
        f"Recovered by fallback parser: "
        f"{len(fallback_recovered)}"
    )

    print(
        f"Recovered by spatial fallback: "
        f"{len(spatial_recovered)}"
    )

    print(
        f"Total published holdings recovered: "
        f"{total_holdings}"
    )

    print(
        "\nRecovery 2 order:"
    )

    print(
        "  1. Normal PDF extraction"
    )

    print(
        "  2. Primary parser"
    )

    print(
        "  3. Fallback parser"
    )

    print(
        "  4. Spatial section reconstruction"
    )

    print(
        "  5. Fallback parser on reconstructed section"
    )

    print(
        "  6. Exact official PDF re-download"
    )

    print(
        "  7. Exact rank/name/weight verification"
    )

    print(
        "\nRecovery 3 is NOT executed by this script."
    )

    print(
        "Recovery 3 must consume the final failed funds from "
        "Recovery 2 separately."
    )

    print(
        "\nNo fuzzy matching."
    )

    print(
        "No proximity name/weight pairing."
    )

    print(
        "No fabricated holdings."
    )

    print(
        "No fabricated percentages."
    )

    print(
        "No third-party holdings."
    )

    print(
        "\nOutput:"
    )

    print(
        f" - {RECOVERY_RUN_SUMMARY_FILE}"
    )

    print(
        f" - {RECOVERY_FUNDS_OUTPUT_DIR}"
    )

    if failed:

        print(
            "\nSTILL FAILED:"
        )

        for item in failed:

            print(
                f" - Row "
                f"{item['excelRow']}: "
                f"{item['error']}"
            )

    print(
        "\nDone."
    )

    # Deliberately return zero even when some recovery funds remain failed.
    # GitHub Actions should preserve and upload diagnostics instead of
    # discarding the recovery artifacts because of a non-zero process exit.
    return 0


# =============================================================================
# ENTRY POINT
# =============================================================================

if __name__ == "__main__":

    try:

        raise SystemExit(
            main()
        )

    except KeyboardInterrupt:

        print(
            "\nInterrupted.",
            file=sys.stderr,
        )

        raise SystemExit(
            130
        )

    except Exception as error:

        print(
            "\nFATAL RECOVERY ERROR:",
            file=sys.stderr,
        )

        print(
            str(error),
            file=sys.stderr,
        )

        raise SystemExit(
            1
        )
