#!/usr/bin/env python3

"""
VGrat FMS - Prudential Dividend ALL-FUND TEST EXTRACTOR
========================================================

MASTER SOURCE
=============

Funds Links.xlsm

Column A:
    Prudential fund URL

Column B:
    Exact PruAccess fund name


PURPOSE
=======

Test official Prudential dividend extraction for EVERY fund listed
in Excel Column A.

Workflow:

    Funds Links.xlsm
          |
          v
    Read every populated URL
          |
          v
    Open official Prudential fund page
          |
          v
    Capture official ilpfunds.json
          |
          v
    Find dividendRate
          |
          v
    Parse published date=value pairs
          |
          v
    Select latest published dividend
          |
          v
    Save dividend record

IMPORTANT RULES
===============

1. Only official Prudential ilpfunds.json data is used.

2. dividendRate is the authoritative dividend history field.

3. dividendRate has the format:

       YYYYMMDD=value&YYYYMMDD=value&...

4. The latest dividend is determined by the latest DATE,
   not by assuming the JSON string is already ordered.

5. The dividend unit is taken directly from Prudential:

       "%"
       "cent per unit"
       etc.

6. No conversion is performed.

7. No dividend yield is calculated.

8. No dividend is inferred from BID/offer prices.

9. No synthetic, estimated, interpolated, or calculated values.

10. If a fund has no usable dividendRate, NO DIVIDEND RECORD
    is created for that fund.

11. The fund remains represented in the run summary so that
    the result can distinguish:

       - dividend found
       - no dividend
       - extraction failed

12. Funds Links.xlsm Column A controls the fund universe.
    No hardcoded fund count is used.

OUTPUT
======

output_dividend_test/
    dividend_records.json
    no_dividend.json
    failed_funds.json
    run_summary.json
    raw/
        <fund number>/
            response_1.json
            response_2.json
            ...
            rendered.html
            visible_text.txt
            requests.json
            summary.json

The main result is:

    output_dividend_test/dividend_records.json

Only funds with an actual published dividendRate appear
in dividend_records.json.
"""

import asyncio
import json
import re
import sys
from datetime import datetime
from pathlib import Path
from urllib.parse import urljoin

from openpyxl import load_workbook
from playwright.async_api import async_playwright, TimeoutError as PlaywrightTimeoutError


# ============================================================
# CONFIGURATION
# ============================================================

EXCEL_FILE = Path("Funds Links.xlsm")

OUTPUT_DIR = Path("output_dividend_test")
RAW_DIR = OUTPUT_DIR / "raw"

DIVIDEND_RECORDS_FILE = OUTPUT_DIR / "dividend_records.json"
NO_DIVIDEND_FILE = OUTPUT_DIR / "no_dividend.json"
FAILED_FUNDS_FILE = OUTPUT_DIR / "failed_funds.json"
RUN_SUMMARY_FILE = OUTPUT_DIR / "run_summary.json"

# Official Prudential JSON endpoint(s) of interest.
IMPORTANT_PARTS = [
    "ilpfunds.json",
]

# Browser settings
PAGE_TIMEOUT_MS = 60000
NETWORK_WAIT_MS = 8000

# Small delay between funds to reduce unnecessary pressure
# on the Prudential website.
DELAY_BETWEEN_FUNDS_SECONDS = 1.0


# ============================================================
# HELPERS
# ============================================================

def safe_filename(value: str, max_length: int = 150) -> str:
    """
    Convert arbitrary text into a filesystem-safe filename.
    """
    value = str(value or "").strip()

    if not value:
        value = "unknown"

    value = re.sub(r'[<>:"/\\|?*]', "_", value)
    value = re.sub(r"\s+", " ", value)
    value = value.strip(" .")

    if not value:
        value = "unknown"

    return value[:max_length]


def normalize_url(url: str) -> str:
    """
    Normalize an Excel URL into a clean string.
    """
    if url is None:
        return ""

    url = str(url).strip()

    if not url:
        return ""

    return url


def load_funds_from_excel():
    """
    Read every populated URL from Column A and the corresponding
    PruAccess fund name from Column B.

    Column A controls the universe.
    """

    if not EXCEL_FILE.exists():
        raise FileNotFoundError(
            f"Excel file not found: {EXCEL_FILE.resolve()}"
        )

    workbook = load_workbook(
        filename=EXCEL_FILE,
        read_only=True,
        data_only=True,
    )

    try:
        worksheet = workbook.active

        funds = []

        for row_number in range(2, worksheet.max_row + 1):

            url_value = worksheet.cell(
                row=row_number,
                column=1,
            ).value

            name_value = worksheet.cell(
                row=row_number,
                column=2,
            ).value

            url = normalize_url(url_value)

            # Column A controls the universe.
            if not url:
                continue

            fund_name = (
                str(name_value).strip()
                if name_value is not None
                else ""
            )

            funds.append(
                {
                    "excelRow": row_number,
                    "prudentialUrl": url,
                    "excelPruAccessName": fund_name,
                }
            )

        return funds

    finally:
        workbook.close()


def save_json(path: Path, data):
    """
    Save JSON using UTF-8 and readable formatting.
    """
    path.parent.mkdir(parents=True, exist_ok=True)

    with path.open(
        "w",
        encoding="utf-8",
    ) as file:
        json.dump(
            data,
            file,
            indent=2,
            ensure_ascii=False,
        )


def parse_dividend_rate_string(dividend_rate):
    """
    Parse Prudential dividendRate.

    Example:

        20260814=0.51&20260715=0.51&20260615=0.51

    Returns a list of:

        {
            "date": "2026-08-14",
            "rawDate": "20260814",
            "rate": "0.51"
        }

    Only valid YYYYMMDD=value entries are accepted.

    No calculations or conversions are performed.
    """

    if dividend_rate is None:
        return []

    if not isinstance(dividend_rate, str):
        return []

    dividend_rate = dividend_rate.strip()

    if not dividend_rate:
        return []

    records = []

    for item in dividend_rate.split("&"):

        item = item.strip()

        if not item:
            continue

        if "=" not in item:
            continue

        raw_date, rate = item.split("=", 1)

        raw_date = raw_date.strip()
        rate = rate.strip()

        if not re.fullmatch(r"\d{8}", raw_date):
            continue

        if not rate:
            continue

        try:
            parsed_date = datetime.strptime(
                raw_date,
                "%Y%m%d",
            )
        except ValueError:
            continue

        records.append(
            {
                "date": parsed_date.strftime("%Y-%m-%d"),
                "rawDate": raw_date,
                "rate": rate,
            }
        )

    return records


def find_dividend_information(data):
    """
    Recursively search the captured Prudential JSON for
    dividendRate.

    The official structure observed is:

        $[0].dividendRate
        $[0].dividendUnit
        $[0].hasDividend
        $[0].payoutFrequency

    We search recursively rather than hardcoding the entire
    JSON path so the extractor remains tolerant of wrapper
    structures.

    Returns the best usable dividend information or None.
    """

    candidates = []

    def walk(value, path="$"):

        if isinstance(value, dict):

            for key, child in value.items():

                child_path = f"{path}.{key}"

                if key == "dividendRate":

                    parsed = parse_dividend_rate_string(child)

                    if parsed:
                        candidates.append(
                            {
                                "path": child_path,
                                "dividendRate": child,
                                "records": parsed,
                                "parent": value,
                            }
                        )

                walk(child, child_path)

        elif isinstance(value, list):

            for index, child in enumerate(value):

                child_path = f"{path}[{index}]"

                walk(child, child_path)

    walk(data)

    if not candidates:
        return None

    # Prefer the candidate containing the greatest number of
    # valid published dividend observations.
    candidates.sort(
        key=lambda item: len(item["records"]),
        reverse=True,
    )

    candidate = candidates[0]

    records = candidate["records"]

    # Determine latest date by actual date.
    latest = max(
        records,
        key=lambda item: item["rawDate"],
    )

    parent = candidate["parent"]

    dividend_unit = parent.get("dividendUnit")

    if dividend_unit is not None:
        dividend_unit = str(dividend_unit).strip()

    if not dividend_unit:
        dividend_unit = None

    has_dividend = parent.get("hasDividend")

    payout_frequency = parent.get("payoutFrequency")

    if payout_frequency is not None:
        payout_frequency = str(payout_frequency).strip()

    return {
        "latestDividendDate": latest["date"],
        "latestDividendRate": latest["rate"],
        "dividendUnit": dividend_unit,
        "dividendRatePath": candidate["path"],
        "publishedDividendCount": len(records),
        "hasDividend": has_dividend,
        "payoutFrequency": payout_frequency,
    }


def extract_dividend_from_responses(response_data):
    """
    Inspect captured JSON responses and return the first usable
    official dividend record.

    Only ilpfunds.json responses are passed into this function.
    """

    for item in response_data:

        data = item.get("data")

        if data is None:
            continue

        result = find_dividend_information(data)

        if result:
            return result

    return None


# ============================================================
# FUND TEST
# ============================================================

async def process_fund(
    browser,
    fund,
    fund_index,
    total_funds,
):
    """
    Process one fund.

    Returns:

        {
            "status": "dividend_found"
            ...
        }

    OR:

        {
            "status": "no_dividend"
            ...
        }

    OR:

        {
            "status": "failed"
            ...
        }
    """

    excel_row = fund["excelRow"]
    prudential_url = fund["prudentialUrl"]
    excel_pruaccess_name = fund["excelPruAccessName"]

    fund_dir_name = (
        f"{fund_index:03d}_row_{excel_row}_"
        f"{safe_filename(excel_pruaccess_name or 'unnamed')}"
    )

    fund_raw_dir = RAW_DIR / fund_dir_name
    fund_raw_dir.mkdir(
        parents=True,
        exist_ok=True,
    )

    print()
    print("=" * 80)
    print(
        f"[{fund_index}/{total_funds}] "
        f"Excel row {excel_row}"
    )
    print(
        f"Fund: {excel_pruaccess_name or '(Column B blank)'}"
    )
    print(f"URL:  {prudential_url}")
    print("=" * 80)

    captured_json = []
    captured_requests = []

    async def handle_response(response):

        url = response.url

        if not any(
            part in url
            for part in IMPORTANT_PARTS
        ):
            return

        print(
            f"  [JSON] {response.status} {url}"
        )

        request_record = {
            "url": url,
            "status": response.status,
            "method": response.request.method,
        }

        captured_requests.append(request_record)

        try:
            text = await response.text()

            response_number = len(captured_json) + 1

            response_filename = (
                f"response_{response_number}.json"
            )

            response_path = (
                fund_raw_dir / response_filename
            )

            response_path.write_text(
                text,
                encoding="utf-8",
            )

            try:
                parsed_json = json.loads(text)

            except json.JSONDecodeError:
                print(
                    "  [WARN] Response was not valid JSON"
                )

                return

            captured_json.append(
                {
                    "url": url,
                    "status": response.status,
                    "data": parsed_json,
                    "file": response_filename,
                }
            )

        except Exception as exc:
            print(
                f"  [WARN] Could not capture response: {exc}"
            )

    context = await browser.new_context()

    context.on(
        "response",
        lambda response: asyncio.create_task(
            handle_response(response)
        ),
    )

    page = await context.new_page()

    try:

        try:
            await page.goto(
                prudential_url,
                wait_until="domcontentloaded",
                timeout=PAGE_TIMEOUT_MS,
            )

        except PlaywrightTimeoutError:
            print(
                "  [WARN] Page navigation timed out; "
                "continuing with captured responses."
            )

        except Exception as exc:
            print(
                f"  [ERROR] Navigation failed: {exc}"
            )

            return {
                "status": "failed",
                "excelRow": excel_row,
                "prudentialUrl": prudential_url,
                "excelPruAccessName": excel_pruaccess_name,
                "error": f"Navigation failed: {exc}",
            }

        # Give the page time to make its API requests.
        await page.wait_for_timeout(
            NETWORK_WAIT_MS
        )

        # Allow outstanding response handlers to finish.
        await page.wait_for_timeout(1000)

        # ----------------------------------------------------
        # Save rendered page
        # ----------------------------------------------------

        try:
            html = await page.content()

            (
                fund_raw_dir / "rendered.html"
            ).write_text(
                html,
                encoding="utf-8",
            )

        except Exception as exc:
            print(
                f"  [WARN] Could not save rendered HTML: {exc}"
            )

        # ----------------------------------------------------
        # Save visible text
        # ----------------------------------------------------

        try:
            visible_text = await page.locator(
                "body"
            ).inner_text()

            (
                fund_raw_dir / "visible_text.txt"
            ).write_text(
                visible_text,
                encoding="utf-8",
            )

        except Exception as exc:
            print(
                f"  [WARN] Could not save visible text: {exc}"
            )

        # ----------------------------------------------------
        # Save requests
        # ----------------------------------------------------

        save_json(
            fund_raw_dir / "requests.json",
            captured_requests,
        )

        # ----------------------------------------------------
        # Wait for any final response tasks
        # ----------------------------------------------------

        await page.wait_for_timeout(1000)

        # ----------------------------------------------------
        # Extract dividend
        # ----------------------------------------------------

        dividend = extract_dividend_from_responses(
            captured_json
        )

        # ----------------------------------------------------
        # Save per-fund summary
        # ----------------------------------------------------

        summary = {
            "excelRow": excel_row,
            "prudentialUrl": prudential_url,
            "excelPruAccessName": excel_pruaccess_name,
            "capturedJsonResponses": len(captured_json),
            "capturedImportantRequests": len(
                captured_requests
            ),
            "dividend": dividend,
        }

        save_json(
            fund_raw_dir / "summary.json",
            summary,
        )

        # ----------------------------------------------------
        # Dividend found
        # ----------------------------------------------------

        if dividend:

            print()
            print("  >>> DIVIDEND FOUND")
            print(
                f"      Date: {dividend['latestDividendDate']}"
            )
            print(
                f"      Rate: {dividend['latestDividendRate']}"
            )
            print(
                f"      Unit: {dividend.get('dividendUnit')}"
            )
            print(
                f"      Published records: "
                f"{dividend['publishedDividendCount']}"
            )

            return {
                "status": "dividend_found",
                "excelRow": excel_row,
                "prudentialUrl": prudential_url,
                "excelPruAccessName": excel_pruaccess_name,
                "latestDividendDate": dividend[
                    "latestDividendDate"
                ],
                "latestDividendRate": dividend[
                    "latestDividendRate"
                ],
                "dividendUnit": dividend.get(
                    "dividendUnit"
                ),
                "publishedDividendCount": dividend[
                    "publishedDividendCount"
                ],
                "dividendRatePath": dividend[
                    "dividendRatePath"
                ],
                "hasDividend": dividend.get(
                    "hasDividend"
                ),
                "payoutFrequency": dividend.get(
                    "payoutFrequency"
                ),
            }

        # ----------------------------------------------------
        # No dividend
        # ----------------------------------------------------

        print()
        print("  >>> NO DIVIDEND RECORD")

        return {
            "status": "no_dividend",
            "excelRow": excel_row,
            "prudentialUrl": prudential_url,
            "excelPruAccessName": excel_pruaccess_name,
        }

    except Exception as exc:

        print()
        print(
            f"  [ERROR] Fund processing failed: {exc}"
        )

        return {
            "status": "failed",
            "excelRow": excel_row,
            "prudentialUrl": prudential_url,
            "excelPruAccessName": excel_pruaccess_name,
            "error": str(exc),
        }

    finally:

        await page.close()
        await context.close()


# ============================================================
# MAIN
# ============================================================

async def main():

    print("=" * 80)
    print("VGrat FMS - PRUDENTIAL DIVIDEND ALL-FUND TEST")
    print("=" * 80)

    # --------------------------------------------------------
    # Prepare output
    # --------------------------------------------------------

    OUTPUT_DIR.mkdir(
        parents=True,
        exist_ok=True,
    )

    RAW_DIR.mkdir(
        parents=True,
        exist_ok=True,
    )

    # --------------------------------------------------------
    # Load Excel
    # --------------------------------------------------------

    try:
        funds = load_funds_from_excel()

    except Exception as exc:

        print(
            f"[FATAL] Could not load Excel: {exc}"
        )

        sys.exit(1)

    total_funds = len(funds)

    print()
    print(
        f"Funds found in Column A: {total_funds}"
    )

    if total_funds == 0:

        print(
            "[FATAL] No populated fund URLs found "
            "in Excel Column A."
        )

        sys.exit(1)

    # --------------------------------------------------------
    # Results
    # --------------------------------------------------------

    dividend_records = []
    no_dividend_records = []
    failed_records = []

    started_at = datetime.now().isoformat()

    # --------------------------------------------------------
    # Browser
    # --------------------------------------------------------

    async with async_playwright() as playwright:

        browser = await playwright.chromium.launch(
            headless=True
        )

        try:

            for index, fund in enumerate(
                funds,
                start=1,
            ):

                result = await process_fund(
                    browser=browser,
                    fund=fund,
                    fund_index=index,
                    total_funds=total_funds,
                )

                status = result.get("status")

                if status == "dividend_found":

                    dividend_records.append(
                        result
                    )

                elif status == "no_dividend":

                    no_dividend_records.append(
                        result
                    )

                else:

                    failed_records.append(
                        result
                    )

                # ------------------------------------------------
                # Save incrementally after every fund.
                #
                # This prevents losing all progress if the workflow
                # is interrupted part-way through the run.
                # ------------------------------------------------

                save_json(
                    DIVIDEND_RECORDS_FILE,
                    dividend_records,
                )

                save_json(
                    NO_DIVIDEND_FILE,
                    no_dividend_records,
                )

                save_json(
                    FAILED_FUNDS_FILE,
                    failed_records,
                )

                current_summary = {
                    "status": "running",
                    "startedAt": started_at,
                    "updatedAt": datetime.now().isoformat(),
                    "excelFile": str(EXCEL_FILE),
                    "totalFunds": total_funds,
                    "processedFunds": (
                        len(dividend_records)
                        + len(no_dividend_records)
                        + len(failed_records)
                    ),
                    "dividendFunds": len(
                        dividend_records
                    ),
                    "noDividendFunds": len(
                        no_dividend_records
                    ),
                    "failedFunds": len(
                        failed_records
                    ),
                }

                save_json(
                    RUN_SUMMARY_FILE,
                    current_summary,
                )

                if (
                    index < total_funds
                    and DELAY_BETWEEN_FUNDS_SECONDS > 0
                ):
                    await asyncio.sleep(
                        DELAY_BETWEEN_FUNDS_SECONDS
                    )

        finally:

            await browser.close()

    # --------------------------------------------------------
    # Final summary
    # --------------------------------------------------------

    completed_at = datetime.now().isoformat()

    final_summary = {
        "status": (
            "completed"
            if not failed_records
            else "completed_with_failures"
        ),
        "startedAt": started_at,
        "completedAt": completed_at,
        "excelFile": str(EXCEL_FILE),
        "totalFunds": total_funds,
        "processedFunds": (
            len(dividend_records)
            + len(no_dividend_records)
            + len(failed_records)
        ),
        "dividendFunds": len(
            dividend_records
        ),
        "noDividendFunds": len(
            no_dividend_records
        ),
        "failedFunds": len(
            failed_records
        ),
        "outputs": {
            "dividendRecords": str(
                DIVIDEND_RECORDS_FILE
            ),
            "noDividend": str(
                NO_DIVIDEND_FILE
            ),
            "failedFunds": str(
                FAILED_FUNDS_FILE
            ),
            "runSummary": str(
                RUN_SUMMARY_FILE
            ),
            "rawDirectory": str(
                RAW_DIR
            ),
        },
    }

    save_json(
        RUN_SUMMARY_FILE,
        final_summary,
    )

    # --------------------------------------------------------
    # Console summary
    # --------------------------------------------------------

    print()
    print("=" * 80)
    print("RUN COMPLETE")
    print("=" * 80)

    print(
        f"Total funds:       {total_funds}"
    )

    print(
        f"Dividend found:    {len(dividend_records)}"
    )

    print(
        f"No dividend:       {len(no_dividend_records)}"
    )

    print(
        f"Failed:             {len(failed_records)}"
    )

    print()
    print(
        f"Dividend records:  {DIVIDEND_RECORDS_FILE}"
    )

    print(
        f"No dividend:       {NO_DIVIDEND_FILE}"
    )

    print(
        f"Failed funds:      {FAILED_FUNDS_FILE}"
    )

    print(
        f"Run summary:       {RUN_SUMMARY_FILE}"
    )

    print("=" * 80)


if __name__ == "__main__":
    asyncio.run(main())
