#!/usr/bin/env python3

"""
PruAccess All-Fund Historical BID Price Validation.

MASTER SOURCE
=============

Funds Links.xlsm

Column A:
    Prudential fund URL

Column B:
    PruAccess fund name


VALIDATION PRINCIPLES
=====================

This validator does NOT create or repair data.

It only validates data already extracted by test_pruaccess.py.

It checks:

1. Excel master universe.
2. Successful/failed fund accounting.
3. Fund identity.
4. Prudential source.
5. PruAccess source.
6. BID-only historical prices.
7. Observation counts.
8. Date validity.
9. Price validity.
10. Strict descending date order.
11. Duplicate observations.
12. Inception coverage.
13. Consecutive BID observation date gaps.
14. Pagination completeness.
15. Page sizes.
16. Consolidated history.
17. Current Prudential BID vs latest PruAccess BID.
18. Suspicious synthetic/generated fields.
19. Missing Excel funds.
20. Duplicate Excel rows.
21. Duplicate fund identifiers.

NO DATA IS FABRICATED.

A failed validation remains a failure.
"""

from __future__ import annotations

import json
import re
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook


# ============================================================
# CONFIG
# ============================================================

EXCEL_FILE = Path(
    "Funds Links.xlsm"
)

OUTPUT_DIR = Path(
    "output_pruaccess"
)

RUN_SUMMARY_FILE = (
    OUTPUT_DIR / "run_summary.json"
)

ALL_FUNDS_FILE = (
    OUTPUT_DIR / "all_funds.json"
)

ALL_BID_HISTORY_FILE = (
    OUTPUT_DIR / "all_bid_history.json"
)

FUNDS_OUTPUT_DIR = (
    OUTPUT_DIR / "funds"
)

VALIDATION_FILE = (
    OUTPUT_DIR / "validation.json"
)

PAGE_SIZE = 20

# Maximum allowed calendar-day gap between consecutive actual
# PruAccess BID observations.
#
# Gaps of 7 calendar days or less are allowed.
# A gap greater than 7 calendar days is a validation failure.
#
# No missing observations are created or filled.
MAX_BID_DATE_GAP_DAYS = 7

PRICE_TOLERANCE = 0.0002


# ============================================================
# HELPERS
# ============================================================

def clean_text(value) -> str:

    if value is None:

        return ""

    return re.sub(
        r"\s+",
        " ",
        str(value),
    ).strip()


def normalize_text(value) -> str:

    value = clean_text(
        value
    ).lower()

    value = value.replace(
        "–",
        "-"
    )

    value = value.replace(
        "—",
        "-"
    )

    value = re.sub(
        r"\s+",
        " ",
        value,
    )

    return value.strip()


def load_json(
    path: Path,
):

    if not path.exists():

        raise FileNotFoundError(
            f"Required JSON file not found: "
            f"{path}"
        )

    with path.open(
        "r",
        encoding="utf-8",
    ) as file:

        return json.load(
            file
        )


def save_json(
    path: Path,
    data,
) -> None:

    path.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    path.write_text(
        json.dumps(
            data,
            indent=2,
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )


def parse_price(value):

    if value is None:

        return None

    text = clean_text(
        value
    )

    text = text.replace(
        "$",
        ""
    )

    text = text.replace(
        ",",
        ""
    )

    if not text:

        return None

    try:

        return float(
            text
        )

    except ValueError:

        return None


def parse_date(value):

    if value is None:

        return None

    text = clean_text(
        value
    )

    formats = [
        "%d-%b-%Y",
        "%d/%m/%Y",
        "%Y-%m-%d",
    ]

    for date_format in formats:

        try:

            return datetime.strptime(
                text,
                date_format,
            )

        except ValueError:

            continue

    return None


def get_fund_directory(
    excel_row: int,
) -> Path | None:

    if not FUNDS_OUTPUT_DIR.exists():

        return None

    prefix = f"{excel_row}_"

    candidates = [
        path
        for path
        in FUNDS_OUTPUT_DIR.iterdir()
        if path.is_dir()
        and path.name.startswith(
            prefix
        )
        and "failed" not in path.name.lower()
    ]

    if not candidates:

        return None

    if len(candidates) > 1:

        raise RuntimeError(
            f"Multiple output directories "
            f"found for Excel row "
            f"{excel_row}: "
            f"{candidates}"
        )

    return candidates[0]


# ============================================================
# EXCEL MASTER UNIVERSE
# ============================================================

def read_excel_universe() -> list[dict]:

    if not EXCEL_FILE.exists():

        raise FileNotFoundError(
            f"Excel file not found: "
            f"{EXCEL_FILE}"
        )

    workbook = load_workbook(
        EXCEL_FILE,
        read_only=True,
        keep_vba=True,
        data_only=True,
    )

    worksheet = workbook.active

    funds = []

    for row_number in range(
        2,
        worksheet.max_row + 1,
    ):

        prudential_url = clean_text(
            worksheet.cell(
                row=row_number,
                column=1,
            ).value
        )

        pruaccess_name = clean_text(
            worksheet.cell(
                row=row_number,
                column=2,
            ).value
        )

        if not prudential_url:

            continue

        funds.append(
            {
                "excelRow":
                    row_number,

                "prudentialUrl":
                    prudential_url,

                "pruAccessName":
                    pruaccess_name,
            }
        )

    workbook.close()

    return funds


# ============================================================
# RECORD NORMALIZATION
# ============================================================

def looks_like_success_record(
    value,
) -> bool:

    if not isinstance(
        value,
        dict,
    ):

        return False

    keys = set(
        value.keys()
    )

    identity_keys = {
        "excelRow",
        "excel_row",
        "prudentialFundName",
        "fundName",
        "fundIdentifier",
        "fundCode",
        "pruAccessFundId",
        "observationCount",
    }

    return bool(
        keys.intersection(
            identity_keys
        )
    )


def normalize_success_record(
    value,
    inherited_identifier=None,
) -> dict | None:

    if not isinstance(
        value,
        dict,
    ):

        return None

    # --------------------------------------------------------
    # Some extractor versions store the useful record under
    # "summary".
    # --------------------------------------------------------

    if (
        "summary" in value
        and isinstance(
            value["summary"],
            dict,
        )
        and looks_like_success_record(
            value["summary"]
        )
    ):

        record = dict(
            value["summary"]
        )

        # Preserve outer identity fields.
        for key in [
            "excelRow",
            "prudentialUrl",
            "excelPruAccessName",
            "prudentialFundName",
            "fundIdentifier",
            "fundCode",
            "pruAccessFundId",
        ]:

            if (
                key not in record
                and key in value
            ):

                record[key] = value[key]

        return record

    record = dict(
        value
    )

    if (
        inherited_identifier
        and not record.get(
            "fundIdentifier"
        )
    ):

        record[
            "fundIdentifier"
        ] = inherited_identifier

    # --------------------------------------------------------
    # Normalize alternative naming.
    # --------------------------------------------------------

    if (
        "excelRow"
        not in record
        and "excel_row"
        in record
    ):

        record[
            "excelRow"
        ] = record[
            "excel_row"
        ]

    if (
        "fundName"
        not in record
        and "prudentialFundName"
        in record
    ):

        record[
            "fundName"
        ] = record[
            "prudentialFundName"
        ]

    if (
        "fundIdentifier"
        not in record
        and "citicode"
        in record
    ):

        record[
            "fundIdentifier"
        ] = record[
            "citicode"
        ]

    if (
        "fundIdentifier"
        not in record
        and "fundIdentifier"
        in record.get(
            "prudential",
            {},
        )
    ):

        record[
            "fundIdentifier"
        ] = record[
            "prudential"
        ][
            "fundIdentifier"
        ]

    if (
        "fundCode"
        not in record
        and isinstance(
            record.get(
                "prudential"
            ),
            dict,
        )
    ):

        record[
            "fundCode"
        ] = record[
            "prudential"
        ].get(
            "fundCode"
        )

    if (
        "fundName"
        not in record
        and isinstance(
            record.get(
                "prudential"
            ),
            dict,
        )
    ):

        record[
            "fundName"
        ] = record[
            "prudential"
        ].get(
            "fundName"
        )

    return record


def flatten_fund_records(
    value,
    inherited_identifier=None,
) -> list[dict]:

    records = []

    if isinstance(
        value,
        dict,
    ):

        # ----------------------------------------------------
        # A dictionary whose values are fund records.
        # ----------------------------------------------------

        if "funds" in value:

            records.extend(
                flatten_fund_records(
                    value["funds"]
                )
            )

        # ----------------------------------------------------
        # Known successful containers.
        # ----------------------------------------------------

        for key in [
            "successfulFunds",
            "successful",
            "success",
        ]:

            if key in value:

                records.extend(
                    flatten_fund_records(
                        value[key]
                    )
                )

        # ----------------------------------------------------
        # If this itself looks like a fund record, retain it.
        # ----------------------------------------------------

        record = normalize_success_record(
            value,
            inherited_identifier,
        )

        if (
            record is not None
            and looks_like_success_record(
                record
            )
        ):

            # Do not add observation dictionaries.
            if (
                "date" not in record
                or "bidPrice" not in record
            ):

                records.append(
                    record
                )

        # ----------------------------------------------------
        # Handle dictionaries keyed by fund identifier.
        # ----------------------------------------------------

        for key, child in value.items():

            if key in {
                "funds",
                "successfulFunds",
                "successful",
                "success",
                "summary",
            }:

                continue

            if isinstance(
                child,
                dict,
            ):

                if (
                    looks_like_success_record(
                        child
                    )
                    or "summary" in child
                    or "fundIdentifier" in child
                ):

                    records.extend(
                        flatten_fund_records(
                            child,
                            inherited_identifier=(
                                key
                                if not child.get(
                                    "fundIdentifier"
                                )
                                else None
                            ),
                        )
                    )

    elif isinstance(
        value,
        list,
    ):

        for item in value:

            if isinstance(
                item,
                dict,
            ):

                records.extend(
                    flatten_fund_records(
                        item,
                        inherited_identifier,
                    )
                )

            elif isinstance(
                item,
                list,
            ):

                records.extend(
                    flatten_fund_records(
                        item,
                        inherited_identifier,
                    )
                )

    return records


def deduplicate_records(
    records: list[dict],
) -> list[dict]:

    result = []

    seen = set()

    for record in records:

        excel_row = record.get(
            "excelRow"
        )

        identifier = clean_text(
            record.get(
                "fundIdentifier"
            )
        )

        fund_code = clean_text(
            record.get(
                "fundCode"
            )
        )

        fund_name = normalize_text(
            record.get(
                "fundName"
                or record.get(
                    "prudentialFundName"
                )
            )
        )

        key = (
            excel_row,
            identifier,
            fund_code,
            fund_name,
        )

        if key in seen:

            continue

        seen.add(
            key
        )

        result.append(
            record
        )

    return result


# ============================================================
# SUCCESSFUL FUND EXTRACTION
# ============================================================

def extract_successful_funds(
    run_summary,
    all_funds,
) -> list[dict]:

    candidates = []

    # --------------------------------------------------------
    # Prefer all_funds because it contains the complete fund
    # records.
    # --------------------------------------------------------

    if isinstance(
        all_funds,
        dict,
    ):

        candidates.extend(
            flatten_fund_records(
                all_funds.get(
                    "funds",
                    [],
                )
            )
        )

    # --------------------------------------------------------
    # Also inspect run_summary.
    # --------------------------------------------------------

    if isinstance(
        run_summary,
        dict,
    ):

        candidates.extend(
            flatten_fund_records(
                run_summary.get(
                    "successfulFunds",
                    [],
                )
            )
        )

        # Some older structures may put them under funds.
        candidates.extend(
            flatten_fund_records(
                run_summary.get(
                    "funds",
                    [],
                )
            )
        )

    records = deduplicate_records(
        candidates
    )

    # --------------------------------------------------------
    # Only records with an Excel row can represent a successful
    # Excel fund.
    # --------------------------------------------------------

    valid = []

    for record in records:

        if record.get(
            "excelRow"
        ) is None:

            continue

        try:

            record[
                "excelRow"
            ] = int(
                record[
                    "excelRow"
                ]
            )

        except (
            ValueError,
            TypeError,
        ):

            continue

        valid.append(
            record
        )

    # --------------------------------------------------------
    # One successful record per Excel row.
    # --------------------------------------------------------

    by_row = {}

    for record in valid:

        row = record[
            "excelRow"
        ]

        if row not in by_row:

            by_row[
                row
            ] = record

    return [
        by_row[row]
        for row
        in sorted(
            by_row
        )
    ]


# ============================================================
# FAILED FUND EXTRACTION
# ============================================================

def extract_failed_funds(
    run_summary,
    all_funds,
) -> list[dict]:

    candidates = []

    for container in [
        run_summary,
        all_funds,
    ]:

        if not isinstance(
            container,
            dict,
        ):

            continue

        for key in [
            "failedFunds",
            "failed",
            "failures",
        ]:

            value = container.get(
                key
            )

            if isinstance(
                value,
                list,
            ):

                for item in value:

                    if isinstance(
                        item,
                        dict,
                    ):

                        candidates.append(
                            item
                        )

    result = []

    seen_rows = set()

    for failure in candidates:

        row = failure.get(
            "excelRow"
        )

        try:

            row = int(
                row
            )

        except (
            ValueError,
            TypeError,
        ):

            continue

        if row in seen_rows:

            continue

        seen_rows.add(
            row
        )

        failure[
            "excelRow"
        ] = row

        result.append(
            failure
        )

    return sorted(
        result,
        key=lambda x: x[
            "excelRow"
        ],
    )


# ============================================================
# SUSPICIOUS DATA CHECK
# ============================================================

SUSPICIOUS_KEY_TERMS = [
    "synthetic",
    "interpolat",
    "simulat",
    "generated",
    "estimated",
    "illustrative",
    "mock",
    "placeholder",
    "carryforward",
    "carry-forward",
    "forecast",
]


def find_suspicious_fields(
    value,
    path="root",
) -> list[str]:

    found = []

    if isinstance(
        value,
        dict,
    ):

        for key, child in value.items():

            key_lower = str(
                key
            ).lower()

            if any(
                term in key_lower
                for term
                in SUSPICIOUS_KEY_TERMS
            ):

                found.append(
                    f"{path}.{key}"
                )

            found.extend(
                find_suspicious_fields(
                    child,
                    f"{path}.{key}",
                )
            )

    elif isinstance(
        value,
        list,
    ):

        for index, child in enumerate(
            value
        ):

            found.extend(
                find_suspicious_fields(
                    child,
                    f"{path}[{index}]",
                )
            )

    return found


# ============================================================
# FUND VALIDATION
# ============================================================

def validate_fund(
    excel_fund: dict,
    successful_summary: dict,
) -> dict:

    errors = []

    warnings = []

    excel_row = excel_fund[
        "excelRow"
    ]

    expected_prudential_url = (
        excel_fund[
            "prudentialUrl"
        ]
    )

    expected_pruaccess_name = (
        excel_fund[
            "pruAccessName"
        ]
    )

    fund_dir = get_fund_directory(
        excel_row
    )

    if fund_dir is None:

        errors.append(
            "Individual fund output directory "
            "not found."
        )

        return {
            "excelRow":
                excel_row,

            "status":
                "failed",

            "errors":
                errors,

            "warnings":
                warnings,
        }

    required_files = [
        "summary.json",
        "bid_history.json",
        "pagination.json",
        "prudential_fund.json",
    ]

    loaded = {}

    for filename in required_files:

        path = (
            fund_dir
            / filename
        )

        if not path.exists():

            errors.append(
                f"Missing required file: "
                f"{filename}"
            )

            continue

        try:

            loaded[
                filename
            ] = load_json(
                path
            )

        except Exception as error:

            errors.append(
                f"Could not load "
                f"{filename}: "
                f"{clean_text(str(error))}"
            )

    if errors:

        return {
            "excelRow":
                excel_row,

            "status":
                "failed",

            "errors":
                errors,

            "warnings":
                warnings,
        }

    summary = loaded[
        "summary.json"
    ]

    bid_history = loaded[
        "bid_history.json"
    ]

    pagination = loaded[
        "pagination.json"
    ]

    prudential = loaded[
        "prudential_fund.json"
    ]

    # ========================================================
    # STRUCTURE
    # ========================================================

    if not isinstance(
        summary,
        dict,
    ):

        errors.append(
            "summary.json is not an object."
        )

    if not isinstance(
        bid_history,
        dict,
    ):

        errors.append(
            "bid_history.json is not an object."
        )

    if not isinstance(
        pagination,
        list,
    ):

        errors.append(
            "pagination.json is not a list."
        )

    if not isinstance(
        prudential,
        dict,
    ):

        errors.append(
            "prudential_fund.json is not an object."
        )

    if errors:

        return {
            "excelRow":
                excel_row,

            "status":
                "failed",

            "errors":
                errors,

            "warnings":
                warnings,
        }

    # ========================================================
    # SUCCESS STATUS
    # ========================================================

    if normalize_text(
        summary.get(
            "status"
        )
    ) != "success":

        errors.append(
            "summary.json status is not success."
        )

    # ========================================================
    # SOURCE
    # ========================================================

    source = normalize_text(
        bid_history.get(
            "source"
        )
    )

    if source != "pruaccess":

        errors.append(
            "Historical source is not PruAccess."
        )

    price_type = normalize_text(
        bid_history.get(
            "priceType"
        )
    )

    if price_type != "bid":

        errors.append(
            "Historical price type is not BID."
        )

    # ========================================================
    # EXCEL IDENTITY
    # ========================================================

    summary_row = summary.get(
        "excelRow"
    )

    try:

        summary_row = int(
            summary_row
        )

    except (
        ValueError,
        TypeError,
    ):

        errors.append(
            "summary.json Excel row is invalid."
        )

        summary_row = None

    if summary_row != excel_row:

        errors.append(
            "Excel row mismatch."
        )

    # ========================================================
    # URL
    # ========================================================

    actual_url = clean_text(
        summary.get(
            "prudentialUrl"
        )
    )

    if actual_url != expected_prudential_url:

        errors.append(
            "Prudential URL does not match "
            "Funds Links.xlsm."
        )

    # ========================================================
    # PRUDENTIAL FUND IDENTITY
    # ========================================================

    prudential_name = clean_text(
        prudential.get(
            "fundName"
        )
    )

    summary_name = clean_text(
        summary.get(
            "prudentialFundName"
        )
    )

    if not prudential_name:

        errors.append(
            "Prudential fundName is missing."
        )

    if not summary_name:

        errors.append(
            "summary prudentialFundName is missing."
        )

    if (
        prudential_name
        and summary_name
        and normalize_text(
            prudential_name
        )
        != normalize_text(
            summary_name
        )
    ):

        errors.append(
            "Prudential fund name mismatch."
        )

    # ========================================================
    # FUND IDENTIFIER
    # ========================================================

    fund_identifier = clean_text(
        prudential.get(
            "fundIdentifier"
        )
    )

    summary_identifier = clean_text(
        summary.get(
            "fundIdentifier"
        )
    )

    if not fund_identifier:

        errors.append(
            "Prudential fundIdentifier missing."
        )

    if (
        fund_identifier
        and summary_identifier
        and fund_identifier
        != summary_identifier
    ):

        errors.append(
            "Fund identifier mismatch."
        )

    # ========================================================
    # FUND CODE
    # ========================================================

    fund_code = clean_text(
        prudential.get(
            "fundCode"
        )
    )

    summary_fund_code = clean_text(
        summary.get(
            "fundCode"
        )
    )

    if (
        fund_code
        and summary_fund_code
        and fund_code
        != summary_fund_code
    ):

        errors.append(
            "Fund code mismatch."
        )

    # ========================================================
    # PRUACCESS NAME
    # ========================================================

    actual_pruaccess_name = clean_text(
        summary.get(
            "matchedPruAccessName"
        )
    )

    if not actual_pruaccess_name:

        errors.append(
            "Matched PruAccess fund name missing."
        )

    if (
        expected_pruaccess_name
        and actual_pruaccess_name
        and normalize_text(
            expected_pruaccess_name
        )
        != normalize_text(
            actual_pruaccess_name
        )
    ):

        errors.append(
            "PruAccess fund name does not "
            "match Excel Column B."
        )

    # ========================================================
    # PRUACCESS ID
    # ========================================================

    pruaccess_id = clean_text(
        summary.get(
            "pruAccessFundId"
        )
    )

    if not pruaccess_id:

        errors.append(
            "PruAccess fund ID missing."
        )

    # ========================================================
    # BID HISTORY OBSERVATIONS
    # ========================================================

    observations = bid_history.get(
        "observations"
    )

    if not isinstance(
        observations,
        list,
    ):

        errors.append(
            "bid_history observations "
            "is not a list."
        )

        observations = []

    actual_count = len(
        observations
    )

    stored_count = bid_history.get(
        "observationCount"
    )

    try:

        stored_count = int(
            stored_count
        )

    except (
        ValueError,
        TypeError,
    ):

        errors.append(
            "bid_history observationCount "
            "is invalid."
        )

        stored_count = None

    if stored_count != actual_count:

        errors.append(
            "bid_history observationCount "
            "does not equal actual observation count."
        )

    summary_count = (
        successful_summary.get(
            "historicalObservationCount"
        )
        or successful_summary.get(
            "observationCount"
        )
        or summary.get(
            "historicalObservationCount"
        )
    )

    if summary_count is not None:

        try:

            summary_count = int(
                summary_count
            )

            if summary_count != actual_count:

                errors.append(
                    "Successful fund summary "
                    "observation count does not "
                    "match actual observations."
                )

        except (
            ValueError,
            TypeError,
        ):

            errors.append(
                "Successful fund summary "
                "observation count is invalid."
            )

    # ========================================================
    # OBSERVATION VALIDATION
    # ========================================================

    parsed_dates = []

    raw_date_strings = []

    date_price_pairs = []

    for index, observation in enumerate(
        observations
    ):

        if not isinstance(
            observation,
            dict,
        ):

            errors.append(
                f"Observation {index + 1} "
                "is not an object."
            )

            continue

        date_value = clean_text(
            observation.get(
                "date"
            )
        )

        bid_value = clean_text(
            observation.get(
                "bidPrice"
            )
        )

        if not date_value:

            errors.append(
                f"Observation {index + 1} "
                "date is missing."
            )

            continue

        if not bid_value:

            errors.append(
                f"Observation {index + 1} "
                "bidPrice is missing."
            )

            continue

        parsed_date = parse_date(
            date_value
        )

        if parsed_date is None:

            errors.append(
                f"Observation {index + 1} "
                f"has invalid date: "
                f"{date_value}"
            )

        else:

            parsed_dates.append(
                parsed_date
            )

        parsed_price = parse_price(
            bid_value
        )

        if (
            parsed_price is None
            or parsed_price <= 0
        ):

            errors.append(
                f"Observation {index + 1} "
                f"has invalid BID price: "
                f"{bid_value}"
            )

        raw_date_strings.append(
            date_value
        )

        date_price_pairs.append(
            (
                date_value,
                bid_value,
            )
        )

    # ========================================================
    # DUPLICATE DATES
    # ========================================================

    duplicate_dates = []

    seen_dates = set()

    for date_value in raw_date_strings:

        if date_value in seen_dates:

            duplicate_dates.append(
                date_value
            )

        seen_dates.add(
            date_value
        )

    if duplicate_dates:

        errors.append(
            "Duplicate historical dates found: "
            + ", ".join(
                sorted(
                    set(
                        duplicate_dates
                    )
                )
            )
        )

    # ========================================================
    # DUPLICATE DATE/PRICE PAIRS
    # ========================================================

    duplicate_pairs = []

    seen_pairs = set()

    for pair in date_price_pairs:

        if pair in seen_pairs:

            duplicate_pairs.append(
                pair
            )

        seen_pairs.add(
            pair
        )

    if duplicate_pairs:

        errors.append(
            "Duplicate date/BID observations found."
        )

    # ========================================================
    # STRICT DESCENDING ORDER
    # ========================================================

    if len(
        parsed_dates
    ) >= 2:

        for index in range(
            1,
            len(parsed_dates),
        ):

            previous = parsed_dates[
                index - 1
            ]

            current = parsed_dates[
                index
            ]

            if not (
                previous > current
            ):

                errors.append(
                    "Historical observations "
                    "are not strictly descending "
                    "by date."
                )

                break

    # ========================================================
    # CONSECUTIVE BID OBSERVATION DATE GAPS
    # ========================================================

    # Historical PruAccess observations do not have to be daily.
    #
    # Weekends, public holidays, non-valuation days, and other
    # valid source gaps are allowed.
    #
    # However, the gap between every two consecutive actual BID
    # observations must not exceed MAX_BID_DATE_GAP_DAYS.
    #
    # No missing observations are created.
    # No carry-forward values are inserted.
    # No interpolation is performed.

    excessive_date_gaps = []

    if len(
        parsed_dates
    ) >= 2:

        for index in range(
            1,
            len(parsed_dates),
        ):

            previous = parsed_dates[
                index - 1
            ]

            current = parsed_dates[
                index
            ]

            gap_days = (
                previous - current
            ).days

            if gap_days > MAX_BID_DATE_GAP_DAYS:

                excessive_date_gaps.append(
                    {
                        "fromDate":
                            previous.strftime(
                                "%Y-%m-%d"
                            ),

                        "toDate":
                            current.strftime(
                                "%Y-%m-%d"
                            ),

                        "gapDays":
                            gap_days,
                    }
                )

        if excessive_date_gaps:

            preview = (
                excessive_date_gaps[:20]
            )

            errors.append(
                "Historical BID observation date "
                "gap exceeds "
                f"{MAX_BID_DATE_GAP_DAYS} "
                "calendar days. "
                f"Found "
                f"{len(excessive_date_gaps)} "
                "excessive gap(s). "
                "First gap(s): "
                f"{json.dumps(preview, ensure_ascii=False)}"
            )

    # ========================================================
    # INCEPTION DATE
    # ========================================================

    inception_value = clean_text(
        prudential.get(
            "inceptionDate"
        )
    )

    inception_date = parse_date(
        inception_value
    )

    if inception_date is None:

        errors.append(
            "Prudential inception date "
            "is invalid."
        )

    elif parsed_dates:

        oldest_date = min(
            parsed_dates
        )

        if oldest_date > inception_date:

            errors.append(
                "Historical PruAccess data does "
                "not reach the Prudential "
                "inception date."
            )

    # ========================================================
    # PAGINATION
    # ========================================================

    expected_page_count = (
        len(pagination)
    )

    if expected_page_count <= 0:

        errors.append(
            "No pagination records found."
        )

    page_numbers = []

    successful_pages = []

    pagination_row_total = 0

    for index, page_record in enumerate(
        pagination
    ):

        if not isinstance(
            page_record,
            dict,
        ):

            errors.append(
                f"Pagination record "
                f"{index + 1} is not an object."
            )

            continue

        page_number = page_record.get(
            "page"
        )

        try:

            page_number = int(
                page_number
            )

        except (
            ValueError,
            TypeError,
        ):

            errors.append(
                f"Pagination record "
                f"{index + 1} has invalid page number."
            )

            continue

        page_numbers.append(
            page_number
        )

        status = normalize_text(
            page_record.get(
                "status"
            )
        )

        if status == "success":

            successful_pages.append(
                page_number
            )

        row_count = page_record.get(
            "rowCount"
        )

        try:

            row_count = int(
                row_count
            )

        except (
            ValueError,
            TypeError,
        ):

            errors.append(
                f"Page {page_number} "
                "has invalid rowCount."
            )

            continue

        if status == "success":

            pagination_row_total += (
                row_count
            )

            # All pages except the final successful
            # page should normally contain PAGE_SIZE.
            if (
                page_number
                < max(
                    successful_pages
                    or [page_number]
                )
                and row_count != PAGE_SIZE
            ):

                warnings.append(
                    f"Page {page_number} "
                    f"contains {row_count} rows."
                )

    if page_numbers:

        expected_sequence = list(
            range(
                1,
                max(
                    page_numbers
                ) + 1,
            )
        )

        if page_numbers != expected_sequence:

            errors.append(
                "Pagination page numbering "
                "is not continuous."
            )

    # --------------------------------------------------------
    # Last page must contain 1..PAGE_SIZE rows.
    # --------------------------------------------------------

    successful_page_records = [
        item
        for item
        in pagination
        if isinstance(
            item,
            dict,
        )
        and normalize_text(
            item.get(
                "status"
            )
        ) == "success"
    ]

    if successful_page_records:

        last_page = successful_page_records[-1]

        last_count = last_page.get(
            "rowCount"
        )

        try:

            last_count = int(
                last_count
            )

            if not (
                1
                <= last_count
                <= PAGE_SIZE
            ):

                errors.append(
                    "Final pagination page has "
                    f"invalid row count: "
                    f"{last_count}"
                )

        except (
            ValueError,
            TypeError,
        ):

            errors.append(
                "Final pagination page "
                "row count is invalid."
            )

    if (
        pagination_row_total
        != actual_count
    ):

        errors.append(
            "Pagination row total does not "
            "match historical observation count."
        )

    # ========================================================
    # CURRENT PRUDENTIAL BID VS LATEST PRUACCESS BID
    # ========================================================

    current_prudential_bid = parse_price(
        prudential.get(
            "bidPrice"
        )
    )

    latest_pruaccess_bid = None

    if observations:

        latest_pruaccess_bid = parse_price(
            observations[0].get(
                "bidPrice"
            )
        )

    if (
        current_prudential_bid is None
        or latest_pruaccess_bid is None
    ):

        errors.append(
            "Unable to compare current "
            "Prudential BID with latest "
            "PruAccess BID."
        )

    else:

        difference = abs(
            current_prudential_bid
            - latest_pruaccess_bid
        )

        if difference > PRICE_TOLERANCE:

            errors.append(
                "Current Prudential BID and "
                "latest PruAccess BID differ "
                f"by {difference:.8f}, "
                f"exceeding tolerance "
                f"{PRICE_TOLERANCE}."
            )

    # ========================================================
    # SUSPICIOUS DATA
    # ========================================================

    suspicious_locations = []

    for filename, data in loaded.items():

        suspicious_locations.extend(
            [
                f"{filename}:{location}"
                for location
                in find_suspicious_fields(
                    data
                )
            ]
        )

    if suspicious_locations:

        errors.append(
            "Suspicious synthetic/generated "
            "field names detected: "
            + ", ".join(
                suspicious_locations[:20]
            )
        )

    # ========================================================
    # FINAL
    # ========================================================

    status = (
        "passed"
        if not errors
        else "failed"
    )

    return {
        "excelRow":
            excel_row,

        "status":
            status,

        "fundName":
            prudential_name,

        "fundIdentifier":
            fund_identifier,

        "fundCode":
            fund_code,

        "pruAccessFundId":
            pruaccess_id,

        "observationCount":
            actual_count,

        "pageCount":
            expected_page_count,

        "errors":
            errors,

        "warnings":
            warnings,
    }


# ============================================================
# MAIN VALIDATION
# ============================================================

def main():

    print(
        "============================================================"
    )

    print(
        "PRUACCESS ALL-FUND HISTORICAL BID PRICE VALIDATION"
    )

    print(
        "============================================================"
    )

    # ========================================================
    # LOAD MASTER FILES
    # ========================================================

    excel_universe = (
        read_excel_universe()
    )

    run_summary = load_json(
        RUN_SUMMARY_FILE
    )

    all_funds = load_json(
        ALL_FUNDS_FILE
    )

    all_bid_history = load_json(
        ALL_BID_HISTORY_FILE
    )

    excel_count = len(
        excel_universe
    )

    print(
        f"\nExcel fund universe: "
        f"{excel_count}"
    )

    # ========================================================
    # EXTRACT SUCCESS/FAILURE RECORDS
    # ========================================================

    successful_records = (
        extract_successful_funds(
            run_summary,
            all_funds,
        )
    )

    failed_records = (
        extract_failed_funds(
            run_summary,
            all_funds,
        )
    )

    print(
        f"Run successful records: "
        f"{len(successful_records)}"
    )

    print(
        f"Run failed records: "
        f"{len(failed_records)}"
    )

    # ========================================================
    # SUCCESS MAP
    # ========================================================

    successful_by_row = {
        record["excelRow"]:
            record
        for record
        in successful_records
        if record.get(
            "excelRow"
        ) is not None
    }

    failed_by_row = {
        record["excelRow"]:
            record
        for record
        in failed_records
        if record.get(
            "excelRow"
        ) is not None
    }

    # ========================================================
    # CHECK DUPLICATE SUCCESS ROWS
    # ========================================================

    duplicate_success_rows = []

    seen_rows = set()

    for record in successful_records:

        row = record.get(
            "excelRow"
        )

        if row in seen_rows:

            duplicate_success_rows.append(
                row
            )

        seen_rows.add(
            row
        )

    global_errors = []

    if duplicate_success_rows:

        global_errors.append(
            "Duplicate successful Excel rows: "
            + ", ".join(
                str(row)
                for row
                in sorted(
                    set(
                        duplicate_success_rows
                    )
                )
            )
        )

    # ========================================================
    # CHECK DUPLICATE FUND IDENTIFIERS
    # ========================================================

    identifier_rows = {}

    for record in successful_records:

        identifier = clean_text(
            record.get(
                "fundIdentifier"
            )
        )

        if not identifier:

            continue

        identifier_rows.setdefault(
            identifier,
            [],
        ).append(
            record.get(
                "excelRow"
            )
        )

    duplicate_identifiers = {
        identifier:
            rows
        for identifier, rows
        in identifier_rows.items()
        if len(rows) > 1
    }

    if duplicate_identifiers:

        global_errors.append(
            "Duplicate fund identifiers detected: "
            + json.dumps(
                duplicate_identifiers,
                ensure_ascii=False,
            )
        )

    # ========================================================
    # CHECK MASTER ACCOUNTING
    # ========================================================

    excel_rows = {
        fund["excelRow"]
        for fund
        in excel_universe
    }

    successful_rows = set(
        successful_by_row
    )

    failed_rows = set(
        failed_by_row
    )

    missing_rows = sorted(
        excel_rows
        - successful_rows
        - failed_rows
    )

    if missing_rows:

        global_errors.append(
            "Excel funds missing from both "
            "successful and failed records: "
            + ", ".join(
                str(row)
                for row
                in missing_rows
            )
        )

    extra_success_rows = sorted(
        successful_rows
        - excel_rows
    )

    if extra_success_rows:

        global_errors.append(
            "Successful records contain Excel "
            "rows not present in master universe: "
            + ", ".join(
                str(row)
                for row
                in extra_success_rows
            )
        )

    extra_failed_rows = sorted(
        failed_rows
        - excel_rows
    )

    if extra_failed_rows:

        global_errors.append(
            "Failed records contain Excel "
            "rows not present in master universe: "
            + ", ".join(
                str(row)
                for row
                in extra_failed_rows
            )
        )

    # ========================================================
    # INDIVIDUAL FUND VALIDATION
    # ========================================================

    fund_results = []

    print(
        "\nValidating individual funds..."
    )

    for index, excel_fund in enumerate(
        excel_universe,
        start=1,
    ):

        row = excel_fund[
            "excelRow"
        ]

        print(
            "\n"
            + "=" * 70
        )

        print(
            f"[{index}/{excel_count}] "
            f"Excel row {row}"
        )

        print(
            "=" * 70
        )

        successful_summary = (
            successful_by_row.get(
                row
            )
        )

        if successful_summary is None:

            failure = failed_by_row.get(
                row
            )

            if failure:

                result = {
                    "excelRow":
                        row,

                    "status":
                        "failed_extraction",

                    "errors": [
                        "Fund extraction failed."
                    ],

                    "warnings": [],

                    "extractionError":
                        failure.get(
                            "error"
                        ),
                }

            else:

                result = {
                    "excelRow":
                        row,

                    "status":
                        "missing",

                    "errors": [
                        "Fund is neither successful "
                        "nor recorded as failed."
                    ],

                    "warnings": [],
                }

            fund_results.append(
                result
            )

            print(
                f"STATUS: "
                f"{result['status']}"
            )

            continue

        result = validate_fund(
            excel_fund,
            successful_summary,
        )

        fund_results.append(
            result
        )

        print(
            f"Fund: "
            f"{result.get('fundName', '-')}"
        )

        print(
            f"Observations: "
            f"{result.get('observationCount', '-')}"
        )

        print(
            f"Pages: "
            f"{result.get('pageCount', '-')}"
        )

        print(
            f"Status: "
            f"{result['status']}"
        )

        if result.get(
            "errors"
        ):

            print(
                "\nErrors:"
            )

            for error in result[
                "errors"
            ]:

                print(
                    f" - {error}"
                )

        if result.get(
            "warnings"
        ):

            print(
                "\nWarnings:"
            )

            for warning in result[
                "warnings"
            ]:

                print(
                    f" - {warning}"
                )

    # ========================================================
    # CONSOLIDATED BID HISTORY
    # ========================================================

    print(
        "\n"
        "============================================================"
    )

    print(
        "VALIDATING CONSOLIDATED BID HISTORY"
    )

    print(
        "============================================================"
    )

    consolidated_errors = []

    if not isinstance(
        all_bid_history,
        dict,
    ):

        consolidated_errors.append(
            "all_bid_history.json is not an object."
        )

    else:

        consolidated_source = normalize_text(
            all_bid_history.get(
                "source"
            )
        )

        consolidated_price_type = normalize_text(
            all_bid_history.get(
                "priceType"
            )
        )

        if consolidated_source != "pruaccess":

            consolidated_errors.append(
                "Consolidated source is not PruAccess."
            )

        if consolidated_price_type != "bid":

            consolidated_errors.append(
                "Consolidated price type is not BID."
            )

        consolidated_funds = (
            all_bid_history.get(
                "funds"
            )
        )

        if not isinstance(
            consolidated_funds,
            dict,
        ):

            consolidated_errors.append(
                "Consolidated funds is not an object."
            )

            consolidated_funds = {}

        expected_consolidated_count = (
            len(
                successful_records
            )
        )

        actual_consolidated_count = len(
            consolidated_funds
        )

        if (
            actual_consolidated_count
            != expected_consolidated_count
        ):

            consolidated_errors.append(
                "Consolidated fund count does "
                "not match successful fund count."
            )

        consolidated_observation_total = 0

        for identifier, fund_history in (
            consolidated_funds.items()
        ):

            if not isinstance(
                fund_history,
                dict,
            ):

                consolidated_errors.append(
                    f"Consolidated fund "
                    f"{identifier} is not an object."
                )

                continue

            if normalize_text(
                fund_history.get(
                    "source"
                )
            ) != "pruaccess":

                consolidated_errors.append(
                    f"Consolidated fund "
                    f"{identifier} source is "
                    "not PruAccess."
                )

            if normalize_text(
                fund_history.get(
                    "priceType"
                )
            ) != "bid":

                consolidated_errors.append(
                    f"Consolidated fund "
                    f"{identifier} price type "
                    "is not BID."
                )

            observations = (
                fund_history.get(
                    "observations"
                )
            )

            if not isinstance(
                observations,
                list,
            ):

                consolidated_errors.append(
                    f"Consolidated fund "
                    f"{identifier} observations "
                    "is not a list."
                )

                continue

            consolidated_observation_total += (
                len(observations)
            )

    # ========================================================
    # TOTAL OBSERVATIONS
    # ========================================================

    expected_total = sum(
        result.get(
            "observationCount",
            0,
        )
        for result
        in fund_results
        if result.get(
            "status"
        ) == "passed"
    )

    actual_total = 0

    if isinstance(
        all_bid_history,
        dict,
    ):

        actual_total = all_bid_history.get(
            "totalHistoricalBidObservations",
            0,
        )

    try:

        actual_total = int(
            actual_total
        )

    except (
        ValueError,
        TypeError,
    ):

        consolidated_errors.append(
            "Consolidated totalHistoricalBidObservations "
            "is invalid."
        )

        actual_total = 0

    if (
        expected_total
        != actual_total
    ):

        consolidated_errors.append(
            "Consolidated total historical BID "
            "observations does not match "
            "validated successful funds."
        )

    # ========================================================
    # OVERALL COUNTS
    # ========================================================

    passed_count = sum(
        1
        for result
        in fund_results
        if result.get(
            "status"
        ) == "passed"
    )

    failed_validation_count = sum(
        1
        for result
        in fund_results
        if result.get(
            "status"
        ) == "failed"
    )

    extraction_failed_count = sum(
        1
        for result
        in fund_results
        if result.get(
            "status"
        ) == "failed_extraction"
    )

    missing_count = sum(
        1
        for result
        in fund_results
        if result.get(
            "status"
        ) == "missing"
    )

    # ========================================================
    # RUN SUMMARY COUNT CROSS-CHECK
    # ========================================================

    run_summary_errors = []

    if isinstance(
        run_summary,
        dict,
    ):

        reported_universe = (
            run_summary.get(
                "fundUniverseCount"
            )
        )

        reported_success = (
            run_summary.get(
                "successfulFundCount"
            )
        )

        reported_failed = (
            run_summary.get(
                "failedFundCount"
            )
        )

        try:

            if (
                reported_universe is not None
                and int(
                    reported_universe
                )
                != excel_count
            ):

                run_summary_errors.append(
                    "run_summary fundUniverseCount "
                    "does not match Excel universe."
                )

        except (
            ValueError,
            TypeError,
        ):

            run_summary_errors.append(
                "run_summary fundUniverseCount "
                "is invalid."
            )

        try:

            if (
                reported_success is not None
                and int(
                    reported_success
                )
                != len(
                    successful_records
                )
            ):

                run_summary_errors.append(
                    "run_summary successfulFundCount "
                    "does not match normalized "
                    "successful records."
                )

        except (
            ValueError,
            TypeError,
        ):

            run_summary_errors.append(
                "run_summary successfulFundCount "
                "is invalid."
            )

        try:

            if (
                reported_failed is not None
                and int(
                    reported_failed
                )
                != len(
                    failed_records
                )
            ):

                run_summary_errors.append(
                    "run_summary failedFundCount "
                    "does not match normalized "
                    "failed records."
                )

        except (
            ValueError,
            TypeError,
        ):

            run_summary_errors.append(
                "run_summary failedFundCount "
                "is invalid."
            )

    # ========================================================
    # OVERALL STATUS
    # ========================================================

    overall_errors = []

    overall_errors.extend(
        global_errors
    )

    overall_errors.extend(
        consolidated_errors
    )

    overall_errors.extend(
        run_summary_errors
    )

    if passed_count != excel_count:

        overall_errors.append(
            "Not every Excel fund passed "
            "validation."
        )

    if failed_validation_count:

        overall_errors.append(
            f"{failed_validation_count} fund(s) "
            "failed individual validation."
        )

    if extraction_failed_count:

        overall_errors.append(
            f"{extraction_failed_count} fund(s) "
            "failed extraction."
        )

    if missing_count:

        overall_errors.append(
            f"{missing_count} fund(s) "
            "are missing from extraction accounting."
        )

    overall_status = (
        "passed"
        if (
            not overall_errors
            and passed_count == excel_count
        )
        else "failed"
    )

    # ========================================================
    # VALIDATION OUTPUT
    # ========================================================

    validation = {
        "status":
            overall_status,

        "validatedAtUtc":
            datetime.utcnow().isoformat()
            + "Z",

        "masterSource":
            str(
                EXCEL_FILE
            ),

        "excelFundUniverseCount":
            excel_count,

        "successfulFundRecordCount":
            len(
                successful_records
            ),

        "failedFundRecordCount":
            len(
                failed_records
            ),

        "passedFundCount":
            passed_count,

        "failedValidationFundCount":
            failed_validation_count,

        "extractionFailedFundCount":
            extraction_failed_count,

        "missingFundCount":
            missing_count,

        "priceType":
            "BID",

        "pageSize":
            PAGE_SIZE,

        "priceTolerance":
            PRICE_TOLERANCE,

        "maximumBidObservationGapDays":
            MAX_BID_DATE_GAP_DAYS,

        "totalValidatedHistoricalBidObservations":
            expected_total,

        "totalConsolidatedHistoricalBidObservations":
            actual_total,

        "globalErrors":
            overall_errors,

        "funds":
            fund_results,
    }

    save_json(
        VALIDATION_FILE,
        validation,
    )

    # ========================================================
    # FINAL REPORT
    # ========================================================

    print(
        "\n\n"
        "============================================================"
    )

    print(
        "VALIDATION COMPLETE"
    )

    print(
        "============================================================"
    )

    print(
        f"Excel fund universe: "
        f"{excel_count}"
    )

    print(
        f"Successful extraction records: "
        f"{len(successful_records)}"
    )

    print(
        f"Failed extraction records: "
        f"{len(failed_records)}"
    )

    print(
        f"Passed validation: "
        f"{passed_count}"
    )

    print(
        f"Failed validation: "
        f"{failed_validation_count}"
    )

    print(
        f"Extraction failures: "
        f"{extraction_failed_count}"
    )

    print(
        f"Missing: "
        f"{missing_count}"
    )

    print(
        f"Validated BID observations: "
        f"{expected_total}"
    )

    print(
        f"Consolidated BID observations: "
        f"{actual_total}"
    )

    print(
        f"Maximum allowed consecutive BID "
        f"observation gap: "
        f"{MAX_BID_DATE_GAP_DAYS} calendar days"
    )

    if overall_errors:

        print(
            "\nGLOBAL ERRORS:"
        )

        for error in overall_errors:

            print(
                f" - {error}"
            )

    print(
        "\nIndividual fund results:"
    )

    for result in fund_results:

        print(
            f" - Row "
            f"{result['excelRow']}: "
            f"{result['status']}"
        )

        for error in result.get(
            "errors",
            [],
        ):

            print(
                f"     ERROR: {error}"
            )

        for warning in result.get(
            "warnings",
            [],
        ):

            print(
                f"     WARNING: {warning}"
            )

    print(
        "\nValidation file:"
    )

    print(
        VALIDATION_FILE
    )

    if overall_status == "passed":

        print(
            "\n============================================================"
        )

        print(
            "VALIDATION PASSED"
        )

        print(
            "============================================================"
        )

    else:

        print(
            "\n============================================================"
        )

        print(
            "VALIDATION FAILED"
        )

        print(
            "============================================================"
        )

        # GitHub Actions must fail when validation fails.
        raise SystemExit(
            1
        )


if __name__ == "__main__":
    main()
