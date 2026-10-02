#!/usr/bin/env python3

"""
VGrat FMS - AI Fund Exposure Research

============================================================
PURPOSE
============================================================

Research every PRULink fund in Funds Links.xlsx and identify:

    - Up to 3 geographic exposures
    - Up to 3 sector exposures

Only category NAMES are returned.

No percentages are stored.

The allowed geography and sector vocabulary is loaded from:

    scripts/exposure_categories.txt

The AI may map source terminology to the canonical
categories in that file.

The AI may NOT invent categories.

============================================================
INPUT
============================================================

Funds Links.xlsx

Column A:
    Prudential URL

Column B:
    Exact Prudential / PruAccess fund name

============================================================
OUTPUT
============================================================

data/fund_exposure.json

============================================================
SOURCE PRIORITY
============================================================

1. Prudential source from Funds Links.xlsx
2. Official Prudential documents
3. Official underlying investment manager documents
4. Reputable secondary sources

============================================================
AI
============================================================

Default model:

    Qwen/Qwen3-1.7B

The model runs locally on GitHub Actions.

No commercial AI API is required.

============================================================
IMPORTANT OUTPUT RULES
============================================================

Maximum 3 geography categories.

Maximum 3 sector categories.

If only one reliable category exists:
    return one.

If only two reliable categories exist:
    return two.

If no reliable category exists:
    return [].

No percentages.

No invented values.

No inference based solely on fund name.

Only canonical category names from
exposure_categories.txt may be returned.
"""

from __future__ import annotations

import argparse
import io
import json
import os
import re
import sys
import time
from pathlib import Path
from typing import Any
from urllib.parse import quote_plus, urlparse

import requests
from bs4 import BeautifulSoup
from openpyxl import load_workbook
from pypdf import PdfReader
from transformers import AutoModelForCausalLM, AutoTokenizer


# ============================================================
# PATHS
# ============================================================

SCRIPT_DIR = Path(__file__).resolve().parent

REPOSITORY_ROOT = SCRIPT_DIR.parent

CATEGORIES_FILE = (
    SCRIPT_DIR / "exposure_categories.txt"
)

EXCEL_PATH = (
    REPOSITORY_ROOT / "Funds Links.xlsx"
)

OUTPUT_PATH = (
    REPOSITORY_ROOT
    / "data"
    / "fund_exposure.json"
)


# ============================================================
# CONFIGURATION
# ============================================================

MODEL_NAME = os.getenv(
    "HF_MODEL",
    "Qwen/Qwen3-1.7B",
)

REQUEST_TIMEOUT = int(
    os.getenv(
        "REQUEST_TIMEOUT",
        "30",
    )
)

MAX_SEARCH_RESULTS = int(
    os.getenv(
        "MAX_SEARCH_RESULTS",
        "8",
    )
)

MAX_DOCUMENTS_PER_FUND = int(
    os.getenv(
        "MAX_DOCUMENTS_PER_FUND",
        "5",
    )
)

MAX_PAGE_CHARS = 18000

MAX_PDF_CHARS = 24000

USER_AGENT = (
    "Mozilla/5.0 "
    "(Windows NT 10.0; Win64; x64) "
    "AppleWebKit/537.36 "
    "(KHTML, like Gecko) "
    "Chrome/154.0 Safari/537.36 "
    "VGrat-FMS-Fund-Exposure-Research"
)


SESSION = requests.Session()

SESSION.headers.update(
    {
        "User-Agent": USER_AGENT,
        "Accept-Language": "en-SG,en;q=0.9",
    }
)


# ============================================================
# TEXT HELPERS
# ============================================================

def clean_text(
    value: Any,
) -> str:

    if value is None:
        return ""

    return re.sub(
        r"\s+",
        " ",
        str(value),
    ).strip()


# ============================================================
# CATEGORY FILE
# ============================================================

def load_categories(
    path: Path,
) -> dict[str, list[str]]:

    if not path.exists():
        raise FileNotFoundError(
            "Exposure category file not found:\n"
            f"{path}"
        )

    geography: list[str] = []

    sector: list[str] = []

    current_section: str | None = None

    with path.open(
        "r",
        encoding="utf-8",
    ) as file:

        for raw_line in file:

            line = raw_line.strip()

            if not line:
                continue

            if line.startswith("#"):
                continue

            upper = line.upper()

            if upper == "[GEOGRAPHY]":
                current_section = "geography"
                continue

            if upper == "[SECTOR]":
                current_section = "sector"
                continue

            if current_section == "geography":
                if line not in geography:
                    geography.append(line)

            elif current_section == "sector":
                if line not in sector:
                    sector.append(line)

    if not geography:
        raise ValueError(
            "No geography categories found in "
            f"{path}"
        )

    if not sector:
        raise ValueError(
            "No sector categories found in "
            f"{path}"
        )

    return {
        "geography": geography,
        "sector": sector,
    }


# ============================================================
# EXCEL FUND UNIVERSE
# ============================================================

def load_funds_from_excel(
    path: Path,
) -> list[dict[str, str]]:

    if not path.exists():
        raise FileNotFoundError(
            "Funds Links.xlsx not found:\n"
            f"{path}"
        )

    workbook = load_workbook(
        filename=path,
        read_only=True,
        data_only=True,
    )

    sheet = workbook.active

    funds: list[dict[str, str]] = []

    for row_number, row in enumerate(
        sheet.iter_rows(
            min_row=2,
            values_only=True,
        ),
        start=2,
    ):

        url = clean_text(
            row[0]
            if len(row) > 0
            else ""
        )

        fund_name = clean_text(
            row[1]
            if len(row) > 1
            else ""
        )

        if not url and not fund_name:
            continue

        if not fund_name:
            fund_name = url

        funds.append(
            {
                "row": str(row_number),
                "fundName": fund_name,
                "prudentialUrl": url,
            }
        )

    workbook.close()

    return funds


# ============================================================
# HTTP
# ============================================================

def fetch_url(
    url: str,
) -> tuple[str, bytes] | None:

    try:

        response = SESSION.get(
            url,
            timeout=REQUEST_TIMEOUT,
            allow_redirects=True,
        )

        if response.status_code >= 400:
            return None

        content_type = (
            response.headers
            .get(
                "content-type",
                "",
            )
            .lower()
        )

        return (
            content_type,
            response.content,
        )

    except Exception as exc:

        print(
            f"        fetch failed: "
            f"{url} -> {exc}"
        )

        return None


# ============================================================
# PDF EXTRACTION
# ============================================================

def extract_pdf_text(
    content: bytes,
) -> str:

    try:

        reader = PdfReader(
            io.BytesIO(content)
        )

        pages: list[str] = []

        total_length = 0

        for page in reader.pages:

            try:
                text = (
                    page.extract_text()
                    or ""
                )
            except Exception:
                text = ""

            if text:

                pages.append(text)

                total_length += len(text)

            if total_length >= MAX_PDF_CHARS:
                break

        return "\n".join(
            pages
        )[:MAX_PDF_CHARS]

    except Exception:

        return ""


# ============================================================
# HTML EXTRACTION
# ============================================================

def extract_html_text(
    content: bytes,
) -> str:

    try:

        soup = BeautifulSoup(
            content,
            "html.parser",
        )

        for tag in soup(
            [
                "script",
                "style",
                "noscript",
                "svg",
                "nav",
                "footer",
                "header",
            ]
        ):
            tag.decompose()

        text = soup.get_text(
            "\n",
            strip=True,
        )

        return text[:MAX_PAGE_CHARS]

    except Exception:

        return ""


# ============================================================
# DOCUMENT EXTRACTION
# ============================================================

def extract_document(
    url: str,
) -> str:

    result = fetch_url(
        url
    )

    if not result:
        return ""

    content_type, content = result

    clean_url = (
        url.lower()
        .split("?")[0]
    )

    if (
        "pdf" in content_type
        or clean_url.endswith(".pdf")
    ):

        return extract_pdf_text(
            content
        )

    return extract_html_text(
        content
    )


# ============================================================
# SEARCH
# ============================================================

def ddg_search(
    query: str,
    max_results: int = MAX_SEARCH_RESULTS,
) -> list[dict[str, str]]:

    """
    DuckDuckGo HTML search.

    No API key required.
    """

    search_url = (
        "https://html.duckduckgo.com/html/?q="
        + quote_plus(query)
    )

    try:

        response = SESSION.get(
            search_url,
            timeout=REQUEST_TIMEOUT,
        )

        if response.status_code >= 400:
            return []

        soup = BeautifulSoup(
            response.text,
            "html.parser",
        )

        results: list[
            dict[str, str]
        ] = []

        for result in soup.select(
            ".result"
        ):

            link = result.select_one(
                ".result__a"
            )

            if not link:
                continue

            href = link.get(
                "href"
            )

            if not href:
                continue

            title = clean_text(
                link.get_text(
                    " ",
                    strip=True,
                )
            )

            snippet_element = (
                result.select_one(
                    ".result__snippet"
                )
            )

            snippet = ""

            if snippet_element:

                snippet = clean_text(
                    snippet_element.get_text(
                        " ",
                        strip=True,
                    )
                )

            results.append(
                {
                    "title": title,
                    "url": href,
                    "snippet": snippet,
                }
            )

            if (
                len(results)
                >= max_results
            ):
                break

        return results

    except Exception as exc:

        print(
            f"        search failed: {exc}"
        )

        return []


# ============================================================
# URL FILTERING
# ============================================================

def is_useful_url(
    url: str,
) -> bool:

    lowered = url.lower()

    blocked_domains = [
        "facebook.com",
        "instagram.com",
        "youtube.com",
        "linkedin.com",
        "twitter.com",
        "x.com",
        "reddit.com",
        "tiktok.com",
    ]

    return not any(
        domain in lowered
        for domain in blocked_domains
    )


def source_priority(
    url: str,
) -> int:

    host = (
        urlparse(url)
        .netloc
        .lower()
    )

    if (
        "prudential.com.sg"
        in host
    ):
        return 1

    official_manager_domains = [
        "eastspring.com",
        "abrdn.com",
        "fullertonfund.com",
        "hsbc.com",
        "blackrock.com",
        "schroders.com",
        "fidelity.com",
        "jpmorgan.com",
        "jpmorganassetmanagement.com",
        "amundi.com",
        "invesco.com",
        "franklintempleton.com",
        "manulifeim.com",
        "pimco.com",
    ]

    if any(
        domain in host
        for domain in official_manager_domains
    ):
        return 2

    return 3


# ============================================================
# SEARCH QUERIES
# ============================================================

def build_search_queries(
    fund_name: str,
) -> list[str]:

    return [
        (
            f'"{fund_name}" '
            "sector allocation geography"
        ),
        (
            f'"{fund_name}" '
            '"geographic allocation" '
            '"sector allocation"'
        ),
        (
            f'"{fund_name}" '
            "factsheet sector geography"
        ),
        (
            f'"{fund_name}" '
            '"underlying fund" factsheet'
        ),
        (
            f'"{fund_name}" '
            "portfolio country sector"
        ),
    ]


# ============================================================
# FUND RESEARCH
# ============================================================

def research_fund(
    fund: dict[str, str],
) -> list[dict[str, Any]]:

    fund_name = fund[
        "fundName"
    ]

    prudential_url = fund[
        "prudentialUrl"
    ]

    print(
        f"      Researching: "
        f"{fund_name}"
    )

    candidates: list[
        dict[str, Any]
    ] = []

    # --------------------------------------------------------
    # Prudential URL from Excel
    # --------------------------------------------------------

    if prudential_url:

        candidates.append(
            {
                "url": prudential_url,
                "title": (
                    "Prudential source "
                    "from Funds Links.xlsx"
                ),
                "snippet": "",
                "priority": 1,
            }
        )

    seen_urls = set()

    if prudential_url:
        seen_urls.add(
            prudential_url
        )

    # --------------------------------------------------------
    # Search
    # --------------------------------------------------------

    for query in build_search_queries(
        fund_name
    ):

        print(
            f"        Search: {query}"
        )

        results = ddg_search(
            query
        )

        for result in results:

            url = result[
                "url"
            ]

            if not is_useful_url(
                url
            ):
                continue

            if url in seen_urls:
                continue

            seen_urls.add(
                url
            )

            candidates.append(
                {
                    **result,
                    "priority": (
                        source_priority(
                            url
                        )
                    ),
                }
            )

            if (
                len(candidates)
                >= MAX_DOCUMENTS_PER_FUND
            ):
                break

        if (
            len(candidates)
            >= MAX_DOCUMENTS_PER_FUND
        ):
            break

        time.sleep(
            0.5
        )

    # --------------------------------------------------------
    # Prioritise sources
    # --------------------------------------------------------

    candidates.sort(
        key=lambda item: item.get(
            "priority",
            99,
        )
    )

    # --------------------------------------------------------
    # Extract documents
    # --------------------------------------------------------

    documents: list[
        dict[str, Any]
    ] = []

    for candidate in candidates[
        :MAX_DOCUMENTS_PER_FUND
    ]:

        url = candidate[
            "url"
        ]

        print(
            f"        Document: {url}"
        )

        text = extract_document(
            url
        )

        if not text:
            continue

        documents.append(
            {
                "url": url,
                "title": candidate.get(
                    "title",
                    "",
                ),
                "snippet": candidate.get(
                    "snippet",
                    "",
                ),
                "text": text,
                "priority": candidate.get(
                    "priority",
                    99,
                ),
            }
        )

    return documents


# ============================================================
# AI PROMPT
# ============================================================

def build_ai_prompt(
    fund_name: str,
    documents: list[
        dict[str, Any]
    ],
    categories: dict[
        str,
        list[str],
    ],
) -> str:

    geography_categories = "\n".join(
        f"- {item}"
        for item in categories[
            "geography"
        ]
    )

    sector_categories = "\n".join(
        f"- {item}"
        for item in categories[
            "sector"
        ]
    )

    source_text_parts: list[str] = []

    for index, document in enumerate(
        documents,
        start=1,
    ):

        source_text_parts.append(
            f"""
SOURCE {index}

TITLE:
{document["title"]}

URL:
{document["url"]}

CONTENT:
{document["text"]}
""".strip()
        )

    source_text = (
        "\n\n".join(
            source_text_parts
        )
    )

    return f"""
You are a financial-document research assistant.

You are researching this fund:

{fund_name}

Your task is to identify the fund's most significant:

1. Geographic exposures
2. Sector exposures

You must use ONLY the supplied source documents.

============================================================
CRITICAL RULES
============================================================

Do NOT use your general knowledge.

Do NOT guess.

Do NOT invent.

Do NOT infer an exposure merely because of the fund name.

Do NOT infer that a country or sector is present unless
the source provides evidence for it.

Do NOT return percentages.

Do NOT return weights.

Do NOT return explanations.

Return names only.

Maximum 3 geography categories.

Maximum 3 sector categories.

If only one valid category is supported:
return one.

If only two valid categories are supported:
return two.

If none are reliably supported:
return [].

============================================================
CANONICAL GEOGRAPHY VOCABULARY
============================================================

You may ONLY return one of these exact names:

{geography_categories}

============================================================
CANONICAL SECTOR VOCABULARY
============================================================

You may ONLY return one of these exact names:

{sector_categories}

============================================================
MAPPING
============================================================

You may map source terminology to the canonical
vocabulary.

Examples:

"USA" -> "United States"

"US" -> "United States"

"UK" -> "United Kingdom"

"Britain" -> "United Kingdom"

"Information Technology" -> "Technology"

"IT" -> "Technology"

"Technology" -> "Technology"

"Communication" or "Telecom" may map to the closest
canonical category only when the source clearly supports
that classification.

"Banking" or "Banks" -> "Financials"

"Insurance" -> "Financials"

"Pharmaceuticals" -> "Healthcare"

"Property" -> "Real Estate"

Only perform a mapping when it is clearly supported.

============================================================
RANKING
============================================================

If the source provides percentages or weights, use them
ONLY to determine which exposures are the largest.

DO NOT output those numbers.

Return the largest exposures first.

If the source explicitly identifies a Top 3 list, use that.

If the source provides an allocation table, identify the
largest valid categories from that table.

If the source does not provide enough information to rank
the categories, return only categories that can be reliably
identified.

============================================================
IMPORTANT
============================================================

Do not treat individual company names as sectors.

Do not treat individual company names as geographies.

Do not convert company names into sectors unless the source
explicitly provides the relevant sector classification.

Do not create "Other" merely because the source contains
other holdings.

Only return "Other" if "Other" is explicitly reported as a
relevant exposure in the source.

Do not return "Global" merely because the fund is globally
invested.

Use "Global" only when the source explicitly presents
Global as the geographic classification.

============================================================
OUTPUT
============================================================

Return ONLY valid JSON.

Exactly this structure:

{{
  "geography": [],
  "sector": []
}}

Example:

{{
  "geography": [
    "United States",
    "Japan",
    "United Kingdom"
  ],
  "sector": [
    "Technology",
    "Financials",
    "Healthcare"
  ]
}}

No markdown.

No commentary.

No percentages.

No additional fields.

============================================================
SOURCE DOCUMENTS
============================================================

{source_text}
""".strip()


# ============================================================
# JSON EXTRACTION
# ============================================================

def extract_json_object(
    text: str,
) -> dict[str, Any] | None:

    text = text.strip()

    text = re.sub(
        r"^```(?:json)?",
        "",
        text,
        flags=re.IGNORECASE,
    )

    text = re.sub(
        r"```$",
        "",
        text,
    )

    text = text.strip()

    try:

        parsed = json.loads(
            text
        )

        if isinstance(
            parsed,
            dict,
        ):
            return parsed

    except Exception:
        pass

    match = re.search(
        r"\{.*\}",
        text,
        flags=re.DOTALL,
    )

    if not match:
        return None

    try:

        parsed = json.loads(
            match.group(0)
        )

        if isinstance(
            parsed,
            dict,
        ):
            return parsed

    except Exception:
        return None

    return None


# ============================================================
# VALIDATE AGAINST TXT VOCABULARY
# ============================================================

def validate_category_list(
    values: Any,
    allowed: list[str],
) -> list[str]:

    if not isinstance(
        values,
        list,
    ):
        return []

    allowed_lookup = {
        item.casefold(): item
        for item in allowed
    }

    output: list[str] = []

    for value in values:

        if not isinstance(
            value,
            str,
        ):
            continue

        value = clean_text(
            value
        )

        if not value:
            continue

        # Remove accidental percentage.
        value = re.sub(
            r"\s*\(?\d+(?:\.\d+)?%\)?",
            "",
            value,
        ).strip()

        canonical = (
            allowed_lookup.get(
                value.casefold()
            )
        )

        # Critical:
        # If the AI returned something that is not
        # in the TXT file, reject it.
        if canonical is None:
            continue

        if canonical not in output:
            output.append(
                canonical
            )

        if len(output) >= 3:
            break

    return output


def validate_ai_result(
    result: dict[str, Any] | None,
    categories: dict[
        str,
        list[str],
    ],
) -> dict[str, list[str]]:

    if not result:
        return {
            "geography": [],
            "sector": [],
        }

    geography = (
        validate_category_list(
            result.get(
                "geography",
                [],
            ),
            categories[
                "geography"
            ],
        )
    )

    sector = (
        validate_category_list(
            result.get(
                "sector",
                [],
            ),
            categories[
                "sector"
            ],
        )
    )

    return {
        "geography": geography,
        "sector": sector,
    }


# ============================================================
# AI MODEL
# ============================================================

def load_ai_model():

    print()
    print(
        "Loading local Hugging Face model:"
    )

    print(
        f"    {MODEL_NAME}"
    )

    tokenizer = (
        AutoTokenizer.from_pretrained(
            MODEL_NAME,
            trust_remote_code=True,
        )
    )

    model = (
        AutoModelForCausalLM.from_pretrained(
            MODEL_NAME,
            torch_dtype="auto",
            device_map="auto",
            trust_remote_code=True,
        )
    )

    return (
        tokenizer,
        model,
    )


# ============================================================
# PROCESS FUND
# ============================================================

def process_fund(
    fund: dict[str, str],
    tokenizer,
    model,
    categories: dict[
        str,
        list[str],
    ],
) -> dict[str, Any]:

    fund_name = fund[
        "fundName"
    ]

    documents = research_fund(
        fund
    )

    if not documents:

        print(
            "      No usable source documents."
        )

        return {
            "fundName": fund_name,
            "prudentialUrl": fund[
                "prudentialUrl"
            ],
            "geography": [],
            "sector": [],
            "status": "no_source",
            "sources": [],
        }

    prompt = build_ai_prompt(
        fund_name,
        documents,
        categories,
    )

    messages = [
        {
            "role": "system",
            "content": (
                "You are a strict financial "
                "document extraction system. "
                "Never invent information. "
                "Return only requested JSON."
            ),
        },
        {
            "role": "user",
            "content": prompt,
        },
    ]

    print(
        "      Running local AI..."
    )

    try:

        inputs = (
            tokenizer.apply_chat_template(
                messages,
                tokenize=True,
                add_generation_prompt=True,
                return_tensors="pt",
            )
        )

        inputs = inputs.to(
            model.device
        )

        generated = model.generate(
            inputs,
            max_new_tokens=180,
            do_sample=False,
        )

        generated_tokens = (
            generated[
                0
            ][
                inputs.shape[-1]:
            ]
        )

        response = tokenizer.decode(
            generated_tokens,
            skip_special_tokens=True,
        )

        print(
            "      AI response:"
        )

        print(
            response[:1000]
        )

        parsed = (
            extract_json_object(
                response
            )
        )

        exposures = (
            validate_ai_result(
                parsed,
                categories,
            )
        )

        if (
            exposures["geography"]
            or exposures["sector"]
        ):

            status = "success"

        else:

            status = (
                "no_exposure_found"
            )

        return {
            "fundName": fund_name,
            "prudentialUrl": fund[
                "prudentialUrl"
            ],
            "geography": exposures[
                "geography"
            ],
            "sector": exposures[
                "sector"
            ],
            "status": status,
            "sources": [
                {
                    "title": document[
                        "title"
                    ],
                    "url": document[
                        "url"
                    ],
                }
                for document in documents
            ],
        }

    except Exception as exc:

        print(
            f"      AI error: {exc}"
        )

        return {
            "fundName": fund_name,
            "prudentialUrl": fund[
                "prudentialUrl"
            ],
            "geography": [],
            "sector": [],
            "status": "ai_error",
            "error": str(exc),
            "sources": [
                {
                    "title": document[
                        "title"
                    ],
                    "url": document[
                        "url"
                    ],
                }
                for document in documents
            ],
        }


# ============================================================
# MAIN
# ============================================================

def main() -> int:

    parser = argparse.ArgumentParser(
        description=(
            "VGrat FMS AI fund geography "
            "and sector research."
        )
    )

    parser.add_argument(
        "--limit",
        type=int,
        default=0,
        help=(
            "Process only the first N funds. "
            "0 = all funds."
        ),
    )

    args = parser.parse_args()

    print(
        "=" * 72
    )

    print(
        "VGrat FMS - AI Fund Exposure Research"
    )

    print(
        "=" * 72
    )

    print(
        f"Repository: {REPOSITORY_ROOT}"
    )

    print(
        f"Funds file: {EXCEL_PATH}"
    )

    print(
        f"Category file: {CATEGORIES_FILE}"
    )

    print(
        f"Output: {OUTPUT_PATH}"
    )

    print(
        f"Model: {MODEL_NAME}"
    )

    # --------------------------------------------------------
    # Load categories
    # --------------------------------------------------------

    categories = load_categories(
        CATEGORIES_FILE
    )

    print()
    print(
        "Allowed geography categories:"
    )

    for item in categories[
        "geography"
    ]:
        print(
            f"    - {item}"
        )

    print()
    print(
        "Allowed sector categories:"
    )

    for item in categories[
        "sector"
    ]:
        print(
            f"    - {item}"
        )

    # --------------------------------------------------------
    # Load funds
    # --------------------------------------------------------

    funds = load_funds_from_excel(
        EXCEL_PATH
    )

    if args.limit > 0:

        funds = funds[
            :args.limit
        ]

    print()
    print(
        f"Funds loaded: {len(funds)}"
    )

    if not funds:

        print(
            "ERROR: No funds found."
        )

        return 1

    # --------------------------------------------------------
    # Load model
    # --------------------------------------------------------

    tokenizer, model = (
        load_ai_model()
    )

    # --------------------------------------------------------
    # Process
    # --------------------------------------------------------

    results: list[
        dict[str, Any]
    ] = []

    for index, fund in enumerate(
        funds,
        start=1,
    ):

        print()
        print(
            "-" * 72
        )

        print(
            f"[{index}/{len(funds)}] "
            f"{fund['fundName']}"
        )

        print(
            "-" * 72
        )

        result = process_fund(
            fund,
            tokenizer,
            model,
            categories,
        )

        results.append(
            result
        )

        print()
        print(
            "      Geography:"
        )

        print(
            f"        "
            f"{result.get('geography', [])}"
        )

        print(
            "      Sector:"
        )

        print(
            f"        "
            f"{result.get('sector', [])}"
        )

        print(
            "      Status:"
        )

        print(
            f"        "
            f"{result.get('status')}"
        )

    # --------------------------------------------------------
    # Summary
    # --------------------------------------------------------

    successful = sum(
        1
        for result in results
        if result.get(
            "status"
        ) == "success"
    )

    no_source = sum(
        1
        for result in results
        if result.get(
            "status"
        ) == "no_source"
    )

    no_exposure = sum(
        1
        for result in results
        if result.get(
            "status"
        ) == "no_exposure_found"
    )

    ai_errors = sum(
        1
        for result in results
        if result.get(
            "status"
        ) == "ai_error"
    )

    # --------------------------------------------------------
    # Output
    # --------------------------------------------------------

    output = {
        "generatedAtUTC": time.strftime(
            "%Y-%m-%dT%H:%M:%SZ",
            time.gmtime(),
        ),
        "model": MODEL_NAME,
        "categoryFile": str(
            CATEGORIES_FILE
            .relative_to(
                REPOSITORY_ROOT
            )
        ),
        "rules": {
            "maximumGeography": 3,
            "maximumSector": 3,
            "percentagesStored": False,
            "inferredValuesAllowed": False,
            "categoriesMustExistInTxt": True,
        },
        "summary": {
            "fundsProcessed": len(
                results
            ),
            "successful": successful,
            "noSource": no_source,
            "noReliableExposure": no_exposure,
            "aiErrors": ai_errors,
        },
        "funds": results,
    }

    OUTPUT_PATH.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    with OUTPUT_PATH.open(
        "w",
        encoding="utf-8",
    ) as file:

        json.dump(
            output,
            file,
            indent=2,
            ensure_ascii=False,
        )

        file.write(
            "\n"
        )

    # --------------------------------------------------------
    # Final validation
    # --------------------------------------------------------

    with OUTPUT_PATH.open(
        "r",
        encoding="utf-8",
    ) as file:

        validation_data = json.load(
            file
        )

    for fund in validation_data[
        "funds"
    ]:

        geography = fund.get(
            "geography",
            [],
        )

        sector = fund.get(
            "sector",
            [],
        )

        if len(geography) > 3:
            raise ValueError(
                "Geography contains more "
                "than 3 categories for "
                f"{fund['fundName']}"
            )

        if len(sector) > 3:
            raise ValueError(
                "Sector contains more "
                "than 3 categories for "
                f"{fund['fundName']}"
            )

        for category in geography:

            if category not in categories[
                "geography"
            ]:

                raise ValueError(
                    "Invalid geography "
                    f"category: {category}"
                )

        for category in sector:

            if category not in categories[
                "sector"
            ]:

                raise ValueError(
                    "Invalid sector "
                    f"category: {category}"
                )

    # --------------------------------------------------------
    # Summary
    # --------------------------------------------------------

    print()
    print(
        "=" * 72
    )

    print(
        "COMPLETE"
    )

    print(
        "=" * 72
    )

    print(
        f"Funds processed: "
        f"{len(results)}"
    )

    print(
        f"Successful: "
        f"{successful}"
    )

    print(
        f"No source: "
        f"{no_source}"
    )

    print(
        f"No reliable exposure: "
        f"{no_exposure}"
    )

    print(
        f"AI errors: "
        f"{ai_errors}"
    )

    print(
        f"Output: "
        f"{OUTPUT_PATH}"
    )

    return 0


if __name__ == "__main__":
    sys.exit(
        main()
    )
